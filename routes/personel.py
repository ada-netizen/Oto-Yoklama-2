from fastapi import APIRouter, UploadFile, File, Form, HTTPException, BackgroundTasks
from fastapi.responses import JSONResponse, FileResponse
from pydantic import BaseModel
from typing import List, Optional
import os, json, re, tempfile, shutil, uuid
import pandas as pd
from datetime import datetime, timedelta
from veritabani import VeritabaniYoneticisi
from sistem_motoru import SistemMotoru
from excel_motoru import ExcelMotoru
from pdf_motoru import PDFYoneticisi
import logging
from araclar import VeriAraclari
from dependencies import db, yollar, ayarlar, islem_logla, create_job, update_job, get_job, GLOBAL_ISLEMLER, islem_durumlari
from utils import *
from utils import _excel_sutunu_bul
from utils import _personel_grup_tahmin_et
from mevzuat_kurallari import ogretmen_mi, mevzuat_bilgisi_al, tum_mevzuat_ozeti, PERSONEL_DURUMLARI_REHBERI

router = APIRouter()

@router.get("/personel-durumlari")
def personel_durumlari_getir():
    return {"basarili": True, "durumlar": PERSONEL_DURUMLARI_REHBERI}

@router.get("/personeller")
def personelleri_getir():
    try:
        db.cursor.execute("""
            SELECT rowid, ad_soyad, unvan, brans, gorev, grup, tc_no, il_ilce, kurum_adi, 
                   kurum_kodu, kurumda_baslama, kariyer, kademe_derece, ogrenim_durumu, 
                   cinsiyet, kan_grubu, dogum_tarihi, goreve_baslama, kurum_sicil_no, 
                   emekli_sicil_no, arsiv_no, telefon, eposta, notlar, aktif 
            FROM personel
        """)
        personel_listesi = []
        for r in db.cursor.fetchall():
            personel_listesi.append({
                'id': r[0],
                'ad': r[1],
                'ad_soyad': r[1],
                'unvan': r[2] or "-",
                'brans': r[3] or "-",
                'gorev': r[4] or "-",
                'grup': r[5] or "Kadrolu",
                'tc_no': r[6] or "",
                'il_ilce': r[7] or "",
                'kurum_adi': r[8] or "",
                'kurum_kodu': r[9] or "",
                'kurumda_baslama': r[10] or "",
                'kariyer': r[11] or "-",
                'kademe_derece': r[12] or "-",
                'ogrenim_durumu': r[13] or "-",
                'cinsiyet': r[14] or "",
                'kan_grubu': r[15] or "",
                'dogum_tarihi': r[16] or "",
                'goreve_baslama': r[17] or "",
                'kurum_sicil_no': r[18] or "",
                'emekli_sicil_no': r[19] or "",
                'arsiv_no': r[20] or "",
                'telefon': r[21] or "",
                'eposta': r[22] or "",
                'notlar': r[23] or "",
                'aktif': int(r[24]) if (len(r) > 24 and r[24] is not None) else 1
            })
        return {"personeller": personel_listesi}
    except Exception as e:
        logging.error(f"personelleri_getir hatasi: {e}")
        return {"personeller": []}


def _personelleri_veritabanina_kaydet(liste, kaynak_adi="Personel"):
    if not liste:
        return {"basarili": False, "mesaj": "Kaydedilecek personel verisi bulunamadı."}

    yedek_basarili, yedek_hatasi = SistemMotoru.yedek_al(db.db_yolu, ayarlari_al(), yollar["YEDEK"])
    if not yedek_basarili:
        return {"basarili": False, "mesaj": f"Aktarım öncesi yedek alınamadı: {yedek_hatasi}"}

    eklenen = 0
    guncellenen = 0
    try:
        db.cursor.execute("BEGIN TRANSACTION")
        for p in liste:
            ad = p.get('ad_soyad', '').strip()
            if not ad:
                continue
            tc = p.get('tc_no', '').strip()
            unvan = p.get('unvan', '-').strip() or '-'
            brans = p.get('brans', '-').strip() or '-'
            gorev = p.get('gorev', '-').strip() or '-'
            grup = p.get('grup', 'Kadrolu').strip() or 'Kadrolu'
            il_ilce = p.get('il_ilce', '').strip()
            kurum_adi = p.get('kurum_adi', '').strip()
            kurum_kodu = p.get('kurum_kodu', '').strip()
            kurumda_baslama = p.get('kurumda_baslama', '').strip()
            kariyer = p.get('kariyer', '-').strip() or '-'
            kademe_derece = p.get('kademe_derece', '-').strip() or '-'
            ogrenim_durumu = p.get('ogrenim_durumu', '-').strip() or '-'
            cinsiyet = p.get('cinsiyet', '').strip()
            kan_grubu = p.get('kan_grubu', '').strip()
            dogum_tarihi = p.get('dogum_tarihi', '').strip()
            goreve_baslama = p.get('goreve_baslama', '').strip()
            kurum_sicil_no = p.get('kurum_sicil_no', '').strip()
            emekli_sicil_no = p.get('emekli_sicil_no', '').strip()
            arsiv_no = p.get('arsiv_no', '').strip()
            telefon = p.get('telefon', '').strip()
            eposta = p.get('eposta', '').strip()
            notlar = p.get('notlar', '').strip()
            aktif = int(p.get('aktif', 1))

            mevcut_id = None
            if tc:
                db.cursor.execute("SELECT rowid FROM personel WHERE tc_no = ?", (tc,))
                row = db.cursor.fetchone()
                if row:
                    mevcut_id = row[0]
            if not mevcut_id and ad:
                db.cursor.execute("SELECT rowid FROM personel WHERE ad_soyad = ?", (ad,))
                row = db.cursor.fetchone()
                if row:
                    mevcut_id = row[0]

            if mevcut_id:
                db.cursor.execute("""
                    UPDATE personel SET
                        ad_soyad = COALESCE(NULLIF(?, ''), ad_soyad),
                        tc_no = COALESCE(NULLIF(?, ''), tc_no),
                        unvan = CASE WHEN ? != '-' AND ? != '' THEN ? ELSE unvan END,
                        brans = CASE WHEN ? != '-' AND ? != '' THEN ? ELSE brans END,
                        gorev = CASE WHEN ? != '-' AND ? != '' THEN ? ELSE gorev END,
                        grup = COALESCE(NULLIF(?, ''), grup),
                        il_ilce = COALESCE(NULLIF(?, ''), il_ilce),
                        kurum_adi = COALESCE(NULLIF(?, ''), kurum_adi),
                        kurum_kodu = COALESCE(NULLIF(?, ''), kurum_kodu),
                        kurumda_baslama = COALESCE(NULLIF(?, ''), kurumda_baslama),
                        kariyer = CASE WHEN ? != '-' AND ? != '' THEN ? ELSE kariyer END,
                        kademe_derece = CASE WHEN ? != '-' AND ? != '' THEN ? ELSE kademe_derece END,
                        ogrenim_durumu = CASE WHEN ? != '-' AND ? != '' THEN ? ELSE ogrenim_durumu END,
                        cinsiyet = COALESCE(NULLIF(?, ''), cinsiyet),
                        kan_grubu = COALESCE(NULLIF(?, ''), kan_grubu),
                        dogum_tarihi = COALESCE(NULLIF(?, ''), dogum_tarihi),
                        goreve_baslama = COALESCE(NULLIF(?, ''), goreve_baslama),
                        kurum_sicil_no = COALESCE(NULLIF(?, ''), kurum_sicil_no),
                        emekli_sicil_no = COALESCE(NULLIF(?, ''), emekli_sicil_no),
                        arsiv_no = COALESCE(NULLIF(?, ''), arsiv_no),
                        telefon = COALESCE(NULLIF(?, ''), telefon),
                        eposta = COALESCE(NULLIF(?, ''), eposta),
                        notlar = COALESCE(NULLIF(?, ''), notlar),
                        aktif = ?
                    WHERE rowid = ?
                """, (
                    ad, tc,
                    unvan, unvan, unvan,
                    brans, brans, brans,
                    gorev, gorev, gorev,
                    grup, il_ilce, kurum_adi, kurum_kodu, kurumda_baslama,
                    kariyer, kariyer, kariyer,
                    kademe_derece, kademe_derece, kademe_derece,
                    ogrenim_durumu, ogrenim_durumu, ogrenim_durumu,
                    cinsiyet, kan_grubu, dogum_tarihi, goreve_baslama,
                    kurum_sicil_no, emekli_sicil_no, arsiv_no,
                    telefon, eposta, notlar, aktif, mevcut_id
                ))
                guncellenen += 1
            else:
                db.cursor.execute("""
                    INSERT INTO personel (
                        ad_soyad, tc_no, unvan, brans, gorev, grup, il_ilce, kurum_adi,
                        kurum_kodu, kurumda_baslama, kariyer, kademe_derece, ogrenim_durumu,
                        cinsiyet, kan_grubu, dogum_tarihi, goreve_baslama, kurum_sicil_no,
                        emekli_sicil_no, arsiv_no, telefon, eposta, notlar, aktif
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """, (
                    ad, tc, unvan, brans, gorev, grup, il_ilce, kurum_adi,
                    kurum_kodu, kurumda_baslama, kariyer, kademe_derece, ogrenim_durumu,
                    cinsiyet, kan_grubu, dogum_tarihi, goreve_baslama, kurum_sicil_no,
                    emekli_sicil_no, arsiv_no, telefon, eposta, notlar, aktif
                ))
                eklenen += 1
        db.conn.commit()
    except Exception as e:
        db.conn.rollback()
        raise

    mesaj = f"{kaynak_adi}: {eklenen} yeni personel eklendi"
    if guncellenen > 0:
        mesaj += f", {guncellenen} mevcut personel güncellendi."
    else:
        mesaj += "."
    return {"basarili": True, "mesaj": mesaj, "eklenen": eklenen, "guncellenen": guncellenen}


def _personel_excel_yukle_dogrudan(temp_yol, uzanti, yukleme_tipi=None):
    """Hazırlanmış personel dosyasını okur ve UPSERT mantığıyla kaydeder."""
    try:
        # 1. MEBBİS Görevlendirme Formatı Kontrolü
        if yukleme_tipi == 'mebbis_gorevlendirme' or ExcelMotoru.mebbis_gorevlendirme_mi(temp_yol):
            mebbis_liste = ExcelMotoru.mebbis_gorevlendirme_personel_oku(temp_yol)
            if not mebbis_liste:
                return {"basarili": False, "mesaj": "MEBBİS Görevlendirme personel tablosu boş veya okunamadı."}
            return _personelleri_veritabanina_kaydet(mebbis_liste, "MEBBİS Görevlendirme")

        # 2. MEBBİS Kadrolu Formatı Kontrolü (Düzensiz / Çok Satırlı / Kayan Sütunlar)
        if yukleme_tipi == 'mebbis_kadrolu' or ExcelMotoru.mebbis_mi_kontrol_et(temp_yol):
            mebbis_liste = ExcelMotoru.mebbis_personel_oku(temp_yol)
            if not mebbis_liste:
                return {"basarili": False, "mesaj": "MEBBİS personel listesi boş veya okunamadı."}
            return _personelleri_veritabanina_kaydet(mebbis_liste, "MEBBİS Kadrolu")

        # 3. e-Okul Formatı
        if yukleme_tipi == 'eokul':
            return {"basarili": False, "mesaj": "e-Okul Personel Şablonu bir sonraki güncellemede eklenecektir."}

        # 4. Standart Düz Tablo Formatı
        df_temp = ExcelMotoru.tabloyu_guvenli_oku(temp_yol, header=None)
        if df_temp.empty:
            return {"basarili": False, "mesaj": "Excel dosyası boş."}
        header_idx = 0
        for i, row in df_temp.iterrows():
            satir_metni = " ".join([str(x).upper() for x in row.values if pd.notna(x)])
            if "AD" in satir_metni and "SOYAD" in satir_metni:
                header_idx = i
                break

        df = ExcelMotoru.tabloyu_guvenli_oku(temp_yol, header=header_idx)
        df.columns = df.columns.str.strip().str.upper()

        ad_sutunu = _excel_sutunu_bul(df.columns, ['AD SOYAD', 'ADI SOYADI'])
        if not ad_sutunu:
            return {"basarili": False, "mesaj": "Excel'de 'Ad Soyad' başlığı bulunamadı!"}

        grup_sutunu = _excel_sutunu_bul(df.columns, ['GRUP / STATÜ', 'GRUP / STATU', 'GRUP/STATÜ', 'GRUP', 'PERSONEL TÜRÜ', 'PERSONEL TURU', 'TÜR', 'TUR', 'STATÜ', 'STATU'])
        unvan_sutunu = _excel_sutunu_bul(df.columns, ['UNVAN', 'UNVANI', 'ÜNVAN', 'ÜNVANI'])
        gorev_sutunu = _excel_sutunu_bul(df.columns, ['GÖREVİ', 'GÖREVI', 'GÖREV', 'GOREV', 'GOREVI'])
        brans_sutunu = _excel_sutunu_bul(df.columns, ['BRANŞI', 'BRANSI', 'BRANŞ', 'BRANS', 'ALANI', 'ALAN'])

        # Tam Yedek / Zengin Tablo Sütunları
        tc_sutunu = _excel_sutunu_bul(df.columns, ['T.C. KİMLİK NO', 'TC KİMLİK NO', 'T.C. KIMLIK NO', 'TC KIMLIK NO', 'T.C. NO', 'TC NO', 'T.C.', 'TC', 'TCKIMLIK', 'TCKİMLİK', 'KİMLİK NO', 'KIMLIK NO'])
        kariyer_sutunu = _excel_sutunu_bul(df.columns, ['KARİYER BASAMAĞI', 'KARIYER BASAMAGI', 'KARİYER', 'KARIYER', 'UNVAN / KARİYER', 'UNVAN/KARIYER'])
        kademe_sutunu = _excel_sutunu_bul(df.columns, ['DERECE / KADEME', 'KADEME / DERECE', 'DERECE/KADEME', 'KADEME/DERECE', 'KADEME DERECE', 'DERECE KADEME', 'KADEME', 'DERECE'])
        ogrenim_sutunu = _excel_sutunu_bul(df.columns, ['ÖĞRENİM DURUMU', 'OGRENIM DURUMU', 'ÖĞRENİM', 'OGRENIM', 'MEZUNİYET', 'EĞİTİM'])
        cinsiyet_sutunu = _excel_sutunu_bul(df.columns, ['CİNSİYET', 'CINSIYET', 'CİNSİYETİ', 'CINSIYETI'])
        kan_sutunu = _excel_sutunu_bul(df.columns, ['KAN GRUBU', 'KAN'])
        dogum_sutunu = _excel_sutunu_bul(df.columns, ['DOĞUM TARİHİ', 'DOGUM TARIHI', 'DOĞUM', 'DOGUM'])
        goreve_baslama_sutunu = _excel_sutunu_bul(df.columns, ['GÖREVE BAŞLAMA', 'GOREVE BASLAMA', 'GÖREVE BAŞLAMA TARİHİ', 'MEMURİYETE BAŞLAMA', 'İLK GÖREVE BAŞLAMA'])
        kurum_adi_sutunu = _excel_sutunu_bul(df.columns, ['KURUM ADI', 'KURUM', 'OKUL ADI', 'OKUL'])
        kurum_kodu_sutunu = _excel_sutunu_bul(df.columns, ['KURUM KODU', 'OKUL KODU'])
        il_ilce_sutunu = _excel_sutunu_bul(df.columns, ['İL / İLÇE', 'İL/İLÇE', 'IL / ILCE', 'IL/ILCE', 'İLÇE', 'ILCE', 'İL', 'IL'])
        kurumda_baslama_sutunu = _excel_sutunu_bul(df.columns, ['KURUMDA BAŞLAMA', 'KURUMDA BASLAMA', 'KURUMDA BAŞLAMA TARİHİ', 'KURUMA BAŞLAMA'])
        kurum_sicil_sutunu = _excel_sutunu_bul(df.columns, ['KURUM SİCİL NO', 'KURUM SICIL NO', 'KURUM SİCİL', 'KURUM SICIL', 'SİCİL NO', 'SICIL NO'])
        emekli_sicil_sutunu = _excel_sutunu_bul(df.columns, ['EMEKLİ / SGK SİCİL NO', 'EMEKLİ SİCİL NO', 'EMEKLİ/SGK SİCİL NO', 'EMEKLİ SİCİL', 'SGK SİCİL NO', 'SGK NO', 'EMEKLİ NO'])
        arsiv_sutunu = _excel_sutunu_bul(df.columns, ['ARŞİV NO', 'ARSIV NO', 'ARŞİV', 'ARSIV'])
        tel_sutunu = _excel_sutunu_bul(df.columns, ['TELEFON', 'TEL', 'GSM', 'CEP TELEFONU', 'CEP NO', 'İLETİŞİM'])
        eposta_sutunu = _excel_sutunu_bul(df.columns, ['E-POSTA', 'EPOSTA', 'EMAIL', 'E-MAIL', 'POSTA'])
        notlar_sutunu = _excel_sutunu_bul(df.columns, ['NOTLAR', 'NOT', 'AÇIKLAMA', 'ACIKLAMA'])
        aktif_sutunu = _excel_sutunu_bul(df.columns, ['AKTİF DURUM (1:AKTİF, 0:PASİF)', 'AKTİF DURUM', 'AKTIF DURUM', 'AKTİF', 'AKTIF', 'DURUM', 'AKTİF Mİ', 'STATÜ DURUMU'])

        def _hucre_temizle(val, varsayilan=""):
            if pd.isna(val) or val is None:
                return varsayilan
            s = str(val).strip()
            if s.lower() in ['nan', 'none', 'nat', '<na>']:
                return varsayilan
            if s.endswith('.0') and s[:-2].isdigit():
                s = s[:-2]
            return s

        def _tarih_metni_temizle(val):
            if pd.isna(val) or val is None:
                return ""
            if isinstance(val, (datetime, pd.Timestamp)):
                return val.strftime("%d/%m/%Y")
            s = str(val).strip()
            if s.lower() in ['nan', 'none', 'nat', '<na>']:
                return ""
            if " 00:00:00" in s:
                s = s.replace(" 00:00:00", "")
            return s

        yeni_personeller = []
        for _, row in df.iterrows():
            ad = _hucre_temizle(row[ad_sutunu])
            if not ad or ad.lower() in ['nan', 'none']:
                continue
            ad = " ".join(ad.split())

            unvan = _hucre_temizle(row[unvan_sutunu], "-") if unvan_sutunu else "-"
            gorev = _hucre_temizle(row[gorev_sutunu], "-") if gorev_sutunu else "-"
            brans = _hucre_temizle(row[brans_sutunu], "-") if brans_sutunu else "-"
            grup = _hucre_temizle(row[grup_sutunu]) if grup_sutunu else ""
            if not grup:
                grup = _personel_grup_tahmin_et(gorev if gorev != '-' else unvan)

            tc = _hucre_temizle(row[tc_sutunu]) if tc_sutunu else ""
            kariyer = _hucre_temizle(row[kariyer_sutunu], "-") if kariyer_sutunu else "-"
            kademe_raw = row[kademe_sutunu] if kademe_sutunu else "-"
            kademe_derece = ExcelMotoru.akilli_kademe_derece(kademe_raw)
            ogrenim_durumu = _hucre_temizle(row[ogrenim_sutunu], "-") if ogrenim_sutunu else "-"
            cinsiyet = _hucre_temizle(row[cinsiyet_sutunu]).upper() if cinsiyet_sutunu else ""
            kan_grubu = _hucre_temizle(row[kan_sutunu]) if kan_sutunu else ""
            dogum_tarihi = _tarih_metni_temizle(row[dogum_sutunu]) if dogum_sutunu else ""
            goreve_baslama = _tarih_metni_temizle(row[goreve_baslama_sutunu]) if goreve_baslama_sutunu else ""
            kurum_adi = _hucre_temizle(row[kurum_adi_sutunu]) if kurum_adi_sutunu else ""
            kurum_kodu = _hucre_temizle(row[kurum_kodu_sutunu]) if kurum_kodu_sutunu else ""
            il_ilce = _hucre_temizle(row[il_ilce_sutunu]) if il_ilce_sutunu else ""
            kurumda_baslama = _tarih_metni_temizle(row[kurumda_baslama_sutunu]) if kurumda_baslama_sutunu else ""
            kurum_sicil_no = _hucre_temizle(row[kurum_sicil_sutunu]) if kurum_sicil_sutunu else ""
            emekli_sicil_no = _hucre_temizle(row[emekli_sicil_sutunu]) if emekli_sicil_sutunu else ""
            arsiv_no = _hucre_temizle(row[arsiv_sutunu]) if arsiv_sutunu else ""
            telefon = _hucre_temizle(row[tel_sutunu]) if tel_sutunu else ""
            eposta = _hucre_temizle(row[eposta_sutunu]) if eposta_sutunu else ""
            notlar = _hucre_temizle(row[notlar_sutunu]) if notlar_sutunu else ""

            aktif = 1
            if aktif_sutunu:
                val_aktif = _hucre_temizle(row[aktif_sutunu]).lower()
                if any(k in val_aktif for k in ['0', 'pasif', 'hayır', 'hayir', 'false', 'başka kurumda', 'baska kurumda', 'görevlendirme']):
                    aktif = 0
                elif any(k in val_aktif for k in ['sağlık', 'saglik', 'raporlu', 'heyet', 'sıhhi']):
                    aktif = 2
                elif any(k in val_aktif for k in ['doğum', 'dogum', 'analık', 'analik']):
                    aktif = 3
                elif any(k in val_aktif for k in ['aylıksız', 'ayliksiz', 'ücretsiz', 'ucretsiz']):
                    aktif = 4
                elif any(k in val_aktif for k in ['refakat']):
                    aktif = 5
                elif any(k in val_aktif for k in ['asker', 'askerlik', 'terhis']):
                    aktif = 6
                elif any(k in val_aktif for k in ['yurtdışı', 'yurtdisi', 'eğitim', 'egitim', 'burslu']):
                    aktif = 7
                elif any(k in val_aktif for k in ['sendika', 'sendikal']):
                    aktif = 8
                elif val_aktif.isdigit() and int(val_aktif) in PERSONEL_DURUMLARI_REHBERI:
                    aktif = int(val_aktif)

            yeni_personeller.append({
                'ad_soyad': ad,
                'tc_no': tc,
                'unvan': unvan,
                'brans': brans,
                'gorev': gorev,
                'grup': grup,
                'il_ilce': il_ilce,
                'kurum_adi': kurum_adi,
                'kurum_kodu': kurum_kodu,
                'kurumda_baslama': kurumda_baslama,
                'kariyer': kariyer,
                'kademe_derece': kademe_derece,
                'ogrenim_durumu': ogrenim_durumu,
                'cinsiyet': cinsiyet,
                'kan_grubu': kan_grubu,
                'dogum_tarihi': dogum_tarihi,
                'goreve_baslama': goreve_baslama,
                'kurum_sicil_no': kurum_sicil_no,
                'emekli_sicil_no': emekli_sicil_no,
                'arsiv_no': arsiv_no,
                'telefon': telefon,
                'eposta': eposta,
                'notlar': notlar,
                'aktif': aktif
            })

        if not yeni_personeller:
            return {"basarili": False, "mesaj": "Excel'de aktarılacak geçerli personel bulunamadı."}

        res = _personelleri_veritabanina_kaydet(yeni_personeller, "Personel Tablosu")
        if res.get("basarili"):
            tam_yedek_mi = (tc_sutunu is not None and kariyer_sutunu is not None and kademe_sutunu is not None)
            if tam_yedek_mi:
                res["mesaj"] = f"{len(yeni_personeller)} personel veritabanı tam yedeğinden (24 alan) eksiksiz geri yüklendi."
            elif not tc_sutunu and not kariyer_sutunu:
                res["mesaj"] = f"{len(yeni_personeller)} personel yüklendi ve gruplandırıldı."
        return res

    except Exception as e:
        return {"basarili": False, "mesaj": str(e)}
    finally:
        if os.path.exists(temp_yol):
            try:
                os.remove(temp_yol)
            except Exception:
                pass


def personel_excel_isleme_gorevi(temp_yol, uzanti, job_id, yukleme_tipi=None):
    try:
        update_job(job_id, durum="isleniyor", mesaj="Personel Excel okunuyor...", yuzde=30)
        sonuc = _personel_excel_yukle_dogrudan(temp_yol, uzanti, yukleme_tipi=yukleme_tipi)
        if sonuc["basarili"]:
            update_job(job_id, durum="tamamlandi", mesaj=sonuc["mesaj"], yuzde=100)
        else:
            update_job(job_id, durum="hata", mesaj=sonuc["mesaj"], yuzde=100)
    except Exception as e:
        update_job(job_id, durum="hata", mesaj=str(e), yuzde=100)
    finally:
        if os.path.exists(temp_yol):
            try:
                os.remove(temp_yol)
            except Exception:
                pass


@router.post("/personel-excel-yukle")
async def personel_excel_yukle(
    background_tasks: BackgroundTasks, 
    dosya: UploadFile = File(...),
    yukleme_tipi: Optional[str] = Form(None)
):
    """Personel Excel aktarımını arka plana alır ve işlem kimliği döndürür."""
    uzanti, hata = excel_yukleme_dogrula(dosya)
    if hata:
        return {"basarili": False, "mesaj": hata}
    job_id = create_job(mesaj="Dosya alınıyor...", yuzde=0)
    temp_yol = f"temp_personel_{job_id}{uzanti}"
    with open(temp_yol, "wb") as buffer:
        shutil.copyfileobj(dosya.file, buffer)
    background_tasks.add_task(personel_excel_isleme_gorevi, temp_yol, uzanti, job_id, yukleme_tipi)
    return {"basarili": True, "job_id": job_id}


class PersonelPanoRequest(BaseModel):
    icerik: str
    yukleme_tipi: Optional[str] = "mebbis_gorevlendirme"


@router.post("/personel-pano-yukle")
def personel_pano_yukle(veri: PersonelPanoRequest):
    """MEBBİS veya diğer kaynaklardan kopyalanan tablo metnini (HTML/TSV) doğrudan çözümler."""
    try:
        icerik = (veri.icerik or "").strip()
        if not icerik:
            return {"basarili": False, "mesaj": "Yapıştırılan içerik boş."}

        yukleme_tipi = (veri.yukleme_tipi or "mebbis_gorevlendirme").lower()
        if "gorev" in yukleme_tipi:
            mebbis_liste = ExcelMotoru.mebbis_gorevlendirme_personel_oku(icerik)
            tip_adi = "MEBBİS Görevlendirme"
        elif "kadro" in yukleme_tipi:
            mebbis_liste = ExcelMotoru.mebbis_personel_oku(icerik)
            tip_adi = "MEBBİS Kadrolu"
        else:
            return {"basarili": False, "mesaj": "e-Okul Personel Şablonu bir sonraki güncellemede eklenecektir."}

        if not mebbis_liste:
            return {"basarili": False, "mesaj": f"{tip_adi} tablosu çözümlenemedi. Lütfen kopyalanan tablonun başlık ve verilerini içerdiğinden emin olun."}

        return _personelleri_veritabanina_kaydet(mebbis_liste, tip_adi)
    except Exception as e:
        import logging
        logging.error(f"personel_pano_yukle hatası: {e}", exc_info=True)
        return {"basarili": False, "mesaj": f"Pano verisi çözümlenirken hata oluştu: {str(e)}"}


@router.get("/personel-sablon-indir/gorevlendirme")
def personel_sablon_indir_gorevlendirme():
    """MEBBİS Görevlendirme Personel için önceden metne kilitlenmiş boş Excel şablonu üretir."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from fastapi.responses import Response
    import io

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Görevlendirme Personel"

    headers = [
        'TC Kimlik No', 'Adı', 'Soyadı', 'İstihdam Tipi', 'Kariyer Basamağı',
        'Ay.Der/Kad.', 'Emekli Sicil No', 'Kurum Sicil No', 'Kadro Ünvanı',
        'Görevi', 'Alanı', 'Görev Yeri', 'İli', 'İlçesi', 'Durumu', 'Cinsiyeti'
    ]
    ws.append(headers)

    header_fill = PatternFill(start_color="0284C7", end_color="0284C7", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")

    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        col_letter = openpyxl.utils.get_column_letter(col_num)
        ws.column_dimensions[col_letter].width = 18
        ws.column_dimensions[col_letter].number_format = '@'

    # Örnek satır ekle
    sample_row = [
        '12345678901', 'Ahmet', 'Yılmaz', 'Kadrolu', 'Uzman Öğretmen',
        '1/4', '12345678', '1234', 'Öğretmen', 'Öğretmen', 'Matematik',
        'Örnek Anadolu Lisesi', 'ESKİŞEHİR', 'ODUNPAZARI', 'Görevde', 'Erkek'
    ]
    ws.append(sample_row)

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)

    return Response(
        content=out.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=MEBBIS_Gorevlendirme_Personel_Sablonu.xlsx"}
    )


@router.post("/personel-ekle")
def personel_ekle(veri: PersonelRequest):
    ad = str(veri.ad or "").strip()
    if not ad:
        return {"basarili": False, "mesaj": "Ad Soyad boş bırakılamaz!"}
    unvan = str(veri.unvan).strip() if veri.unvan else "-"
    gorev = str(veri.gorev or "-").strip()
    brans = str(veri.brans or "-").strip()
    grup = veri.grup or _personel_grup_tahmin_et(gorev)

    db.cursor.execute("PRAGMA table_info(personel)")
    mevcut_sutunlar = [col[1] for col in db.cursor.fetchall()]
    if 'grup' not in mevcut_sutunlar:
        db.cursor.execute("ALTER TABLE personel ADD COLUMN grup TEXT DEFAULT 'Diğer Personel'")
    if 'unvan' not in mevcut_sutunlar:
        db.cursor.execute("ALTER TABLE personel ADD COLUMN unvan TEXT DEFAULT '-'")

    db.cursor.execute("INSERT INTO personel (ad_soyad, unvan, brans, gorev, grup) VALUES (?, ?, ?, ?, ?)", (ad, unvan, brans, gorev, grup))
    db.conn.commit()
    return {"basarili": True, "mesaj": f"{ad} eklendi."}



@router.delete("/personel-sil/{ad}")
def personel_sil(ad: str):
    db.cursor.execute("DELETE FROM personel WHERE ad_soyad = ?", (ad,))
    db.conn.commit()
    
    # Eşleşmeleri sil (ad silindiği için eski eşleşmeleri de temizlemek iyi olabilir, ama şimdilik bırakıyoruz)
    return {"basarili": True, "mesaj": f"{ad} silindi."}

@router.put("/personel-guncelle/{eski_ad}")
def personel_guncelle(eski_ad: str, veri: PersonelRequest):
    yeni_ad = str(veri.ad or "").strip()
    if not yeni_ad:
        return {"basarili": False, "mesaj": "Ad Soyad boş bırakılamaz!"}
    unvan = str(veri.unvan).strip() if veri.unvan else "-"
    gorev = str(veri.gorev or "-").strip()
    brans = str(veri.brans or "-").strip()
    grup = veri.grup or _personel_grup_tahmin_et(gorev)

    db.cursor.execute("UPDATE personel SET ad_soyad=?, unvan=?, brans=?, gorev=?, grup=? WHERE ad_soyad=?", (yeni_ad, unvan, brans, gorev, grup, eski_ad))
    db.conn.commit()
    
    # Oto-eşleşmeleri güncelle
    ayar = ayarlari_al()
    eslesmeler = ayar.get("oto_eslesmeler", [])
    degisiklik_var = False
    for eslesme in eslesmeler:
        if eslesme.get("hedef") == eski_ad:
            eslesme["hedef"] = yeni_ad
            degisiklik_var = True
    
    if degisiklik_var:
        ayar["oto_eslesmeler"] = eslesmeler
        SistemMotoru.ayarlari_kaydet(yollar["AYARLAR"], ayar)
        
    return {"basarili": True, "mesaj": f"Personel güncellendi."}

@router.post("/personel-profil-guncelle")
def personel_profil_guncelle(veri: PersonelGuncelleRequest):
    try:
        ad = str(veri.ad_soyad or "").strip()
        if not ad:
            return {"basarili": False, "mesaj": "Ad Soyad boş bırakılamaz!"}
            
        tc = str(veri.tc_no or "").strip()
        unvan = str(veri.unvan or "-").strip()
        brans = str(veri.brans or "-").strip()
        gorev = str(veri.gorev or "-").strip()
        grup = str(veri.grup or "Kadrolu").strip()
        kariyer = str(veri.kariyer or "-").strip()

        # Öğretmen değilse branş ve kariyer alanları gerekmez, '-' atanır
        if not ogretmen_mi(grup=grup, unvan=unvan, gorev=gorev):
            if not brans or brans == "":
                brans = "-"
            if not kariyer or kariyer == "":
                kariyer = "-"

        kademe_derece = str(veri.kademe_derece or "-").strip()
        # İUP, TYP, Sürekli İşçi ve Hizmet Alımı için derece-kademe geçerli değildir
        if grup in ["İUP", "TYP", "Sürekli İşçi", "Hizmet Alımı"]:
            kademe_derece = "-"
        ogrenim_durumu = str(veri.ogrenim_durumu or "-").strip()
        cinsiyet = str(veri.cinsiyet or "").strip().upper()
        kan_grubu = str(veri.kan_grubu or "").strip()
        dogum_tarihi = str(veri.dogum_tarihi or "").strip()
        goreve_baslama = str(veri.goreve_baslama or "").strip()
        kurum_sicil_no = str(veri.kurum_sicil_no or "").strip()
        emekli_sicil_no = str(veri.emekli_sicil_no or "").strip()
        arsiv_no = str(veri.arsiv_no or "").strip()
        telefon = str(veri.telefon or "").strip()
        eposta = str(veri.eposta or "").strip()
        notlar = str(veri.notlar or "").strip()
        il_ilce = str(veri.il_ilce or "").strip()
        kurum_adi = str(veri.kurum_adi or "").strip()
        kurum_kodu = str(veri.kurum_kodu or "").strip()
        kurumda_baslama = str(veri.kurumda_baslama or "").strip()
        try:
            aktif_int = int(veri.aktif) if veri.aktif is not None else 1
            aktif = aktif_int if aktif_int in PERSONEL_DURUMLARI_REHBERI else 1
        except (ValueError, TypeError):
            aktif = 1

        if veri.id:
            db.cursor.execute("""
                UPDATE personel SET
                    ad_soyad=?, tc_no=?, unvan=?, brans=?, gorev=?, grup=?, il_ilce=?, kurum_adi=?, 
                    kurum_kodu=?, kurumda_baslama=?, kariyer=?, kademe_derece=?, ogrenim_durumu=?, 
                    cinsiyet=?, kan_grubu=?, dogum_tarihi=?, goreve_baslama=?, kurum_sicil_no=?, 
                    emekli_sicil_no=?, arsiv_no=?, telefon=?, eposta=?, notlar=?, aktif=?
                WHERE rowid=?
            """, (
                ad, tc, unvan, brans, gorev, grup, il_ilce, kurum_adi,
                kurum_kodu, kurumda_baslama, kariyer, kademe_derece, ogrenim_durumu,
                cinsiyet, kan_grubu, dogum_tarihi, goreve_baslama, kurum_sicil_no,
                emekli_sicil_no, arsiv_no, telefon, eposta, notlar, aktif, veri.id
            ))
        else:
            db.cursor.execute("SELECT rowid FROM personel WHERE ad_soyad=?", (ad,))
            mevcut = db.cursor.fetchone()
            if mevcut:
                db.cursor.execute("""
                    UPDATE personel SET
                        tc_no=?, unvan=?, brans=?, gorev=?, grup=?, il_ilce=?, kurum_adi=?, 
                        kurum_kodu=?, kurumda_baslama=?, kariyer=?, kademe_derece=?, ogrenim_durumu=?, 
                        cinsiyet=?, kan_grubu=?, dogum_tarihi=?, goreve_baslama=?, kurum_sicil_no=?, 
                        emekli_sicil_no=?, arsiv_no=?, telefon=?, eposta=?, notlar=?, aktif=?
                    WHERE ad_soyad=?
                """, (
                    tc, unvan, brans, gorev, grup, il_ilce, kurum_adi,
                    kurum_kodu, kurumda_baslama, kariyer, kademe_derece, ogrenim_durumu,
                    cinsiyet, kan_grubu, dogum_tarihi, goreve_baslama, kurum_sicil_no,
                    emekli_sicil_no, arsiv_no, telefon, eposta, notlar, aktif, ad
                ))
            else:
                db.cursor.execute("""
                    INSERT INTO personel (
                        ad_soyad, tc_no, unvan, brans, gorev, grup, il_ilce, kurum_adi, 
                        kurum_kodu, kurumda_baslama, kariyer, kademe_derece, ogrenim_durumu, 
                        cinsiyet, kan_grubu, dogum_tarihi, goreve_baslama, kurum_sicil_no, 
                        emekli_sicil_no, arsiv_no, telefon, eposta, notlar, aktif
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """, (
                    ad, tc, unvan, brans, gorev, grup, il_ilce, kurum_adi,
                    kurum_kodu, kurumda_baslama, kariyer, kademe_derece, ogrenim_durumu,
                    cinsiyet, kan_grubu, dogum_tarihi, goreve_baslama, kurum_sicil_no,
                    emekli_sicil_no, arsiv_no, telefon, eposta, notlar, aktif
                ))

        db.conn.commit()
        return {"basarili": True, "mesaj": f"{ad} profil bilgileri kaydedildi."}
    except Exception as e:
        return {"basarili": False, "mesaj": str(e)}

@router.get("/mevzuat-bilgisi/{grup}")
def mevzuat_bilgisi_endpoint(grup: str, unvan: Optional[str] = None):
    try:
        return {"basarili": True, "veri": mevzuat_bilgisi_al(grup=grup, unvan=unvan)}
    except Exception as e:
        return {"basarili": False, "mesaj": str(e)}

@router.get("/mevzuat-ozeti")
def tum_mevzuat_ozeti_endpoint():
    try:
        return {"basarili": True, "veri": tum_mevzuat_ozeti()}
    except Exception as e:
        return {"basarili": False, "mesaj": str(e)}

@router.get("/personel-excel-indir")
def personel_excel_indir(format: str = "ozet"):
    ana_klasor, ayar = pdf_klasoru_hazirla()

    if format == "tam":
        sutun_haritasi = [
            ("ad_soyad", "Ad Soyad"),
            ("tc_no", "T.C. Kimlik No"),
            ("unvan", "Unvan"),
            ("brans", "Branş"),
            ("gorev", "Görev"),
            ("grup", "Grup / Statü"),
            ("kariyer", "Kariyer Basamağı"),
            ("kademe_derece", "Derece / Kademe"),
            ("ogrenim_durumu", "Öğrenim Durumu"),
            ("cinsiyet", "Cinsiyet"),
            ("kan_grubu", "Kan Grubu"),
            ("dogum_tarihi", "Doğum Tarihi"),
            ("goreve_baslama", "Göreve Başlama"),
            ("kurum_adi", "Kurum Adı"),
            ("kurum_kodu", "Kurum Kodu"),
            ("il_ilce", "İl / İlçe"),
            ("kurumda_baslama", "Kurumda Başlama"),
            ("kurum_sicil_no", "Kurum Sicil No"),
            ("emekli_sicil_no", "Emekli / SGK Sicil No"),
            ("arsiv_no", "Arşiv No"),
            ("telefon", "Telefon"),
            ("eposta", "E-Posta"),
            ("notlar", "Notlar"),
            ("aktif", "Aktif Durum (1:Aktif, 0:Pasif)")
        ]

        db_alanlari = [col for col, _ in sutun_haritasi]
        excel_basliklari = [hdr for _, hdr in sutun_haritasi]

        db.cursor.execute(f"SELECT {', '.join(db_alanlari)} FROM personel ORDER BY ad_soyad")
        satirlar = db.cursor.fetchall()

        if not satirlar:
            return {"basarili": False, "mesaj": "Personel listesi boş!"}

        df = pd.DataFrame(satirlar, columns=excel_basliklari)
        yol = os.path.join(ana_klasor, "Personel_Veritabani_Tam_Yedek.xlsx")

        try:
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
            with pd.ExcelWriter(yol, engine='openpyxl') as writer:
                df.to_excel(writer, index=False, sheet_name='Personel_Yedek')
                ws = writer.sheets['Personel_Yedek']

                header_fill = PatternFill(start_color="0284C7", end_color="0284C7", fill_type="solid")
                header_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
                thin_border = Border(
                    left=Side(style='thin', color='CBD5E1'),
                    right=Side(style='thin', color='CBD5E1'),
                    top=Side(style='thin', color='CBD5E1'),
                    bottom=Side(style='thin', color='CBD5E1')
                )

                for col_num in range(1, len(excel_basliklari) + 1):
                    cell = ws.cell(row=1, column=col_num)
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                    cell.border = thin_border

                ws.freeze_panes = 'A2'
                ws.row_dimensions[1].height = 28

                for col in ws.columns:
                    col_letter = col[0].column_letter
                    max_len = max(len(str(cell.value or '')) for cell in col)
                    ws.column_dimensions[col_letter].width = max(max_len + 4, 12)
        except Exception as e:
            logging.error(f"Excel openpyxl biçimlendirme hatası: {e}")
            df.to_excel(yol, index=False)

        try:
            os.startfile(yol)
        except Exception:
            pass

        return {"basarili": True, "mesaj": "Veritabanı tam yedeği (24 sütun) Excel dosyası olarak oluşturuldu."}

    else:
        db.cursor.execute("SELECT ad_soyad, brans, gorev, grup FROM personel ORDER BY ad_soyad")
        satirlar = db.cursor.fetchall()

        if not satirlar:
            return {"basarili": False, "mesaj": "Personel listesi boş!"}

        personel_listesi = [{'Ad Soyad': r[0], 'Branş': r[1], 'Görev': r[2], 'Grup': r[3]} for r in satirlar]
        df = pd.DataFrame(personel_listesi)
        yol = os.path.join(ana_klasor, "Personel_Listesi.xlsx")
        df.to_excel(yol, index=False)

        try:
            os.startfile(yol)
        except Exception:
            pass

        return {"basarili": True, "mesaj": "Özet personel listesi Excel dosyası olarak oluşturuldu."}


@router.get("/personel-pdf-indir")
def personel_pdf_indir():
    db.cursor.execute("SELECT ad_soyad, brans, gorev, grup FROM personel ORDER BY ad_soyad")
    veri = [[r[0], r[1], r[2], r[3]] for r in db.cursor.fetchall()]
    
    if not veri:
        return {"basarili": False, "mesaj": "Personel listesi boş!"}
        
    ana_klasor, ayar = pdf_klasoru_hazirla()
    yol = os.path.join(ana_klasor, "Personel_Listesi.pdf")
    motor = PDFYoneticisi(ayar)
    
    # Yeni eklenecek PDF methodu: personel_raporu_ciz
    motor.personel_raporu_ciz(veri, yol)
    
    os.startfile(yol)
    return {"basarili": True, "mesaj": "PDF oluşturuldu."}



@router.delete("/personel-sifirla")
def personel_sifirla():
    try:
        basarili, hata = SistemMotoru.yedek_al(db.db_yolu, ayarlari_al(), yollar["YEDEK"])
        if not basarili:
            return {"basarili": False, "mesaj": f"Sıfırlama öncesi yedek alınamadı: {hata}"}
        db.cursor.execute("DELETE FROM personel")
        db.conn.commit()
        islem_logla("uyarı", "Personel sıfırlama", "Tüm personel kayıtları sıfırlandı.")
        return {"basarili": True, "mesaj": "Tüm personel silindi."}
    except Exception as e:
        return {"basarili": False, "mesaj": str(e)}

