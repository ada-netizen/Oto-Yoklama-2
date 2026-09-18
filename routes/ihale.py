from fastapi import APIRouter, UploadFile, File, Form, HTTPException, BackgroundTasks, Body, Request
from fastapi.responses import JSONResponse, FileResponse
from pydantic import BaseModel
from typing import List, Optional
import os, json, re, tempfile, shutil, uuid, platform
import pandas as pd
from datetime import datetime, timedelta
from veritabani import VeritabaniYoneticisi
from sistem_motoru import SistemMotoru
from excel_motoru import ExcelMotoru
from pdf_motoru import PDFYoneticisi
import logging
from araclar import VeriAraclari
from dependencies import db, yollar, ayarlar, islem_logla, create_job, update_job, get_job, GLOBAL_ISLEMLER, islem_durumlari
from utils import *
from utils import _excel_sutunu_bul, _sablon_excel_yolu_olustur
from typing import Dict

router = APIRouter()

class SablonVerisi(BaseModel):
    kalemler: List[Dict]
    firmaVergiler: List[str] = []
    firmalar: List[str] = []
    firmaAdresleri: List[str] = []

VARSAYILAN_OLCU_BIRIMLERI = [
    "ADET_BIRIM",
    "AFIF_BIRIM_FIYATI",
    "ATV_BIRIM_FIYATI",
    "ALTIN_AYARI",
    "KG_METREKARE",
    "TON_BASINA_TASIMA_KAPASITESI",
    "ADET_CIFT",
    "BRUT_KALORI_DEGERI",
    "BIN_LITRE",
    "GUMUS",
    "GRAM",
    "GROS_TON",
    "YUZ_ADET",
    "KILOGRAM_ADET",
    "KILOWATT_SAAT",
    "KILOWATT",
    "LITRE",
    "METRE",
    "METREKUP",
    "METREKARE",
    "TON",
    "KALEM",
    "PUAN",
]

@router.get("/olcu-birimleri")
def olcu_birimleri_getir():
    return {"birimler": VARSAYILAN_OLCU_BIRIMLERI}

@router.post("/sablon-hazirla")
def sablon_hazirla(veri: SablonVerisi):
    try:
        logging.info(f"SABLON HAZIRLA: {len(veri.kalemler)} kalem, {len(veri.firmalar)} firma")
        import openpyxl, shutil as _shutil

        # Şablon dosyasını bul
        sablon_yolu = _sablon_excel_yolu_olustur()
        if not sablon_yolu:
            return {"basarili": False, "mesaj": "Şablon dosyası (sablonlar/sablon.xlsx) bulunamadı."}

        # Hedef dosyayı şablondan kopyala (açık olsa bile farklı isim)
        ayar = ayarlari_al()
        ana_klasor = ayar.get("pdf_kayit_klasoru", yollar["PDF"])
        os.makedirs(ana_klasor, exist_ok=True)
        zaman_damgasi = datetime.now().strftime("%Y%m%d_%H%M%S")
        dosya_yolu = os.path.join(ana_klasor, f"Yaklasik_Maliyet_Sablon_{zaman_damgasi}.xlsx")
        _shutil.copy2(sablon_yolu, dosya_yolu)

        # Kopyayı aç ve doldur
        wb = openpyxl.load_workbook(dosya_yolu)
        ws = wb.active

        # Başlık satırını koru, 2. satırdan itibaren temizle
        if ws.max_row > 1:
            ws.delete_rows(2, ws.max_row - 1)

        # Firma ve fiyat verilerini hazırla
        firmalar     = veri.firmalar     or []
        vergiler     = veri.firmaVergiler or []
        gecerli_kalemler = [k for k in veri.kalemler if str(k.get('cins', '')).strip()]

        row = 2
        for kalem in gecerli_kalemler:
            cins   = kalem.get('cins', '')
            miktar = kalem.get('miktar', '')
            birim  = kalem.get('birim', '')
            fiyatlar = kalem.get('fiyatlar', [])

            # Kalem ana satırı: A=Ürün Adı, B=Miktar, C=Ölçü Birimi
            ws.cell(row=row, column=1, value=cins)
            try:
                ws.cell(row=row, column=2, value=float(miktar) if miktar else None)
            except Exception:
                ws.cell(row=row, column=2, value=miktar)
            ws.cell(row=row, column=3, value=birim)
            row += 1

            # Her firma için alt satır: D=ÜrünNo(boş), E=Model(boş), F=Marka(boş), G=Birim Tutar, H=Firma VKN/TCKN
            for fi, firma_adi in enumerate(firmalar):
                if not str(firma_adi).strip():
                    continue
                vergi = vergiler[fi] if fi < len(vergiler) else ""
                bf = ""
                if fi < len(fiyatlar) and fiyatlar[fi] not in ("", None):
                    try:
                        bf = float(fiyatlar[fi])
                    except Exception:
                        bf = fiyatlar[fi]
                ws.cell(row=row, column=4, value=None)   # Ürün No
                ws.cell(row=row, column=5, value=None)   # Model
                ws.cell(row=row, column=6, value=None)   # Marka
                ws.cell(row=row, column=7, value=bf)     # Birim Tutar
                ws.cell(row=row, column=8, value=vergi)  # Firma VKN/TCKN
                row += 1

        wb.save(dosya_yolu)
        dosyayi_otomatik_ac(dosya_yolu)
        return {"basarili": True, "mesaj": f"Şablon hazırlandı ve açıldı."}
    except Exception as e:
        logging.error(f"Şablon hazırlanırken hata: {e}", exc_info=True)
        return {"basarili": False, "mesaj": str(e)}

@router.post("/excel-onizle")
async def excel_onizle(dosya: UploadFile = File(...), tur: str = Form("personel")):
    """İçe aktarmadan önce dosyanın ilk satırlarını ve temel hatalarını döndürür."""
    uzanti, hata = excel_yukleme_dogrula(dosya)
    if hata:
        return {"basarili": False, "mesaj": hata}
    if tur not in {"ogrenci", "devamsizlik", "personel"}:
        return {"basarili": False, "mesaj": "Geçersiz içe aktarma türü."}

    temp_yol = os.path.join(tempfile.gettempdir(), f"onizleme_{uuid.uuid4().hex}{uzanti}")
    try:
        with open(temp_yol, "wb") as buffer:
            shutil.copyfileobj(dosya.file, buffer)

        if uzanti == ".csv":
            try:
                df_ham = pd.read_csv(temp_yol, header=None, encoding="utf-8-sig")
            except UnicodeDecodeError:
                df_ham = pd.read_csv(temp_yol, header=None, encoding="cp1254")
        else:
            try:
                df_ham = pd.read_excel(temp_yol, header=None)
            except Exception:
                try:
                    df_list = pd.read_html(temp_yol, header=None)
                    df_ham = max(df_list, key=len) if df_list else pd.DataFrame()
                except Exception:
                    df_ham = pd.read_csv(temp_yol, header=None, on_bad_lines='skip')
        if df_ham.empty:
            return {"basarili": False, "mesaj": "Dosya boş."}

        header_idx = 0
        if tur == "personel":
            for i, row in df_ham.iterrows():
                basliklar = {_excel_sutunu_bul([deger], ["AD SOYAD", "ADI SOYADI"]) for deger in row.values}
                if any(basliklar):
                    header_idx = i
                    break

        if uzanti == ".csv":
            try:
                df = pd.read_csv(temp_yol, header=header_idx, encoding="utf-8-sig")
            except UnicodeDecodeError:
                df = pd.read_csv(temp_yol, header=header_idx, encoding="cp1254")
        else:
            try:
                df = pd.read_excel(temp_yol, header=header_idx)
            except Exception:
                try:
                    df_list = pd.read_html(temp_yol, header=header_idx)
                    df = max(df_list, key=len) if df_list else pd.DataFrame()
                except Exception:
                    df = pd.read_csv(temp_yol, header=header_idx, on_bad_lines='skip')
        df.columns = [str(sutun).strip() for sutun in df.columns]

        hatalar = []
        if tur == "personel":
            ad_sutunu = _excel_sutunu_bul(df.columns, ["AD SOYAD", "ADI SOYADI"])
            if not ad_sutunu:
                hatalar.append("Ad Soyad sütunu bulunamadı.")
            else:
                eksik_ad = int(df[ad_sutunu].isna().sum())
                if eksik_ad:
                    hatalar.append(f"{eksik_ad} satırda ad-soyad eksik.")

        onizleme_kayitlar = df.head(20).to_dict(orient="records")
        onizleme = [{str(k): (str(v) if pd.notna(v) else "") for k, v in row.items()} for row in onizleme_kayitlar]
        return {
            "basarili": True,
            "tur": tur,
            "toplam_satir": int(len(df)),
            "gosterilen_satir": len(onizleme),
            "sutunlar": list(df.columns),
            "onizleme": onizleme,
            "hatalar": hatalar,
        }
    except Exception as e:
        logging.error(f"Excel önizleme hatası: {e}")
        return {"basarili": False, "mesaj": f"Dosya önizlenemedi: {e}"}
    finally:
        if os.path.exists(temp_yol):
            try:
                os.remove(temp_yol)
            except Exception:
                pass

@router.post("/ihale-excel-oku")
async def ihale_excel_oku(dosya: UploadFile = File(...)):
    try:
        temp_path = os.path.join(yollar["TEMP"], f"ihale_{uuid.uuid4().hex}.xls")
        with open(temp_path, "wb") as f:
            shutil.copyfileobj(dosya.file, f)
        
        kalemler, hata = ExcelMotoru.ihale_oku(temp_path)
        if hata:
            return {"basarili": False, "mesaj": hata}
            
        return {"basarili": True, "kalemler": kalemler}
    except Exception as e:
        return {"basarili": False, "mesaj": str(e)}

@router.post("/ihale-tekli-belge")
async def ihale_tekli_belge(request: Request, background_tasks: BackgroundTasks = None):
    veri = await request.json()
    ayar = ayarlari_al()
    pdf_yol = ayar.get("pdf_kayit_klasoru", yollar["PDF"])
    
    mudur_adi = "Okul Müdürü"
    try:
        db.cursor.execute("SELECT ad_soyad FROM personel WHERE gorev LIKE '%MÜDÜR%' AND gorev NOT LIKE '%MÜDÜR YARDIMCISI%' LIMIT 1")
        row = db.cursor.fetchone()
        if row:
            mudur_adi = row[0]
    except: pass

    veri["okul_adi"] = ayar.get("okul_adi", "Okul")
    baslik = str(veri.get("resmi_baslik") or "").strip()
    if not baslik:
        veri["resmi_baslik"] = f"{veri['okul_adi'].strip()} Müdürlüğü"
    else:
        if not re.search(r"(MÜDÜRLÜĞÜ|MÜDÜRLÜK|OKULU|LİSESİ|ORTAOKULU|İLKOKULU|ANAOKULU|KAYMAKAMLIĞI|VALİLİĞİ|BAŞKANLIĞI)$", baslik, flags=re.IGNORECASE):
            veri["resmi_baslik"] = f"{baslik} Müdürlüğü"
    veri["okul_muduru"] = mudur_adi
    
    # Komisyon üyelerinin görevlerini (unvanlarını) veritabanından çekip ekleyelim
    if "komisyon" in veri and isinstance(veri["komisyon"], dict):
        for k_key, k_name in list(veri["komisyon"].items()):
            if k_name and isinstance(k_name, str):
                try:
                    db.cursor.execute("SELECT gorev FROM personel WHERE ad_soyad = ?", (k_name,))
                    row = db.cursor.fetchone()
                    if row:
                        veri["komisyon"][f"{k_key}_gorev"] = row[0]
                    else:
                        veri["komisyon"][f"{k_key}_gorev"] = "Üye"
                except:
                    veri["komisyon"][f"{k_key}_gorev"] = "Üye"
    
    try:
        import ihale_motoru
        sonuc_dosyasi = ihale_motoru.belge_uret(veri, pdf_yol)
        if platform.system() == "Windows" and sonuc_dosyasi and os.path.exists(sonuc_dosyasi):
            os.startfile(sonuc_dosyasi)
        return {"basarili": True, "mesaj": "Belge başarıyla üretildi!", "dosya": sonuc_dosyasi}
    except Exception as e:
        import traceback
        logging.error(f"Tekli belge hatasi: {e}\n{traceback.format_exc()}")
        return {"basarili": False, "mesaj": f"Hata: {e}"}

