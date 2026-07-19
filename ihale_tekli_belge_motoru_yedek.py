import os
import datetime
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side

font_name = 'Helvetica-TR'
try:
    pdfmetrics.registerFont(TTFont('Helvetica-TR', 'C:\\Windows\\Fonts\\arial.ttf'))
except:
    font_name = 'Helvetica'

def format_date(date_str):
    if not date_str:
        return datetime.datetime.now().strftime("%d.%m.%Y")
    try:
        dt = datetime.datetime.strptime(date_str, "%Y-%m-%d")
        return dt.strftime("%d.%m.%Y")
    except:
        return date_str

def get_komisyon(veri):
    kom = veri.get("komisyon", {})
    uyeler = [
        kom.get("ihale_kom_yaklasik_1", ""),
        kom.get("ihale_kom_yaklasik_2", ""),
        kom.get("ihale_kom_yaklasik_3", "")
    ]
    uyeler = [u for u in uyeler if u]
    return uyeler

def belge_uret(veri, pdf_yol):
    belge_tipi = veri.get("belge_tipi", "")
    fmt = veri.get("format", "pdf")
    
    if not os.path.exists(pdf_yol):
        os.makedirs(pdf_yol, exist_ok=True)
        
    klasor_adi = f"Ihale_{datetime.datetime.now().strftime('%Y%m%d_%H%M%S')}"
    hedef_klasor = os.path.join(pdf_yol, klasor_adi)
    os.makedirs(hedef_klasor, exist_ok=True)
    
    if belge_tipi == "fiyat_isteme":
        if fmt == "pdf":
            return uret_fiyat_isteme_pdf(veri, hedef_klasor)
        elif fmt == "excel":
            return uret_fiyat_isteme_excel(veri, hedef_klasor)
    return None

def uret_fiyat_isteme_pdf(veri, hedef_klasor):
    dosya_adi = os.path.join(hedef_klasor, "Yaklasik_Maliyet_Fiyat_Isteme.pdf")
    doc = SimpleDocTemplate(
        dosya_adi,
        pagesize=A4,
        rightMargin=25*mm,
        leftMargin=25*mm,
        topMargin=25*mm,
        bottomMargin=25*mm
    )
    
    # Tüm PDF 9 punto
    style_normal = ParagraphStyle('Normal_TR', fontName=font_name, fontSize=9, leading=11)
    style_justify = ParagraphStyle('Justify_TR', fontName=font_name, fontSize=9, alignment=4, leading=11, firstLineIndent=12*mm)
    style_center = ParagraphStyle('Center_TR', fontName=font_name, fontSize=9, alignment=1, leading=11)
    style_right = ParagraphStyle('Right_TR', fontName=font_name, fontSize=9, alignment=2, leading=11)
    
    elements = []
    
    # 1. Resmi Yazı Başlığı
    baslik_satirlari = veri.get("resmi_baslik", "").replace("\\n", "\n").split("\n")
    for satir in baslik_satirlari:
        elements.append(Paragraph(satir.strip(), style_center))
    
    elements.append(Spacer(1, 8*mm))
    
    # 2. Sayı, Tarih, Konu
    tarih = format_date(veri.get("belge_tarihi", ""))
    sayi = veri.get('yazisma_kodu', '')
    
    tbl_sayi_konu = Table([
        [Paragraph("Sayı", style_normal), Paragraph(":", style_normal), Paragraph(sayi, style_normal), Paragraph(tarih, style_right)],
        [Paragraph("Konu", style_normal), Paragraph(":", style_normal), Paragraph("Yaklaşık Maliyet Fiyatları", style_normal), ""]
    ], colWidths=[12*mm, 3*mm, 95*mm, 50*mm])
    
    tbl_sayi_konu.setStyle(TableStyle([
        ('ALIGN', (0,0), (0,-1), 'LEFT'),
        ('ALIGN', (1,0), (1,-1), 'LEFT'),
        ('ALIGN', (2,0), (2,-1), 'LEFT'),
        ('ALIGN', (3,0), (3,0), 'RIGHT'),
        ('VALIGN', (0,0), (-1,-1), 'TOP'),
        ('BOTTOMPADDING', (0,0), (-1,-1), 1),
        ('TOPPADDING', (0,0), (-1,-1), 1),
        ('LEFTPADDING', (0,0), (-1,-1), 0),
        ('RIGHTPADDING', (0,0), (-1,-1), 0),
    ]))
    elements.append(tbl_sayi_konu)
    
    elements.append(Spacer(1, 8*mm))
    
    # 3. Hitap ve Gövde
    elements.append(Paragraph("Sayın Yetkili", style_center))
    elements.append(Spacer(1, 3*mm))
    metin = "İdaremizce satın alınması düşünülen aşağıda cinsi, miktarı, özellikleri ve diğer şartları yazılı mal, hizmet ya da yapım işlerinin 4734 Sayılı Kamu İhale Kanunu gereğince, yaklaşık maliyetinin tesbit edilmesinde değerlendirilmek ve KDV hariç olmak üzere piyasada satış fiyatlarının bildirilmesini rica ederim. / ederiz."
    elements.append(Paragraph(metin, style_justify))
    elements.append(Spacer(1, 8*mm))
    
    # 4. Komisyon İmzaları (Boşluk kaldırıldı, parantezler kaldırıldı)
    komisyon_uyeleri = get_komisyon(veri)
    if len(komisyon_uyeleri) > 0:
        imza_data = [[]]
        for idx, uye in enumerate(komisyon_uyeleri):
            title = "Başkan" if idx == 0 else "Üye"
            imza_data[0].append(Paragraph(f"{uye}<br/>{title}", style_center))
        
        col_w = 160*mm / len(komisyon_uyeleri)
        sig_table = Table(imza_data, colWidths=[col_w] * len(komisyon_uyeleri))
        sig_table.setStyle(TableStyle([
            ('ALIGN', (0,0), (-1,-1), 'CENTER'),
            ('VALIGN', (0,0), (-1,-1), 'TOP'),
        ]))
        elements.append(sig_table)
    
    elements.append(Spacer(1, 8*mm))
    
    # 5. Tablo
    kalemler = veri.get("kalemler", [])
    gecerli_kalemler = [k for k in kalemler if k.get("cins", "").strip()]
    if not gecerli_kalemler:
        gecerli_kalemler = [{}] 
        
    table_data = []
    
    # Super Header
    table_data.append(["SATIN ALINACAK MALIN / HİZMETİN / YAPIM İŞİNİN", "", "", "", "", ""])
    
    # Headers
    table_data.append(["Sıra", "Cinsi", "Özellikleri", "Miktarı\n(Adet)", "Birim Fiyatı\n(KDV Hariç)", "Toplam Fiyat\n(KDV Hariç)"])
    
    # Data Rows
    for idx, k in enumerate(gecerli_kalemler):
        miktar_text = f"{k.get('miktar', '')} {k.get('birim', '')}".strip()
        table_data.append([
            str(idx+1),
            k.get("cins", ""),
            k.get("ozellik", ""),
            miktar_text,
            "", ""
        ])
        
    # KDV Hariç Toplam
    start_bottom = len(table_data)
    table_data.append(["KDV Hariç Toplam", "", "", "", "", ""])
    
    # Diğer Şartlar Nested Table
    sartlar_data = [
        [Paragraph("DİĞER ŞARTLAR", ParagraphStyle('S', fontName=font_name, fontSize=9)), ""],
        [Paragraph("1- Teslim Süresi", ParagraphStyle('S', fontName=font_name, fontSize=9)), Paragraph(": 1 gün", ParagraphStyle('S', fontName=font_name, fontSize=9))],
        [Paragraph("2- Teslim Edilecek Parti Miktarı", ParagraphStyle('S', fontName=font_name, fontSize=9)), Paragraph(": 1", ParagraphStyle('S', fontName=font_name, fontSize=9))],
        [Paragraph("3- Nakliye ve Sigortanın kime ait olduğu", ParagraphStyle('S', fontName=font_name, fontSize=9)), Paragraph(": Satıcıya", ParagraphStyle('S', fontName=font_name, fontSize=9))],
        [Paragraph("4- Diğer Özel Şartlar", ParagraphStyle('S', fontName=font_name, fontSize=9)), Paragraph(": YOK", ParagraphStyle('S', fontName=font_name, fontSize=9))],
        [Paragraph("5- Uyulması Gereken Standartlar", ParagraphStyle('S', fontName=font_name, fontSize=9)), Paragraph(": TSE", ParagraphStyle('S', fontName=font_name, fontSize=9))],
        [Paragraph("6- Teknik Şartname", ParagraphStyle('S', fontName=font_name, fontSize=9)), Paragraph(": YOK", ParagraphStyle('S', fontName=font_name, fontSize=9))],
        [Paragraph("7- Diğer Hususlar", ParagraphStyle('S', fontName=font_name, fontSize=9)), Paragraph(": YOK", ParagraphStyle('S', fontName=font_name, fontSize=9))]
    ]
    sartlar_table = Table(sartlar_data, colWidths=[65*mm, 35*mm])
    sartlar_table.setStyle(TableStyle([
        ('ALIGN', (0,0), (-1,-1), 'LEFT'),
        ('VALIGN', (0,0), (-1,-1), 'TOP'),
        ('TOPPADDING', (0,0), (-1,-1), 1),
        ('BOTTOMPADDING', (0,0), (-1,-1), 1),
        ('LEFTPADDING', (0,0), (-1,-1), 0),
        ('RIGHTPADDING', (0,0), (-1,-1), 0),
        ('SPAN', (0,0), (1,0)),
    ]))
    
    sig_text = f"Piyasada satış fiyatları<br/>yukarıda gösterilmiştir.<br/><br/><br/>Tasdik Eden<br/>..../..../........<br/><br/>Kişi / Oda / Firmanın<br/>Adı veya Ticaret<br/>Ünvanı - Kaşe İmza"
    
    table_data.append([sartlar_table, "", "", "", Paragraph(sig_text, style_center), ""])
    
    col_widths = [12*mm, 38*mm, 45*mm, 20*mm, 22*mm, 23*mm]
    t = Table(table_data, colWidths=col_widths)
    
    table_style = [
        ('FONTNAME', (0,0), (-1,-1), font_name),
        ('FONTSIZE', (0,0), (-1,-1), 9),
        ('ALIGN', (0,0), (-1,-1), 'LEFT'),
        ('ALIGN', (0,0), (-1,0), 'CENTER'),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('INNERGRID', (0,0), (-1,-1), 0.25, colors.black),
        ('BOX', (0,0), (-1,-1), 0.25, colors.black),
        ('WORDWRAP', (0,0), (-1,-1), True),
        ('SPAN', (0,0), (5,0)),
        
        ('SPAN', (0, start_bottom), (3, start_bottom)),
        ('ALIGN', (0, start_bottom), (3, start_bottom), 'RIGHT'),
        
        ('SPAN', (0, start_bottom+1), (3, start_bottom+1)),
        ('VALIGN', (0, start_bottom+1), (3, start_bottom+1), 'TOP'),
        ('SPAN', (4, start_bottom+1), (5, start_bottom+1)),
        ('ALIGN', (4, start_bottom+1), (5, start_bottom+1), 'CENTER'),
        ('VALIGN', (4, start_bottom+1), (5, start_bottom+1), 'TOP'),
    ]
    t.setStyle(TableStyle(table_style))
    elements.append(t)
    
    elements.append(Spacer(1, 2*mm))
    elements.append(Paragraph("Not :Yaklaşık Maliyet Hesap Cetveline Eklenecektir. ( 4734 Sayılı Kanun Md.9 )", ParagraphStyle('Not', fontName=font_name, fontSize=9)))
    
    doc.build(elements)
    return dosya_adi

def uret_fiyat_isteme_excel(veri, hedef_klasor):
    dosya_adi = os.path.join(hedef_klasor, "Yaklasik_Maliyet_Fiyat_Isteme.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Fiyat Isteme"
    
    ws.page_margins.top = 0.984
    ws.page_margins.bottom = 0.984
    ws.page_margins.left = 0.984
    ws.page_margins.right = 0.984
    
    # Tüm Excel 9 punto
    normal_font = Font(name='Arial', size=9)
    center_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
    left_align = Alignment(horizontal='left', vertical='center', wrap_text=True)
    top_left_align = Alignment(horizontal='left', vertical='top', wrap_text=True)
    top_center_align = Alignment(horizontal='center', vertical='top', wrap_text=True)
    right_align = Alignment(horizontal='right', vertical='center', wrap_text=True)
    justify_align = Alignment(horizontal='justify', vertical='center', wrap_text=True)
    thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
    no_bottom = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'))
    no_top = Border(left=Side(style='thin'), right=Side(style='thin'), bottom=Side(style='thin'))
    no_top_bottom = Border(left=Side(style='thin'), right=Side(style='thin'))
    
    # 1. Resmi Yazı Başlığı
    baslik = veri.get("resmi_baslik", "").replace("\\n", "\n")
    ws.merge_cells('A1:F3')
    cell = ws['A1']
    cell.value = baslik
    cell.font = normal_font
    cell.alignment = center_align
    
    # 2. Sayı, Tarih, Konu
    tarih = format_date(veri.get("belge_tarihi", ""))
    
    ws['A5'] = "Sayı"
    ws['A5'].font = normal_font
    ws['B5'] = f": {veri.get('yazisma_kodu', '')}"
    ws['B5'].font = normal_font
    ws.merge_cells('B5:D5')
    
    ws['F5'] = tarih
    ws['F5'].font = normal_font
    ws['F5'].alignment = right_align
    
    ws['A6'] = "Konu"
    ws['A6'].font = normal_font
    ws['B6'] = ": Yaklaşık Maliyet Fiyatları"
    ws['B6'].font = normal_font
    ws.merge_cells('B6:F6')
    
    # 3. Hitap ve Gövde
    ws.merge_cells('A8:F8')
    ws['A8'] = "Sayın Yetkili"
    ws['A8'].font = normal_font
    ws['A8'].alignment = center_align
    
    metin = "        İdaremizce satın alınması düşünülen aşağıda cinsi, miktarı, özellikleri ve diğer şartları yazılı mal, hizmet ya da yapım işlerinin 4734 Sayılı Kamu İhale Kanunu gereğince, yaklaşık maliyetinin tesbit edilmesinde değerlendirilmek ve KDV hariç olmak üzere piyasada satış fiyatlarının bildirilmesini rica ederim. / ederiz."
    ws.merge_cells('A9:F10')
    ws['A9'] = metin
    ws['A9'].font = normal_font
    ws['A9'].alignment = justify_align
    
    # 4. Komisyon Üyeleri
    row_idx = 12
    komisyon_uyeleri = get_komisyon(veri)
    if len(komisyon_uyeleri) > 0:
        cols = [1, 3, 5]
        ws.merge_cells(start_row=row_idx, start_column=3, end_row=row_idx, end_column=4)
        ws.merge_cells(start_row=row_idx+1, start_column=3, end_row=row_idx+1, end_column=4)
        ws.merge_cells(start_row=row_idx, start_column=5, end_row=row_idx, end_column=6)
        ws.merge_cells(start_row=row_idx+1, start_column=5, end_row=row_idx+1, end_column=6)
        ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=2)
        ws.merge_cells(start_row=row_idx+1, start_column=1, end_row=row_idx+1, end_column=2)
        
        for idx, uye in enumerate(komisyon_uyeleri):
            if idx < 3:
                c_isim = ws.cell(row=row_idx, column=cols[idx])
                c_isim.value = uye
                c_isim.font = normal_font
                c_isim.alignment = center_align
                
                title = "Başkan" if idx == 0 else "Üye"
                c_unvan = ws.cell(row=row_idx+1, column=cols[idx])
                c_unvan.value = title
                c_unvan.font = normal_font
                c_unvan.alignment = center_align
    
    row_idx += 3
    
    # 5. Tablo
    ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=6)
    c_super = ws.cell(row=row_idx, column=1)
    c_super.value = "SATIN ALINACAK MALIN / HİZMETİN / YAPIM İŞİNİN"
    c_super.font = normal_font
    c_super.alignment = center_align
    for col in range(1, 7):
        ws.cell(row=row_idx, column=col).border = thin_border
    row_idx += 1
    
    headers = ["Sıra", "Cinsi", "Özellikleri", "Miktarı\n(Adet)", "Birim Fiyatı\n(KDV Hariç)", "Toplam Fiyat\n(KDV Hariç)"]
    for col_num, header in enumerate(headers, 1):
        c = ws.cell(row=row_idx, column=col_num)
        c.value = header
        c.font = normal_font
        c.alignment = left_align
        c.border = thin_border
        
    kalemler = veri.get("kalemler", [])
    gecerli_kalemler = [k for k in kalemler if k.get("cins", "").strip()]
    if not gecerli_kalemler:
        gecerli_kalemler = [{}]
        
    row_idx += 1
    for i, k in enumerate(gecerli_kalemler):
        miktar_text = f"{k.get('miktar', '')} {k.get('birim', '')}".strip()
        data = [i+1, k.get("cins", ""), k.get("ozellik", ""), miktar_text, "", ""]
        for col_num, val in enumerate(data, 1):
            c = ws.cell(row=row_idx, column=col_num)
            c.value = val if str(val).strip() else ""
            c.font = normal_font
            c.alignment = left_align
            c.border = thin_border
        row_idx += 1
        
    # KDV Hariç Toplam
    ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=4)
    c_tot = ws.cell(row=row_idx, column=1, value="KDV Hariç Toplam")
    c_tot.font = normal_font
    c_tot.alignment = right_align
    for col in range(1, 7):
        ws.cell(row=row_idx, column=col).border = thin_border
        
    row_idx += 1
    
    # Diğer Şartlar & Sig Box
    start_sart_row = row_idx
    
    ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=4)
    c_ds = ws.cell(row=row_idx, column=1, value="DİĞER ŞARTLAR")
    c_ds.font = normal_font
    c_ds.alignment = top_left_align
    row_idx += 1
    
    sartlar = [
        ("1- Teslim Süresi", ": 1 gün"),
        ("2- Teslim Edilecek Parti Miktarı", ": 1"),
        ("3- Nakliye ve Sigortanın kime ait olduğu", ": Satıcıya"),
        ("4- Diğer Özel Şartlar", ": YOK"),
        ("5- Uyulması Gereken Standartlar", ": TSE"),
        ("6- Teknik Şartname", ": YOK"),
        ("7- Diğer Hususlar", ": YOK")
    ]
    
    for i, s in enumerate(sartlar):
        ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=3)
        c_n = ws.cell(row=row_idx, column=1, value=s[0])
        c_n.font = normal_font
        c_n.alignment = top_left_align
        
        c_v = ws.cell(row=row_idx, column=4, value=s[1])
        c_v.font = normal_font
        c_v.alignment = top_left_align
        row_idx += 1
        
    for r in range(start_sart_row, row_idx):
        for c in range(1, 5):
            b = no_top_bottom
            if r == start_sart_row: b = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'))
            if r == row_idx - 1: b = Border(left=Side(style='thin'), right=Side(style='thin'), bottom=Side(style='thin'))
            ws.cell(row=r, column=c).border = b
            
    # Signature Box
    ws.merge_cells(start_row=start_sart_row, start_column=5, end_row=row_idx-1, end_column=6)
    sig_text = f"Piyasada satış fiyatları\nyukarıda gösterilmiştir.\n\n\nTasdik Eden\n..../..../........\n\nKişi / Oda / Firmanın\nAdı veya Ticaret\nÜnvanı - Kaşe İmza"
    c_sig = ws.cell(row=start_sart_row, column=5, value=sig_text)
    c_sig.alignment = top_center_align
    c_sig.font = normal_font
    
    for r in range(start_sart_row, row_idx):
        for c in range(5, 7):
            b = no_top_bottom
            if r == start_sart_row: b = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'))
            if r == row_idx - 1: b = Border(left=Side(style='thin'), right=Side(style='thin'), bottom=Side(style='thin'))
            ws.cell(row=r, column=c).border = b

    ws.cell(row=row_idx, column=1, value="Not :Yaklaşık Maliyet Hesap Cetveline Eklenecektir. ( 4734 Sayılı Kanun Md.9 )").font = normal_font
    
    ws.column_dimensions['A'].width = 8
    ws.column_dimensions['B'].width = 25
    ws.column_dimensions['C'].width = 25
    ws.column_dimensions['D'].width = 15
    ws.column_dimensions['E'].width = 15
    ws.column_dimensions['F'].width = 18
    
    wb.save(dosya_adi)
    return dosya_adi
