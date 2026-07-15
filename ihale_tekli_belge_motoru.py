import os
import datetime
import platform
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side

# Font Setup
font_name = 'Helvetica-TR'
font_name_bold = 'Helvetica-TR-Bold'
try:
    pdfmetrics.registerFont(TTFont('Helvetica-TR', 'C:\\Windows\\Fonts\\arial.ttf'))
    pdfmetrics.registerFont(TTFont('Helvetica-TR-Bold', 'C:\\Windows\\Fonts\\arialbd.ttf'))
except:
    font_name = 'Helvetica'
    font_name_bold = 'Helvetica-Bold'

def format_date(date_str):
    if not date_str:
        return datetime.datetime.now().strftime("%d.%m.%Y")
    try:
        dt = datetime.datetime.strptime(date_str, "%Y-%m-%d")
        return dt.strftime("%d.%m.%Y")
    except:
        return date_str

def belge_uret(veri, pdf_yol):
    belge_tipi = veri.get("belge_tipi", "")
    fmt = veri.get("format", "pdf")
    
    if not os.path.exists(pdf_yol):
        os.makedirs(pdf_yol, exist_ok=True)
        
    klasor_adi = f"Ihale_{datetime.datetime.now().strftime('%Y%m%d_%H%M%S')}"
    hedef_klasor = os.path.join(pdf_yol, klasor_adi)
    os.makedirs(hedef_klasor, exist_ok=True)
    
    if belge_tipi == "fiyat_isteme":
        if fmt == "pdf":
            return uret_fiyat_isteme_pdf(veri, hedef_klasor)
        elif fmt == "excel":
            return uret_fiyat_isteme_excel(veri, hedef_klasor)
    return None

def get_komisyon(veri):
    kom = veri.get("komisyon", {})
    uyeler = [
        kom.get("ihale_kom_yaklasik_1", ""),
        kom.get("ihale_kom_yaklasik_2", ""),
        kom.get("ihale_kom_yaklasik_3", "")
    ]
    # Filter empty
    uyeler = [u for u in uyeler if u]
    return uyeler

def uret_fiyat_isteme_pdf(veri, hedef_klasor):
    dosya_adi = os.path.join(hedef_klasor, "Fiyat_Isteme_Belgesi.pdf")
    doc = SimpleDocTemplate(
        dosya_adi,
        pagesize=A4,
        rightMargin=15*mm,
        leftMargin=15*mm,
        topMargin=15*mm,
        bottomMargin=15*mm
    )
    
    styles = getSampleStyleSheet()
    style_normal = ParagraphStyle('Normal_TR', fontName=font_name, fontSize=11, leading=14)
    style_bold = ParagraphStyle('Bold_TR', fontName=font_name_bold, fontSize=11, leading=14)
    style_center = ParagraphStyle('Center_TR', fontName=font_name_bold, fontSize=12, alignment=1, leading=16)
    
    elements = []
    
    # 1. Resmi Yazı Başlığı
    baslik_satirlari = veri.get("resmi_baslik", "").split("\\n")
    for satir in baslik_satirlari:
        elements.append(Paragraph(satir.strip(), style_center))
    
    elements.append(Spacer(1, 15*mm))
    
    # 2. Sayı, Konu, Tarih
    tarih = format_date(veri.get("belge_tarihi", ""))
    elements.append(Paragraph(f"<b>Sayı :</b> {veri.get('yazisma_kodu', '')}", style_normal))
    elements.append(Paragraph(f"<b>Konu :</b> {veri.get('ihale_konusu', '')}", style_normal))
    
    elements.append(Spacer(1, 5*mm))
    # Tarih sağa hizalı
    p_tarih = Paragraph(f"{tarih}", ParagraphStyle('Right', fontName=font_name, fontSize=11, alignment=2))
    elements.append(p_tarih)
    
    elements.append(Spacer(1, 10*mm))
    
    # 3. İçerik Metni
    metin = "Aşağıda cins, özellik ve miktarları belirtilen mal/hizmet alımı işi için piyasada satış fiyatlarının bildirilmesini rica ederiz."
    elements.append(Paragraph(metin, style_normal))
    elements.append(Spacer(1, 5*mm))
    
    # 4. Tablo
    kalemler = veri.get("kalemler", [])
    table_data = [["Sıra", "Malın/Hizmetin Cinsi", "Özelliği", "Miktarı", "Birim", "Birim Fiyat\\n(KDV Hariç)", "Tutar"]]
    
    for idx, k in enumerate(kalemler):
        table_data.append([
            str(idx+1),
            k.get("cins", ""),
            k.get("ozellik", ""),
            str(k.get("miktar", "")),
            k.get("birim", ""),
            "", ""
        ])
    
    col_widths = [10*mm, 50*mm, 45*mm, 15*mm, 15*mm, 22*mm, 23*mm]
    t = Table(table_data, colWidths=col_widths)
    t.setStyle(TableStyle([
        ('FONTNAME', (0,0), (-1,-1), font_name),
        ('FONTSIZE', (0,0), (-1,-1), 9),
        ('ALIGN', (0,0), (-1,-1), 'CENTER'),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('INNERGRID', (0,0), (-1,-1), 0.25, colors.black),
        ('BOX', (0,0), (-1,-1), 0.25, colors.black),
        ('FONTNAME', (0,0), (-1,0), font_name_bold),
        ('BACKGROUND', (0,0), (-1,0), colors.lightgrey),
        ('WORDWRAP', (0,0), (-1,-1), True)
    ]))
    elements.append(t)
    elements.append(Spacer(1, 20*mm))
    
    # 5. Komisyon
    komisyon_uyeleri = get_komisyon(veri)
    elements.append(Paragraph("<b>YAKLAŞIK MALİYET TESPİT KOMİSYONU</b>", style_center))
    elements.append(Spacer(1, 10*mm))
    
    if len(komisyon_uyeleri) > 0:
        imza_data = [[]]
        imza_gorev = [[]]
        for idx, uye in enumerate(komisyon_uyeleri):
            title = "Okul Müdürü\\n(Başkan)" if idx == 0 else "Müdür Yardımcısı\\n(Üye)"
            imza_data[0].append(Paragraph(f"<b>{uye}</b>", ParagraphStyle('C', fontName=font_name_bold, fontSize=10, alignment=1)))
            imza_gorev[0].append(Paragraph(title, ParagraphStyle('C', fontName=font_name, fontSize=10, alignment=1)))
        
        col_w = 180*mm / len(komisyon_uyeleri)
        sig_table = Table(imza_data + imza_gorev, colWidths=[col_w] * len(komisyon_uyeleri))
        sig_table.setStyle(TableStyle([
            ('ALIGN', (0,0), (-1,-1), 'CENTER'),
            ('VALIGN', (0,0), (-1,-1), 'TOP'),
        ]))
        elements.append(sig_table)
        
    doc.build(elements)
    return dosya_adi

def uret_fiyat_isteme_excel(veri, hedef_klasor):
    dosya_adi = os.path.join(hedef_klasor, "Fiyat_Isteme_Belgesi.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Fiyat Isteme"
    
    bold_font = Font(name='Arial', size=11, bold=True)
    normal_font = Font(name='Arial', size=11)
    center_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
    left_align = Alignment(horizontal='left', vertical='center', wrap_text=True)
    thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
    
    baslik = veri.get("resmi_baslik", "").replace("\\n", "\n")
    ws.merge_cells('A1:G3')
    cell = ws['A1']
    cell.value = baslik
    cell.font = bold_font
    cell.alignment = center_align
    
    ws['A5'] = f"Sayı : {veri.get('yazisma_kodu', '')}"
    ws['A5'].font = bold_font
    ws['A6'] = f"Konu : {veri.get('ihale_konusu', '')}"
    ws['A6'].font = bold_font
    
    ws['G5'] = format_date(veri.get("belge_tarihi", ""))
    ws['G5'].font = normal_font
    ws['G5'].alignment = Alignment(horizontal='right')
    
    metin = "Aşağıda cins, özellik ve miktarları belirtilen mal/hizmet alımı işi için piyasada satış fiyatlarının bildirilmesini rica ederiz."
    ws.merge_cells('A8:G8')
    ws['A8'] = metin
    ws['A8'].font = normal_font
    ws['A8'].alignment = left_align
    
    headers = ["Sıra", "Malın/Hizmetin Cinsi", "Özelliği", "Miktarı", "Birim", "Birim Fiyat (KDV Hariç)", "Tutar"]
    for col_num, header in enumerate(headers, 1):
        c = ws.cell(row=10, column=col_num)
        c.value = header
        c.font = bold_font
        c.alignment = center_align
        c.border = thin_border
        
    kalemler = veri.get("kalemler", [])
    row_idx = 11
    for i, k in enumerate(kalemler):
        data = [i+1, k.get("cins", ""), k.get("ozellik", ""), k.get("miktar", ""), k.get("birim", ""), "", ""]
        for col_num, val in enumerate(data, 1):
            c = ws.cell(row=row_idx, column=col_num)
            c.value = val
            c.font = normal_font
            c.alignment = center_align
            c.border = thin_border
        row_idx += 1
        
    ws.column_dimensions['A'].width = 8
    ws.column_dimensions['B'].width = 30
    ws.column_dimensions['C'].width = 25
    ws.column_dimensions['D'].width = 12
    ws.column_dimensions['E'].width = 12
    ws.column_dimensions['F'].width = 20
    ws.column_dimensions['G'].width = 20
    
    row_idx += 3
    ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=7)
    c = ws.cell(row=row_idx, column=1)
    c.value = "YAKLAŞIK MALİYET TESPİT KOMİSYONU"
    c.font = bold_font
    c.alignment = center_align
    
    row_idx += 2
    komisyon_uyeleri = get_komisyon(veri)
    if len(komisyon_uyeleri) > 0:
        cols = [2, 4, 6]
        for idx, uye in enumerate(komisyon_uyeleri):
            if idx < 3:
                c_isim = ws.cell(row=row_idx, column=cols[idx])
                c_isim.value = uye
                c_isim.font = bold_font
                c_isim.alignment = center_align
                
                title = "Okul Müdürü\n(Başkan)" if idx == 0 else "Müdür Yardımcısı\n(Üye)"
                c_unvan = ws.cell(row=row_idx+1, column=cols[idx])
                c_unvan.value = title
                c_unvan.font = normal_font
                c_unvan.alignment = center_align

    wb.save(dosya_adi)
    return dosya_adi
