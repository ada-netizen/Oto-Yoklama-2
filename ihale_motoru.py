import os
import re
import datetime
from reportlab.lib.pagesizes import A4, landscape
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side

def format_date(date_str):
    if not date_str:
        return ""
    try:
        if "T" in date_str:
            d = datetime.datetime.strptime(date_str.split("T")[0], "%Y-%m-%d")
        else:
            d = datetime.datetime.strptime(date_str, "%Y-%m-%d")
        return d.strftime("%d.%m.%Y")
    except:
        return date_str

def get_komisyon(veri):
    kom = veri.get("komisyon", {})
    uyeler = []
    if isinstance(kom, list):
        for item in kom:
            ad = item.get("ad_soyad", "").strip()
            if ad: uyeler.append(ad)
    elif isinstance(kom, dict):
        belge_tipi = veri.get("belge_tipi", "")
        if belge_tipi in ["fiyat_isteme", "ozel_fiyat_isteme", "piyasa_arastirmasi"]:
            keys = ["ihale_kom_piyasa_1", "ihale_kom_piyasa_2", "ihale_kom_piyasa_3"]
        else:
            keys = ["ihale_kom_yaklasik_1", "ihale_kom_yaklasik_2", "ihale_kom_yaklasik_3"]
        
        for k in keys:
            ad = kom.get(k, "")
            if ad and ad.strip(): 
                uyeler.append(ad.strip())
    return uyeler


def get_ihale_konu(veri):
    for key in ("konu", "ihale_konusu", "ihale_adi", "is_adi"):
        val = veri.get(key)
        if isinstance(val, str) and val.strip():
            return val.strip()
    return "Kırtasiye Alımı"


def _ilk_harfleri_buyuk_yap(metin):
    if not metin:
        return ""
    kelimeler = []
    for kelime in str(metin).replace("\\n", " ").split():
        if not kelime:
            continue
        if len(kelime) == 1:
            kelimeler.append(kelime.upper())
        else:
            kelimeler.append(kelime[:1].upper() + kelime[1:].lower())
    return " ".join(kelimeler)


def _idare_adi_olustur(baslik=None, okul_adi=None):
    adaylar = []
    for deger in [baslik, okul_adi]:
        if isinstance(deger, str):
            metin = deger.replace("\\n", "\n").strip()
            if metin:
                adaylar.append(" ".join(s.strip() for s in metin.split("\n") if s.strip()))
    if not adaylar:
        return "İdare Müdürlüğü"

    idare = _ilk_harfleri_buyuk_yap(adaylar[0].strip())
    idare_lower = idare.lower()

    if "müdürlüğü" in idare_lower or "müdürlük" in idare_lower:
        return idare
    if re.search(r"(okulu|lisesi|ortaokulu|ilkokulu|anaokulu|merkezi|idaresi|birimi|kurumu|kaymakamlığı|valiliği|başkanlığı)$", idare, flags=re.IGNORECASE):
        return f"{idare} Müdürlüğü"
    if any(pat in idare_lower for pat in ["okul", "lise", "ortaokul", "ilkokul", "anaokul", "kaymakamlık", "valilik", "başkanlık", "kurum", "daire", "bölge", "ofis"]):
        return f"{idare} Müdürlüğü"
    return idare


def _ay_adi_tr(ay_no):
    adlar = {
        1: "Ocak", 2: "Şubat", 3: "Mart", 4: "Nisan", 5: "Mayıs", 6: "Haziran",
        7: "Temmuz", 8: "Ağustos", 9: "Eylül", 10: "Ekim", 11: "Kasım", 12: "Aralık"
    }
    return adlar.get(ay_no, str(ay_no))


def _temiz_klasor_adi(metin):
    temiz = re.sub(r"[^\w\sçÇğĞıİöÖşŞüÜ-]", " ", metin, flags=re.UNICODE)
    temiz = re.sub(r"\s+", " ", temiz).strip()
    return temiz


def get_ihale_klasoru(veri, base_klasor):
    tarih_metin = veri.get("belge_tarihi") or veri.get("tarih") or datetime.datetime.now().strftime("%Y-%m-%d")
    tarih_dt = None
    for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%Y/%m/%d"):
        try:
            tarih_dt = datetime.datetime.strptime(str(tarih_metin).split("T")[0], fmt)
            break
        except ValueError:
            continue
    if tarih_dt is None:
        tarih_dt = datetime.datetime.now()

    konu = get_ihale_konu(veri)
    klasor_adi = f"{tarih_dt.year} {_ay_adi_tr(tarih_dt.month)} {_temiz_klasor_adi(konu)}"
    ihale_ana = os.path.join(base_klasor, "Ihale")
    os.makedirs(ihale_ana, exist_ok=True)
    hedef = os.path.join(ihale_ana, klasor_adi)
    os.makedirs(hedef, exist_ok=True)
    return hedef


def uret_fiyat_isteme_pdf(veri, hedef_klasor):
    timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    dosya_adi = os.path.join(hedef_klasor, f"Fiyat_Isteme_{timestamp}.pdf")
    
    try:
        pdfmetrics.registerFont(TTFont('Arial_TR', r'C:\Windows\Fonts\arial.ttf'))
        font_name = 'Arial_TR'
    except:
        font_name = 'Helvetica'
        
    doc = SimpleDocTemplate(
        dosya_adi,
        pagesize=A4,
        rightMargin=25*mm,
        leftMargin=25*mm,
        topMargin=12.5*mm,
        bottomMargin=12.5*mm
    )
    
    # Tüm PDF 9 punto
    style_normal = ParagraphStyle('Normal_TR', fontName=font_name, fontSize=9, leading=11)
    style_justify = ParagraphStyle('Justify_TR', fontName=font_name, fontSize=9, alignment=4, leading=11, firstLineIndent=12*mm)
    style_center = ParagraphStyle('Center_TR', fontName=font_name, fontSize=9, alignment=1, leading=11)
    style_right = ParagraphStyle('Right_TR', fontName=font_name, fontSize=9, alignment=2, leading=11)
    
    elements = []
    
    # 1. Resmi Yazı Başlığı
    baslik = veri.get("resmi_baslik", "")
    baslik_satirlari = baslik.split("\n")
    for satir in baslik_satirlari:
        elements.append(Paragraph(satir.strip(), style_center))
    
    elements.append(Spacer(1, 8*mm))
    
    # 2. Sayı, Tarih, Konu
    tarih = format_date(veri.get("belge_tarihi", ""))
    sayi = veri.get("yazisma_kodu", "")
    
    tbl_sayi_konu = Table([
        [Paragraph("Sayı", style_normal), Paragraph(":", style_normal), Paragraph(sayi, style_normal), Paragraph(tarih, style_right)],
        [Paragraph("Konu", style_normal), Paragraph(":", style_normal), Paragraph("Yaklaşık Maliyet Fiyatları", style_normal), ""]
    ], colWidths=[12*mm, 3*mm, 95*mm, 50*mm])
    
    tbl_sayi_konu.setStyle(TableStyle([
        ('ALIGN', (0,0), (0,-1), 'LEFT'),
        ('ALIGN', (1,0), (1,-1), 'LEFT'),
        ('ALIGN', (2,0), (2,-1), 'LEFT'),
        ('ALIGN', (3,0), (3,0), 'RIGHT'),
        ('VALIGN', (0,0), (-1,-1), 'TOP'),
        ('BOTTOMPADDING', (0,0), (-1,-1), 1),
        ('TOPPADDING', (0,0), (-1,-1), 1),
        ('LEFTPADDING', (0,0), (-1,-1), 0),
        ('RIGHTPADDING', (0,0), (-1,-1), 0),
    ]))
    elements.append(tbl_sayi_konu)
    
    elements.append(Spacer(1, 8*mm))
    
    # 3. Hitap ve Gövde
    elements.append(Paragraph("Sayın Yetkili", style_center))
    elements.append(Spacer(1, 4*mm))
    metin = "İdaremizce satın alınması düşünülen aşağıda cinsi, miktarı, özellikleri ve diğer şartları yazılı mal, hizmet ya da yapım işlerinin 4734 Sayılı Kamu İhale Kanunu gereğince, yaklaşık maliyetinin tesbit edilmesinde değerlendirilmek ve KDV hariç olmak üzere piyasada satış fiyatlarının bildirilmesini rica ederim. / ederiz."
    elements.append(Paragraph(metin, style_justify))
    elements.append(Spacer(1, 8*mm))
    
    # 4. Komisyon İmzaları (Boşluk kaldırıldı, parantezler kaldırıldı)
    komisyon_uyeleri = get_komisyon(veri)
    if len(komisyon_uyeleri) > 0:
        imza_data = [[]]
        for idx, uye in enumerate(komisyon_uyeleri):
            title = "Başkan" if idx == 0 else "Üye"
            imza_data[0].append(Paragraph(f"{uye}<br/>{title}", style_center))
        
        col_w = 160*mm / len(komisyon_uyeleri)
        sig_table = Table(imza_data, colWidths=[col_w] * len(komisyon_uyeleri))
        sig_table.setStyle(TableStyle([
            ('ALIGN', (0,0), (-1,-1), 'CENTER'),
            ('VALIGN', (0,0), (-1,-1), 'TOP'),
        ]))
        elements.append(sig_table)
    
    elements.append(Spacer(1, 8*mm))
    
    # 5. Tablo
    kalemler = veri.get("kalemler", [])
    gecerli_kalemler = [k for k in kalemler if k.get("cins", "").strip()]
    if not gecerli_kalemler:
        gecerli_kalemler = [{}] 
        
    table_data = []
    
    # Super Header
    table_data.append(["SATIN ALINACAK MALIN / HİZMETİN / YAPIM İŞİNİN", "", "", "", "", ""])
    
    # Headers
    table_data.append(["Sıra", "Cinsi", "Özellikleri", "Miktarı\n(Adet)", "Birim Fiyatı\n(KDV Hariç)", "Toplam Fiyat\n(KDV Hariç)"])
    
    # Data Rows
    for idx, k in enumerate(gecerli_kalemler):
        val = str(k.get("miktar", ""))
        try:
            if float(val) == int(float(val)): val = str(int(float(val)))
        except:
            pass
        table_data.append([
            str(idx+1),
            k.get("cins", ""),
            k.get("ozellik", ""),
            val,
            "", ""
        ])
        
    # KDV Hariç Toplam
    start_bottom = len(table_data)
    table_data.append(["KDV Hariç Toplam", "", "", "", "", ""])
    
    # Diğer Şartlar Nested Table
    sartlar_data = [
        [Paragraph("DİĞER ŞARTLAR", ParagraphStyle('S', fontName=font_name, fontSize=9)), ""],
        [Paragraph("1- Teslim Süresi", ParagraphStyle('S', fontName=font_name, fontSize=9)), Paragraph(": 1 gün", ParagraphStyle('S', fontName=font_name, fontSize=9))],
        [Paragraph("2- Teslim Edilecek Parti Miktarı", ParagraphStyle('S', fontName=font_name, fontSize=9)), Paragraph(": 1", ParagraphStyle('S', fontName=font_name, fontSize=9))],
        [Paragraph("3- Nakliye ve Sigortanın kime ait olduğu", ParagraphStyle('S', fontName=font_name, fontSize=9)), Paragraph(": Satıcıya", ParagraphStyle('S', fontName=font_name, fontSize=9))],
        [Paragraph("4- Diğer Özel Şartlar", ParagraphStyle('S', fontName=font_name, fontSize=9)), Paragraph(": YOK", ParagraphStyle('S', fontName=font_name, fontSize=9))],
        [Paragraph("5- Uyulması Gereken Standartlar", ParagraphStyle('S', fontName=font_name, fontSize=9)), Paragraph(": TSE", ParagraphStyle('S', fontName=font_name, fontSize=9))],
        [Paragraph("6- Teknik Şartname", ParagraphStyle('S', fontName=font_name, fontSize=9)), Paragraph(": YOK", ParagraphStyle('S', fontName=font_name, fontSize=9))],
        [Paragraph("7- Diğer Hususlar", ParagraphStyle('S', fontName=font_name, fontSize=9)), Paragraph(": YOK", ParagraphStyle('S', fontName=font_name, fontSize=9))]
    ]
    sartlar_table = Table(sartlar_data, colWidths=[65*mm, 35*mm])
    sartlar_table.setStyle(TableStyle([
        ('ALIGN', (0,0), (-1,-1), 'LEFT'),
        ('VALIGN', (0,0), (-1,-1), 'TOP'),
        ('TOPPADDING', (0,0), (-1,-1), 1),
        ('BOTTOMPADDING', (0,0), (-1,-1), 1),
        ('LEFTPADDING', (0,0), (-1,-1), 0),
        ('RIGHTPADDING', (0,0), (-1,-1), 0),
        ('SPAN', (0,0), (1,0)),
    ]))
    
    sig_text = f"Piyasada satış fiyatları<br/>yukarıda gösterilmiştir.<br/><br/><br/>Tasdik Eden<br/>..../..../........<br/><br/>Kişi / Oda / Firmanın<br/>Adı veya Ticaret<br/>Ünvanı - Kaşe İmza"
    
    table_data.append([sartlar_table, "", "", "", Paragraph(sig_text, style_center), ""])
    
    col_widths = [12*mm, 38*mm, 45*mm, 20*mm, 22*mm, 23*mm]
    t = Table(table_data, colWidths=col_widths)
    
    table_style = [
        ('FONTNAME', (0,0), (-1,-1), font_name),
        ('FONTSIZE', (0,0), (-1,-1), 9),
        ('ALIGN', (0,0), (-1,-1), 'LEFT'),
        ('ALIGN', (0,0), (-1,0), 'CENTER'),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('INNERGRID', (0,0), (-1,-1), 0.25, colors.black),
        ('BOX', (0,0), (-1,-1), 0.25, colors.black),
        
        ('SPAN', (0,0), (5,0)),
        
        ('SPAN', (0, start_bottom), (4, start_bottom)),
        ('ALIGN', (0, start_bottom), (4, start_bottom), 'RIGHT'),
        
        ('SPAN', (0, start_bottom+1), (3, start_bottom+1)),
        ('VALIGN', (0, start_bottom+1), (3, start_bottom+1), 'TOP'),
        ('SPAN', (4, start_bottom+1), (5, start_bottom+1)),
        ('ALIGN', (4, start_bottom+1), (5, start_bottom+1), 'CENTER'),
        ('VALIGN', (4, start_bottom+1), (5, start_bottom+1), 'TOP'),
    ]
    t.setStyle(TableStyle(table_style))
    elements.append(t)
    
    elements.append(Spacer(1, 2*mm))
    elements.append(Paragraph("Not :Yaklaşık Maliyet Hesap Cetveline Eklenecektir. ( 4734 Sayılı Kanun Md.9 )", ParagraphStyle('Not', fontName=font_name, fontSize=9)))
    
    doc.build(elements)
    return dosya_adi

def uret_fiyat_isteme_excel(veri, hedef_klasor):
    timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    dosya_adi = os.path.join(hedef_klasor, f"Fiyat_Isteme_{timestamp}.xlsx")
    
    wb = Workbook()
    ws = wb.active
    ws.title = "Fiyat Isteme"
    
    ws.page_margins.top = 0.492
    ws.page_margins.bottom = 0.492
    ws.page_margins.left = 0.984
    ws.page_margins.right = 0.984
    
    # Tüm Excel 9 punto
    normal_font = Font(name='Arial', size=9)
    center_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
    left_align = Alignment(horizontal='left', vertical='center', wrap_text=True)
    top_left_align = Alignment(horizontal='left', vertical='top', wrap_text=True)
    top_center_align = Alignment(horizontal='center', vertical='top', wrap_text=True)
    right_align = Alignment(horizontal='right', vertical='center', wrap_text=True)
    justify_align = Alignment(horizontal='justify', vertical='center', wrap_text=True)
    thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
    no_bottom = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'))
    no_top = Border(left=Side(style='thin'), right=Side(style='thin'), bottom=Side(style='thin'))
    no_top_bottom = Border(left=Side(style='thin'), right=Side(style='thin'))
    
    # 1. Resmi Yazı Başlığı
    baslik = veri.get("resmi_baslik", "").replace("\\n", "\n")
    baslik_satirlari = [s.strip() for s in baslik.split("\n") if s.strip()]
    
    row_idx = 1
    for satir in baslik_satirlari:
        ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=6)
        c = ws.cell(row=row_idx, column=1)
        c.value = satir
        c.font = normal_font
        c.alignment = center_align
        row_idx += 1
        
    row_idx += 1
    
    # 2. Sayı, Tarih, Konu
    tarih = format_date(veri.get("belge_tarihi", ""))
    
    ws[f'A{row_idx}'] = "Sayı"
    ws[f'A{row_idx}'].font = normal_font
    ws[f'B{row_idx}'] = f": {veri.get('yazisma_kodu', '')}"
    ws[f'B{row_idx}'].font = normal_font
    ws.merge_cells(f'B{row_idx}:D{row_idx}')
    
    ws[f'F{row_idx}'] = tarih
    ws[f'F{row_idx}'].font = normal_font
    ws[f'F{row_idx}'].alignment = right_align
    row_idx += 1
    
    ws[f'A{row_idx}'] = "Konu"
    ws[f'A{row_idx}'].font = normal_font
    ws[f'B{row_idx}'] = ": Yaklaşık Maliyet Fiyatları"
    ws[f'B{row_idx}'].font = normal_font
    ws.merge_cells(f'B{row_idx}:F{row_idx}')
    row_idx += 2
    
    # 3. Hitap ve Gövde
    ws.merge_cells(f'A{row_idx}:F{row_idx}')
    ws[f'A{row_idx}'] = "Sayın Yetkili"
    ws[f'A{row_idx}'].font = normal_font
    ws[f'A{row_idx}'].alignment = center_align
    row_idx += 1
    
    metin = "        İdaremizce satın alınması düşünülen aşağıda cinsi, miktarı, özellikleri ve diğer şartları yazılı mal, hizmet ya da yapım işlerinin 4734 Sayılı Kamu İhale Kanunu gereğince, yaklaşık maliyetinin tesbit edilmesinde değerlendirilmek ve KDV hariç olmak üzere piyasada satış fiyatlarının bildirilmesini rica ederim. / ederiz."
    ws.merge_cells(f'A{row_idx}:F{row_idx+1}')
    ws[f'A{row_idx}'] = metin
    ws[f'A{row_idx}'].font = normal_font
    ws[f'A{row_idx}'].alignment = justify_align
    ws.row_dimensions[row_idx].height = 30
    row_idx += 3
    
    # 4. Komisyon İmzaları
    komisyon_uyeleri = get_komisyon(veri)
    if len(komisyon_uyeleri) > 0:
        if len(komisyon_uyeleri) == 1:
            cols = [3]
        elif len(komisyon_uyeleri) == 2:
            cols = [2, 5]
        elif len(komisyon_uyeleri) == 3:
            cols = [2, 4, 6]
        else:
            cols = [1, 3, 5, 6]
            
        for idx, uye in enumerate(komisyon_uyeleri[:4]):
            c_isim = ws.cell(row=row_idx, column=cols[idx])
            c_isim.value = uye
            c_isim.font = normal_font
            c_isim.alignment = center_align
            
            title = "Başkan" if idx == 0 else "Üye"
            c_unvan = ws.cell(row=row_idx+1, column=cols[idx])
            c_unvan.value = title
            c_unvan.font = normal_font
            c_unvan.alignment = center_align
    row_idx += 3
    
    # 5. Tablo
    ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=6)
    c_super = ws.cell(row=row_idx, column=1)
    c_super.value = "SATIN ALINACAK MALIN / HİZMETİN / YAPIM İŞİNİN"
    c_super.font = normal_font
    c_super.alignment = center_align
    for col in range(1, 7):
        ws.cell(row=row_idx, column=col).border = thin_border
    row_idx += 1
    
    headers = ["Sıra", "Cinsi", "Özellikleri", "Miktarı\n(Adet)", "Birim Fiyatı\n(KDV Hariç)", "Toplam Fiyat\n(KDV Hariç)"]
    for col_num, header in enumerate(headers, 1):
        c = ws.cell(row=row_idx, column=col_num)
        c.value = header
        c.font = normal_font
        c.alignment = left_align
        c.border = thin_border
        
    kalemler = veri.get("kalemler", [])
    gecerli_kalemler = [k for k in kalemler if k.get("cins", "").strip()]
    if not gecerli_kalemler:
        gecerli_kalemler = [{}]
        
    row_idx += 1
    for idx, k in enumerate(gecerli_kalemler):
        val = str(k.get("miktar", ""))
        try:
            if float(val) == int(float(val)): val = str(int(float(val)))
        except: pass
        row_data = [str(idx+1), k.get("cins", ""), k.get("ozellik", ""), val, "", ""]
        for col_num, val in enumerate(row_data, 1):
            c = ws.cell(row=row_idx, column=col_num)
            c.value = val if str(val).strip() else ""
            c.font = normal_font
            c.alignment = left_align
            c.border = thin_border
        row_idx += 1
        
    # KDV Hariç Toplam
    ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=4)
    c_tot = ws.cell(row=row_idx, column=1, value="KDV Hariç Toplam")
    c_tot.font = normal_font
    c_tot.alignment = right_align
    for col in range(1, 7):
        ws.cell(row=row_idx, column=col).border = thin_border
        
    row_idx += 1
    
    # Diğer Şartlar & Sig Box
    start_sart_row = row_idx
    
    ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=4)
    c_ds = ws.cell(row=row_idx, column=1, value="DİĞER ŞARTLAR")
    c_ds.font = normal_font
    c_ds.alignment = top_left_align
    row_idx += 1
    
    sartlar = [
        ("1- Teslim Süresi", ": 1 gün"),
        ("2- Teslim Edilecek Parti Miktarı", ": 1"),
        ("3- Nakliye ve Sigortanın kime ait olduğu", ": Satıcıya"),
        ("4- Diğer Özel Şartlar", ": YOK"),
        ("5- Uyulması Gereken Standartlar", ": TSE"),
        ("6- Teknik Şartname", ": YOK"),
        ("7- Diğer Hususlar", ": YOK")
    ]
    
    for i, s in enumerate(sartlar):
        ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=3)
        c_n = ws.cell(row=row_idx, column=1, value=s[0])
        c_n.font = normal_font
        c_n.alignment = top_left_align
        
        c_v = ws.cell(row=row_idx, column=4, value=s[1])
        c_v.font = normal_font
        c_v.alignment = top_left_align
        row_idx += 1
        
    for r in range(start_sart_row, row_idx):
        for c in range(1, 5):
            b = no_top_bottom
            if r == start_sart_row: b = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'))
            if r == row_idx - 1: b = Border(left=Side(style='thin'), right=Side(style='thin'), bottom=Side(style='thin'))
            ws.cell(row=r, column=c).border = b
            
    # Signature Box
    ws.merge_cells(start_row=start_sart_row, start_column=5, end_row=row_idx-1, end_column=6)
    sig_text = f"Piyasada satış fiyatları\nyukarıda gösterilmiştir.\n\n\nTasdik Eden\n..../..../........\n\nKişi / Oda / Firmanın\nAdı veya Ticaret\nÜnvanı - Kaşe İmza"
    c_sig = ws.cell(row=start_sart_row, column=5, value=sig_text)
    c_sig.alignment = top_center_align
    c_sig.font = normal_font
    
    for r in range(start_sart_row, row_idx):
        for c in range(5, 7):
            b = no_top_bottom
            if r == start_sart_row: b = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'))
            if r == row_idx - 1: b = Border(left=Side(style='thin'), right=Side(style='thin'), bottom=Side(style='thin'))
            ws.cell(row=r, column=c).border = b

    ws.cell(row=row_idx, column=1, value="Not :Yaklaşık Maliyet Hesap Cetveline Eklenecektir. ( 4734 Sayılı Kanun Md.9 )").font = normal_font
    
    ws.column_dimensions['A'].width = 8
    ws.column_dimensions['B'].width = 25
    ws.column_dimensions['C'].width = 25
    ws.column_dimensions['D'].width = 15
    ws.column_dimensions['E'].width = 15
    ws.column_dimensions['F'].width = 18
    
    wb.save(dosya_adi)
    return dosya_adi

# ----- YENI EKLENEN YAKLAŞIK MALİYET HESAP CETVELİ -----

def uret_yaklasik_maliyet_pdf(veri, hedef_klasor):
    import os
    import datetime
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.pdfgen import canvas
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.lib import colors
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm

    timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    dosya_adi = os.path.join(hedef_klasor, f"Yaklasik_Maliyet_Hesap_Cetveli_{timestamp}.pdf")
    
    try:
        pdfmetrics.registerFont(TTFont('Arial_TR', r'C:\Windows\Fonts\arial.ttf'))
        font_name = 'Arial_TR'
    except:
        font_name = 'Helvetica'
        
    doc = SimpleDocTemplate(
        dosya_adi,
        pagesize=landscape(A4),
        rightMargin=15*mm,
        leftMargin=15*mm,
        topMargin=12.5*mm,
        bottomMargin=12.5*mm
    )
    
    style_normal = ParagraphStyle('Normal_TR', fontName=font_name, fontSize=8, leading=10)
    style_center = ParagraphStyle('Center_TR', fontName=font_name, fontSize=8, alignment=1, leading=10)
    style_right = ParagraphStyle('Right_TR', fontName=font_name, fontSize=8, alignment=2, leading=10)
    style_title = ParagraphStyle('Title_TR', fontName=font_name, fontSize=11, alignment=1, leading=13)
    
    elements = []
    
    elements.append(Paragraph("YAKLAŞIK MALİYET HESAP CETVELİ", style_title))
    elements.append(Spacer(1, 4*mm))
    
    tarih = ""
    if "belge_tarihi" in veri:
        date_str = veri["belge_tarihi"]
        if "T" in date_str:
            d = datetime.datetime.strptime(date_str.split("T")[0], "%Y-%m-%d")
        elif "-" in date_str:
            d = datetime.datetime.strptime(date_str, "%Y-%m-%d")
        else:
            try: d = datetime.datetime.strptime(date_str, "%d.%m.%Y")
            except: d = datetime.datetime.now()
        tarih = d.strftime("%d.%m.%Y")
        
    konu = get_ihale_konu(veri)
    
    baslik = veri.get("resmi_baslik", "")
    lines = [L.strip() for L in baslik.split("\n") if L.strip()]
    if len(lines) >= 3:
        idare_adi = lines[2]
    elif lines:
        idare_adi = lines[-1]
    else:
        idare_adi = "Gazi Mustafa Kemal Anadolu Lisesi Müdürlüğü"
    
    table_data = []
    
    row0 = ["İdarenin Adı", "", "", "", "", idare_adi, "", "", "", "", "", "", "", Paragraph("İdarece Tesbit<br/>ve Takdir Edilen<br/>Yaklaşık Maliyet<br/>( KDV Hariç )", style_center), ""]
    table_data.append([Paragraph(x, style_normal) if isinstance(x, str) and x else x for x in row0])
    
    row1 = ["Yaklaşık Maliyeti Yapılan İş / Mal / Hizmetin Adı, Niteliği", "", "", "", "", konu, "", "", "", "", "", "", "", "", ""]
    table_data.append([Paragraph(x, style_normal) if isinstance(x, str) and x else x for x in row1])
    
    row2 = ["Düzenleme Tarihi", "", "", tarih, "", Paragraph("Kişiler / Kurumlar / Firmalar ve Yaklaşık Maliyet İçin<br/>Bildirilen Fiyatlar ( KDV Hariç )", style_center), "", "", "", "", "", "", "", "", ""]
    table_data.append([Paragraph(x, style_normal) if isinstance(x, str) and x else x for x in row2])
    
    firmalar = veri.get("firmalar", ["", "", "", ""])
    firma_vergiler = veri.get("firma_vergiler", ["", "", "", ""])
    while len(firmalar) < 4: firmalar.append("")
    while len(firma_vergiler) < 4: firma_vergiler.append("")
    
    f_headers = []
    for i in range(4):
        f = firmalar[i].strip()
        v = firma_vergiler[i].strip()
        f_text = f"{f}" if f else ""
        if v:
            if len(v) == 10:
                f_text += f"<br/>(Vergi No: {v})"
            elif len(v) == 11:
                f_text += f"<br/>(T.C. No: {v})"
            else:
                f_text += f"<br/>({v})"
        f_headers.append(Paragraph(f_text, style_center))
        
    row3 = [Paragraph("Sıra<br/>No", style_center), Paragraph("Yaklaşık Maliyeti Hesaplanacak Mal ve Hizmetin", style_center), "", "", "", f_headers[0], "", f_headers[1], "", f_headers[2], "", f_headers[3], "", Paragraph("Birim Yaklaşık<br/>Maliyet<br/>Fiyatı", style_center), Paragraph("Toplam<br/>Yaklaşık<br/>Maliyet<br/>Fiyatı", style_center)]
    table_data.append(row3)
    
    row4 = ["", Paragraph("Cinsi", style_center), Paragraph("Özelliği", style_center), Paragraph("Miktarı", style_center), Paragraph("Ölçüsü", style_center)]
    for i in range(4):
        row4.extend([Paragraph("Birim<br/>Fiyat", style_center), Paragraph("Toplam<br/>Fiyat", style_center)])
    row4.extend(["", ""])
    table_data.append(row4)
    
    kalemler = veri.get("kalemler", [])
    gecerli_kalemler = [k for k in kalemler if k.get("cins", "").strip()]
    if not gecerli_kalemler:
        gecerli_kalemler = [{}]
        
    col_sums = [0.0, 0.0, 0.0, 0.0, 0.0]
    
    for idx, k in enumerate(gecerli_kalemler):
        val = str(k.get("miktar", ""))
        miktar_float = 0.0
        try:
            miktar_float = float(val)
            if miktar_float == int(miktar_float): val = str(int(miktar_float))
        except: pass
        
        row = [
            Paragraph(str(idx+1), style_center),
            Paragraph(k.get("cins", ""), style_normal),
            Paragraph(k.get("ozellik", ""), style_normal),
            Paragraph(val, style_center),
            Paragraph(k.get("birim", ""), style_center)
        ]
        
        fiyatlar = k.get("fiyatlar", [])
        gecerli_fiyatlar = []
        
        for fidx in range(4):
            bf = ""
            tf = ""
            if fidx < len(fiyatlar) and fiyatlar[fidx] not in ["", None]:
                try:
                    fiyat_val = float(fiyatlar[fidx])
                    if fiyat_val > 0:
                        bf = f"{fiyat_val:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
                        tf_float = fiyat_val * miktar_float
                        tf = f"{tf_float:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
                        gecerli_fiyatlar.append(fiyat_val)
                        col_sums[fidx] += tf_float
                except:
                    pass
            row.extend([Paragraph(bf, style_right), Paragraph(tf, style_right)])
            
        yak_maliyet_birim = ""
        yak_maliyet_toplam = ""
        if gecerli_fiyatlar:
            ymb = sum(gecerli_fiyatlar) / len(gecerli_fiyatlar)
            ymt = ymb * miktar_float
            yak_maliyet_birim = f"{ymb:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
            yak_maliyet_toplam = f"{ymt:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
            col_sums[4] += ymt
            
        row.extend([Paragraph(yak_maliyet_birim, style_right), Paragraph(yak_maliyet_toplam, style_right)])
        table_data.append(row)
        
    t_row = [
        "", Paragraph("Kişiler / Firmalarca Bildirilen Yaklaşık Maliyet Fiyatları Toplamı ( KDV HARİÇ )", style_normal), "", "", "", 
        Paragraph("TOPLAM", style_center), Paragraph(f"{col_sums[0]:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".") if col_sums[0]>0 else "", style_center), 
        Paragraph("TOPLAM", style_center), Paragraph(f"{col_sums[1]:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".") if col_sums[1]>0 else "", style_center), 
        Paragraph("TOPLAM", style_center), Paragraph(f"{col_sums[2]:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".") if col_sums[2]>0 else "", style_center), 
        Paragraph("TOPLAM", style_center), Paragraph(f"{col_sums[3]:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".") if col_sums[3]>0 else "", style_center), 
        Paragraph("TOPLAM", style_center), Paragraph(f"{col_sums[4]:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".") if col_sums[4]>0 else "", style_center)
    ]
    table_data.append(t_row)
        
    col_widths = [8*mm, 25*mm, 30*mm, 12*mm, 15*mm, 17*mm, 18*mm, 17*mm, 18*mm, 17*mm, 18*mm, 17*mm, 18*mm, 20*mm, 20*mm]
    t = Table(table_data, colWidths=col_widths, repeatRows=5)
    
    t_style = [
        ('FONTNAME', (0,0), (-1,-1), font_name),
        ('FONTSIZE', (0,0), (-1,-1), 8),
        ('ALIGN', (0,0), (-1,-1), 'CENTER'),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('GRID', (0,0), (-1,-1), 0.25, colors.black),
        
        # Row 0 merges
        ('SPAN', (0,0), (4,0)),
        ('SPAN', (5,0), (12,0)),
        ('SPAN', (13,0), (14,2)), # Idarece Tesbit spans rows 0,1,2 and cols 13,14
        
        # Row 1 merges
        ('SPAN', (0,1), (4,1)),
        ('SPAN', (5,1), (12,1)),
        
        # Row 2 merges
        ('SPAN', (0,2), (2,2)),
        ('SPAN', (3,2), (4,2)),
        ('SPAN', (5,2), (12,2)),
        
        # Row 3 merges
        ('SPAN', (0,3), (0,4)), # Sira No
        ('ROTATION', (0,3), (0,4), 90),
        ('SPAN', (1,3), (4,3)), # Mal ve Hizmet
        ('SPAN', (5,3), (6,3)), # Firm 1
        ('SPAN', (7,3), (8,3)), # Firm 2
        ('SPAN', (9,3), (10,3)), # Firm 3
        ('SPAN', (11,3), (12,3)), # Firm 4
        ('SPAN', (13,3), (13,4)), # Birim Yaklasik
        ('SPAN', (14,3), (14,4)), # Toplam Yaklasik
    ]
    
    # Dashed lines
    data_row_start = 4
    data_row_end = len(table_data) - 2 # Before total row
    if data_row_end >= data_row_start:
        for r in range(data_row_start, data_row_end + 1):
            t_style.append(('LINEBEFORE', (6, r), (6, r), 0.5, colors.black, None, (2, 2)))
            t_style.append(('LINEBEFORE', (8, r), (8, r), 0.5, colors.black, None, (2, 2)))
            t_style.append(('LINEBEFORE', (10, r), (10, r), 0.5, colors.black, None, (2, 2)))
            t_style.append(('LINEBEFORE', (12, r), (12, r), 0.5, colors.black, None, (2, 2)))
    
    # Total row merges
    last_row = len(table_data) - 1
    t_style.extend([
        ('SPAN', (1, last_row), (4, last_row)),
    ])
    
    t.setStyle(TableStyle(t_style))
    elements.append(t)
    
    elements.append(Spacer(1, 4*mm))
    
    p1 = "&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;İdaremizce ihtiyaç duyulan ve satın alınması düşünülen aşağıda cinsi, özellikleri ve miktarları yazılı malların / hizmetlerin 4734 Sayılı Kamu İhale Kanunu'nun 9'uncu Maddesi gereğince yaklaşık maliyetinin tesbitine esas olmak üzere; ilgili kişi kurum ve firmalardan yaklaşık maliyetinin tesbitine esas olmak üzere, her türlü fiyat araştırması yapılmıştır. Araştırma sonuçları yukarıda tabloda gösterilmiştir."
    p2 = "&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;Yukarıda açıklandığı üzere, ihaleye çıkılması düşünülen mal / hizmetlerin fiyat araştırması neticesinde; 4734 Sayılı İhale Kanunu'nun 9. Maddesi gereğince yaklaşık maliyetinin KDV hariç yukarıda belirtildiği gibi takdir ve tesbit edilerek iş bu Hesap Cetveli tarafımca / tarafımızca düzenlenerek imza altına alınmıştır."
    
    elements.append(Paragraph(p1, style_normal))
    elements.append(Spacer(1, 2*mm))
    elements.append(Paragraph(p2, style_normal))
    elements.append(Spacer(1, 5*mm))
    
    elements.append(Paragraph("YAKLAŞIK MALİYET TESPİTİ YAPAN GÖREVLİLER", style_center))
    elements.append(Spacer(1, 5*mm))
    
    def get_komisyon_local(veri):
        secimler = veri.get("komisyon", {})
        ids = ["ihale_kom_yaklasik_1", "ihale_kom_yaklasik_2", "ihale_kom_yaklasik_3"]
        uyeler = []
        for i in ids:
            u = secimler.get(i)
            if u: uyeler.append(u)
        return uyeler
        
    komisyon_uyeleri = get_komisyon_local(veri)
    if len(komisyon_uyeleri) > 0:
        imza_data = [[]]
        for idx, uye in enumerate(komisyon_uyeleri):
            title = "Başkan" if idx == 0 else "Üye"
            imza_data[0].append(Paragraph(f"{uye}<br/>{title}", style_center))
        
        col_w = 260*mm / len(komisyon_uyeleri)
        sig_table = Table(imza_data, colWidths=[col_w] * len(komisyon_uyeleri))
        sig_table.setStyle(TableStyle([
            ('ALIGN', (0,0), (-1,-1), 'CENTER'),
            ('VALIGN', (0,0), (-1,-1), 'TOP'),
        ]))
        elements.append(sig_table)
    
    elements.append(Spacer(1, 5*mm))
    elements.append(Paragraph("Not:Yaklaşık Maliyet Hesap Cetveli ve Ekleri İhale Onay Belgesi'ne Eklenecektir.", style_normal))
    
    doc.build(elements)
    return dosya_adi


def uret_yaklasik_maliyet_excel(veri, hedef_klasor):
    timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    dosya_adi = os.path.join(hedef_klasor, f"Yaklasik_Maliyet_{timestamp}.xlsx")
    
    wb = Workbook()
    ws = wb.active
    ws.title = "Yaklasik Maliyet"
    
    ws.page_margins.top = 0.492
    ws.page_margins.bottom = 0.492
    ws.page_margins.left = 0.984
    ws.page_margins.right = 0.984
    
    normal_font = Font(name='Arial', size=9)
    bold_title = Font(name='Arial', size=11, bold=False)
    
    center_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
    left_align = Alignment(horizontal='left', vertical='center', wrap_text=True)
    right_align = Alignment(horizontal='right', vertical='center', wrap_text=True)
    thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
    
    ws.merge_cells('A1:O1')
    c1 = ws.cell(row=1, column=1, value="Y A K L A ┼Ş I K   M A L ─░ Y E T   H E S A P   C E T V E L ─░")
    c1.font = bold_title
    c1.alignment = center_align
    
    tarih = format_date(veri.get("belge_tarihi", ""))
    konu = veri.get("konu", "K─▒rtasiye Al─▒m─▒")
    
    ws.cell(row=3, column=1, value="─░darenin Ad─▒").font = normal_font
    ws.cell(row=3, column=2, value=": Gazi Mustafa Kemal Anadolu Lisesi M├╝d├╝rl├╝─ş├╝").font = normal_font
    ws.merge_cells('B3:E3')
    ws.cell(row=3, column=14, value="─░darece Tespit\nve Takdir Edilen\nYakla┼ş─▒k Maliyet\n(KDV Hari├ğ)").font = normal_font
    ws.cell(row=3, column=14).alignment = center_align
    ws.merge_cells('N3:O5')
    
    ws.cell(row=4, column=1, value="Yakla┼ş─▒k Maliyeti Yap─▒lan ─░┼ş / Mal / Hizmetin Ad─▒, Niteli─şi").font = normal_font
    ws.cell(row=4, column=2, value=f": {konu}").font = normal_font
    ws.merge_cells('B4:E4')
    
    ws.cell(row=5, column=1, value="D├╝zenleme Tarihi").font = normal_font
    ws.cell(row=5, column=2, value=f": {tarih}").font = normal_font
    ws.merge_cells('B5:E5')
    
    ws.merge_cells('F5:M5')
    ws.cell(row=5, column=6, value="Ki┼şiler / Kurumlar / Firmalar ve Yakla┼ş─▒k Maliyet ─░├ğin\nBildirilen Fiyatlar ( KDV Hari├ğ )").font = normal_font
    ws.cell(row=5, column=6).alignment = center_align
    
    # Headers
    r = 7
    firmalar = veri.get("firmalar", [])
    f1 = firmalar[0].strip() if len(firmalar) > 0 and firmalar[0] else ""
    f2 = firmalar[1].strip() if len(firmalar) > 1 and firmalar[1] else ""
    f3 = firmalar[2].strip() if len(firmalar) > 2 and firmalar[2] else ""
    f4 = firmalar[3].strip() if len(firmalar) > 3 and firmalar[3] else ""
    
    h1 = ["S─▒ra\nNo", "Yakla┼ş─▒k Maliyeti Hesaplanacak Mal ve Hizmetin", "", "", "", f"1\n{f1}", "", f"2\n{f2}", "", f"3\n{f3}", "", f"4\n{f4}", "", "Birim Yakla┼ş─▒k\nMaliyet Fiyat─▒", "Toplam Yakla┼ş─▒k\nMaliyet Fiyat─▒"]
    ws.row_dimensions[r].height = 30
    for i, h in enumerate(h1, 1):
        c = ws.cell(row=r, column=i, value=h)
        c.font = normal_font
        c.alignment = center_align
        c.border = thin_border
        
    ws.merge_cells(start_row=r, start_column=1, end_row=r+1, end_column=1)
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
    ws.merge_cells(start_row=r, start_column=6, end_row=r, end_column=7)
    ws.merge_cells(start_row=r, start_column=8, end_row=r, end_column=9)
    ws.merge_cells(start_row=r, start_column=10, end_row=r, end_column=11)
    ws.merge_cells(start_row=r, start_column=12, end_row=r, end_column=13)
    ws.merge_cells(start_row=r, start_column=14, end_row=r+1, end_column=14)
    ws.merge_cells(start_row=r, start_column=15, end_row=r+1, end_column=15)
    
    r += 1
    h2 = ["", "Cinsi", "├ûzelli─şi", "Miktar─▒", "├ûl├ğ├╝s├╝", "Birim Fiyat", "Toplam Fiyat", "Birim Fiyat", "Toplam Fiyat", "Birim Fiyat", "Toplam Fiyat", "Birim Fiyat", "Toplam Fiyat", "", ""]
    for i, h in enumerate(h2, 1):
        if type(ws.cell(row=r, column=i)).__name__ == 'MergedCell':
            continue
        c = ws.cell(row=r, column=i, value=h)
        c.font = normal_font
        c.alignment = center_align
        c.border = thin_border
        
    kalemler = veri.get("kalemler", [])
    gecerli_kalemler = [k for k in kalemler if k.get("cins", "").strip()]
    if not gecerli_kalemler: gecerli_kalemler = [{}]
    
    r += 1
    col_sums = [0.0, 0.0, 0.0, 0.0, 0.0]
    
    for idx, k in enumerate(gecerli_kalemler):
        val = str(k.get("miktar", ""))
        miktar_float = 0.0
        try:
            miktar_float = float(val)
            if miktar_float == int(miktar_float): val = str(int(miktar_float))
        except: pass
        
        row_data = [
            str(idx+1), k.get("cins", ""), k.get("ozellik", ""), val, k.get("birim", "")
        ]
        
        fiyatlar = k.get("fiyatlar", [])
        gecerli_fiyatlar = []
        
        for fidx in range(4):
            bf = ""
            tf = ""
            if fidx < len(fiyatlar) and fiyatlar[fidx] not in ["", None]:
                try:
                    fiyat_val = float(fiyatlar[fidx])
                    if fiyat_val > 0:
                        bf = fiyat_val
                        tf = bf * miktar_float
                        gecerli_fiyatlar.append(fiyat_val)
                        col_sums[fidx] += tf
                except:
                    pass
            row_data.extend([bf, tf])
            
        yak_maliyet_birim = ""
        yak_maliyet_toplam = ""
        if gecerli_fiyatlar:
            yak_maliyet_birim = sum(gecerli_fiyatlar) / len(gecerli_fiyatlar)
            yak_maliyet_toplam = yak_maliyet_birim * miktar_float
            col_sums[4] += yak_maliyet_toplam
            
        row_data.extend([yak_maliyet_birim, yak_maliyet_toplam])

        for col_num, cell_val in enumerate(row_data, 1):
            if type(ws.cell(row=r, column=col_num)).__name__ == 'MergedCell':
                continue
            c = ws.cell(row=r, column=col_num, value=cell_val)
            c.font = normal_font
            c.alignment = center_align
            c.border = thin_border
            if isinstance(cell_val, float):
                c.number_format = '#,##0.00'
        r += 1
        
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
    ws.merge_cells(start_row=r, start_column=6, end_row=r, end_column=7)
    ws.merge_cells(start_row=r, start_column=8, end_row=r, end_column=9)
    ws.merge_cells(start_row=r, start_column=10, end_row=r, end_column=11)
    ws.merge_cells(start_row=r, start_column=12, end_row=r, end_column=13)
    
    t_row = [
        "", "Ki┼şiler / Firmalarca Bildirilen Yakla┼ş─▒k Maliyet Fiyatlar─▒ Toplam─▒ ( KDV HAR─░├ç )", "", "", "", 
        col_sums[0] if col_sums[0]>0 else "TOPLAM", "", 
        col_sums[1] if col_sums[1]>0 else "TOPLAM", "", 
        col_sums[2] if col_sums[2]>0 else "TOPLAM", "", 
        col_sums[3] if col_sums[3]>0 else "TOPLAM", "", 
        "", col_sums[4] if col_sums[4]>0 else "TOPLAM"
    ]
    
    for i, val in enumerate(t_row, 1):
        if type(ws.cell(row=r, column=i)).__name__ == 'MergedCell':
            continue
        c = ws.cell(row=r, column=i, value=val)
        c.font = normal_font
        c.border = thin_border
        c.alignment = right_align if i==2 else center_align
        
    r += 2
    p1 = "           ─░daremizce ihtiya├ğ duyulan ve sat─▒n al─▒nmas─▒ d├╝┼ş├╝n├╝len a┼şa─ş─▒da cinsi, ├Âzellikleri ve miktarlar─▒ yaz─▒l─▒ mallar─▒n / hizmetlerin 4734 Say─▒l─▒ Kamu ─░hale Kanunu'nun 9'uncu Maddesi gere─şince yakla┼ş─▒k maliyetinin tesbitine esas olmak ├╝zere; ilgili ki┼şi kurum ve firmalardan yakla┼ş─▒k maliyetinin tesbitine esas olmak ├╝zere, her t├╝rl├╝ fiyat ara┼şt─▒rmas─▒ yap─▒lm─▒┼şt─▒r. Ara┼şt─▒rma sonu├ğlar─▒ yukar─▒da tabloda g├Âsterilmi┼ştir."
    ws.merge_cells(start_row=r, start_column=1, end_row=r+1, end_column=15)
    c_p1 = ws.cell(row=r, column=1, value=p1)
    c_p1.font = normal_font
    c_p1.alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)
    
    r += 2
    p2 = "          Yukar─▒da a├ğ─▒kland─▒─ş─▒ ├╝zere, ihaleye ├ğ─▒k─▒lmas─▒ d├╝┼ş├╝n├╝len mal / hizmetlerin fiyat ara┼şt─▒rmas─▒ neticesinde; 4734 Say─▒l─▒ ─░hale Kanunu'nun 9. Maddesi gere─şince yakla┼ş─▒k maliyetinin KDV hari├ğ yukar─▒da belirtildi─şi gibi takdir ve tesbit edilerek i┼ş bu Hesap Cetveli taraf─▒mca / taraf─▒m─▒zca d├╝zenlenerek imza alt─▒na al─▒nm─▒┼şt─▒r."
    ws.merge_cells(start_row=r, start_column=1, end_row=r+1, end_column=15)
    c_p2 = ws.cell(row=r, column=1, value=p2)
    c_p2.font = normal_font
    c_p2.alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)
    
    r += 3
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=15)
    c_s = ws.cell(row=r, column=1, value="Y A K L A ┼Ş I K   M A L ─░ Y E T ─░   Y A P A N   G ├û R E V L ─░ / G ├û R E V L ─░ L E R")
    c_s.font = normal_font
    c_s.alignment = center_align
    
    r += 2
    komisyon_uyeleri = get_komisyon(veri)
    if len(komisyon_uyeleri) > 0:
        if len(komisyon_uyeleri) == 1: cols = [8]
        elif len(komisyon_uyeleri) == 2: cols = [4, 12]
        elif len(komisyon_uyeleri) == 3: cols = [2, 8, 14]
        else: cols = [2, 6, 10, 14]
            
        for idx, uye in enumerate(komisyon_uyeleri[:4]):
            c_isim = ws.cell(row=r, column=cols[idx], value=uye)
            c_isim.font = normal_font
            c_isim.alignment = center_align
            
            title = "Ba┼şkan" if idx == 0 else "├£ye"
            c_unvan = ws.cell(row=r+1, column=cols[idx], value=title)
            c_unvan.font = normal_font
            c_unvan.alignment = center_align
            
    r += 4
    ws.cell(row=r, column=1, value="Not:Yakla┼ş─▒k Maliyet Hesap Cetveli ve Ekleri ─░hale Onay Belgesi'ne Eklenecektir.").font = normal_font
    
    cols_width = [5, 15, 15, 10, 10, 10, 10, 10, 10, 10, 10, 10, 10, 15, 15]
    for i, w in enumerate(cols_width, 1):
        ws.column_dimensions[chr(64+i)].width = w
        
    wb.save(dosya_adi)
    return dosya_adi


def uret_ozel_fiyat_isteme_pdf(veri, hedef_klasor):
    import datetime
    import os
    from reportlab.pdfgen import canvas
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib import colors
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    
    timestamp = datetime.datetime.now().strftime('%Y%m%d_%H%M%S')
    dosya_adi = os.path.join(hedef_klasor, f"Fiyat_Isteme_{timestamp}.pdf")
    
    try:
        pdfmetrics.registerFont(TTFont('Arial_TR', r'C:\Windows\Fonts\arial.ttf'))
        font_name = 'Arial_TR'
    except:
        font_name = 'Helvetica'
        
    doc = SimpleDocTemplate(
        dosya_adi,
        pagesize=A4,
        rightMargin=15*mm,
        leftMargin=15*mm,
        topMargin=15*mm,
        bottomMargin=15*mm
    )
    
    style_normal = ParagraphStyle('Normal_TR', fontName=font_name, fontSize=9, leading=11)
    style_justify = ParagraphStyle('Justify_TR', fontName=font_name, fontSize=9, alignment=4, leading=11, firstLineIndent=12*mm)
    style_center = ParagraphStyle('Center_TR', fontName=font_name, fontSize=9, alignment=1, leading=11)
    style_right = ParagraphStyle('Right_TR', fontName=font_name, fontSize=9, alignment=2, leading=11)
    
    style_normal_10 = ParagraphStyle('Normal10_TR', fontName=font_name, fontSize=10, leading=12)
    style_right_10 = ParagraphStyle('Right10_TR', fontName=font_name, fontSize=10, alignment=2, leading=12)
    style_center_10 = ParagraphStyle('Center10_TR', fontName=font_name, fontSize=10, alignment=1, leading=12)
    
    kalemler = veri.get("kalemler", [])
    gecerli_kalemler = [k for k in kalemler if k.get("cins", "").strip()]
    if not gecerli_kalemler:
        gecerli_kalemler = [{}] 
        
    komisyon_uyeleri = get_komisyon(veri)
        
    elements = []
    
    def tr_capitalize(text):
        tr_lower = {'I': 'ı', 'İ': 'i'}
        tr_upper = {'ı': 'I', 'i': 'İ'}
        
        words = text.split()
        res = []
        for w in words:
            if not w: continue
            first = w[0]
            rest = w[1:]
            
            first = tr_upper.get(first, first.upper())
            rest_lower = ""
            for c in rest:
                rest_lower += tr_lower.get(c, c.lower())
                
            res.append(first + rest_lower)
        return " ".join(res)
    
    baslik = veri.get("resmi_baslik", "").replace("\\n", "\n")
    baslik_satirlari = [s.strip() for s in baslik.split("\n") if s.strip()]
    for i, satir in enumerate(baslik_satirlari):
        if i == 2:
            satir = tr_capitalize(satir)
        elements.append(Paragraph(satir.strip(), style_center_10))
    
    elements.append(Spacer(1, 8*mm))
    
    tarih = format_date(veri.get("belge_tarihi", ""))
    sayi = veri.get("yazisma_kodu", "")
    
    tbl_sayi_konu = Table([
        [Paragraph("Sayı", style_normal_10), Paragraph(":", style_normal_10), Paragraph(sayi, style_normal_10), Paragraph(tarih, style_right_10)],
        [Paragraph("Konu", style_normal_10), Paragraph(":", style_normal_10), Paragraph("Teklifiniz", style_normal_10), ""]
    ], colWidths=[12*mm, 3*mm, 95*mm, 70*mm])
    
    tbl_sayi_konu.setStyle(TableStyle([
        ('ALIGN', (0,0), (0,-1), 'LEFT'),
        ('ALIGN', (1,0), (1,-1), 'LEFT'),
        ('ALIGN', (2,0), (2,-1), 'LEFT'),
        ('ALIGN', (3,0), (3,0), 'RIGHT'),
        ('VALIGN', (0,0), (-1,-1), 'TOP'),
        ('BOTTOMPADDING', (0,0), (-1,-1), 1),
        ('TOPPADDING', (0,0), (-1,-1), 1),
        ('LEFTPADDING', (0,0), (-1,-1), 0),
        ('RIGHTPADDING', (0,0), (-1,-1), 0),
    ]))
    elements.append(tbl_sayi_konu)
    
    elements.append(Spacer(1, 8*mm))
    
    hitap = "Sayın Yetkili"
    elements.append(Paragraph(hitap, style_center_10))
    elements.append(Spacer(1, 4*mm))
    metin = "Aşağıda cinsi, özellikleri ve miktarları yazılı mallar / hizmetler 4734 sayılı Kamu İhale Kanunu'nun 22/d Maddesi gereğince Doğrudan Temin Usulüyle satın alınacaktır. İlgilenmeniz halinde KDV hariç teklifinizin bildirilmesini rica ederim / ederiz."
    elements.append(Paragraph(metin, style_justify))
    elements.append(Spacer(1, 8*mm))
    
    if len(komisyon_uyeleri) > 0:
        imza_data = [[]]
        for uidx, uye in enumerate(komisyon_uyeleri):
            title = "Başkan" if uidx == 0 else "Üye"
            imza_data[0].append(Paragraph(f"{uye}<br/>{title}", style_center))
        
        col_w = 180*mm / len(komisyon_uyeleri)
        sig_table = Table(imza_data, colWidths=[col_w] * len(komisyon_uyeleri))
        sig_table.setStyle(TableStyle([
            ('ALIGN', (0,0), (-1,-1), 'CENTER'),
            ('VALIGN', (0,0), (-1,-1), 'TOP'),
        ]))
        elements.append(sig_table)
    
    elements.append(Spacer(1, 8*mm))
    
    # 5. Tablo (Font 8, Yeni Kolonlar)
    table_data = []
    # Super Header
    super_header_style = ParagraphStyle('SuperHeader_TR', fontName=font_name, fontSize=10, alignment=1, leading=12)
    table_data.append([Paragraph("Satın Alınacak Malın", super_header_style), "", "", "", "", Paragraph("Teklif Edilen KDV Hariç", super_header_style), ""])
    
    # Headers
    headers = ["S.No", "Cinsi", "Özellikleri", "Ölçüsü", "Miktarı", "Birim Fiyatı\n(TL)", "Toplam Fiyatı\n(TL)"]
    table_data.append([Paragraph(h, ParagraphStyle('TH_TR', fontName=font_name, fontSize=8, alignment=0, leading=10)) for h in headers])
    
    style_cell = ParagraphStyle('Cell_TR', fontName=font_name, fontSize=8, leading=10, alignment=1)
    style_cell_left = ParagraphStyle('Cell_Left_TR', fontName=font_name, fontSize=8, leading=10, alignment=0)
    for i_k, k in enumerate(gecerli_kalemler):
        val = str(k.get("miktar", ""))
        try:
            if float(val) == int(float(val)): val = str(int(float(val)))
        except: pass
        
        row_data = [
            Paragraph(str(i_k+1), style_cell),
            Paragraph(k.get("cins", ""), style_cell_left),
            Paragraph(k.get("ozellik", ""), style_cell_left),
            Paragraph(k.get("birim", ""), style_cell),
            Paragraph(val, style_cell),
            "", ""
        ]
        table_data.append(row_data)
        
    start_bottom = len(table_data)
    table_data.append([Paragraph("KDV Hariç Teklif Edilen Toplam Fiyat:", ParagraphStyle('TR_R', fontName=font_name, fontSize=8, alignment=2, leading=10)), "", "", "", "", "", ""])
    
    col_widths = [10*mm, 35*mm, 45*mm, 15*mm, 15*mm, 30*mm, 30*mm]
    t = Table(table_data, colWidths=col_widths, repeatRows=2)
    
    table_style = [
        ('FONTNAME', (0,0), (-1,-1), font_name),
        ('FONTSIZE', (0,0), (-1,-1), 8),
        ('ALIGN', (0,0), (-1,-1), 'LEFT'),
        ('ALIGN', (0,0), (-1,1), 'CENTER'),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('LEFTPADDING', (0,0), (-1,-1), 2),
        ('RIGHTPADDING', (0,0), (-1,-1), 2),
        ('INNERGRID', (0,0), (-1, start_bottom), 0.25, colors.black),
        ('BOX', (0,0), (-1, start_bottom), 0.25, colors.black),
        
        # Super Header span
        ('SPAN', (0,0), (4,0)),
        ('SPAN', (5,0), (6,0)),
        
        # KDV Hariç Toplam
        ('SPAN', (0, start_bottom), (5, start_bottom)),
        ('ALIGN', (0, start_bottom), (5, start_bottom), 'RIGHT'),
        ('BOX', (0, start_bottom), (-1, start_bottom), 0.25, colors.black),
        ('INNERGRID', (0, start_bottom), (-1, start_bottom), 0.25, colors.black),
        ('BOTTOMPADDING', (0,0), (-1,-1), 1),
        ('TOPPADDING', (0,0), (-1,-1), 1),
        ('BOTTOMPADDING', (0, start_bottom), (-1, start_bottom), 8),
        ('TOPPADDING', (0, start_bottom), (-1, start_bottom), 8),
    ]
    t.setStyle(TableStyle(table_style))
    elements.append(t)
    
    # Sartlar Table (Font 8)
    sartlar_data = [
        [Paragraph("DİĞER ŞARTLAR", ParagraphStyle('S1', fontName=font_name, fontSize=8, alignment=1)), ""],
        [Paragraph("1- TESLİM SÜRESİ", ParagraphStyle('S2', fontName=font_name, fontSize=8)), Paragraph("1 gün", ParagraphStyle('S3', fontName=font_name, fontSize=8))],
        [Paragraph("2- TESLİM EDİLECEK PARTİ MİKTARI", ParagraphStyle('S2', fontName=font_name, fontSize=8)), Paragraph("1", ParagraphStyle('S3', fontName=font_name, fontSize=8))],
        [Paragraph("3- NAKLİYE VE SİGORTANIN KİME AİT OLDUĞU", ParagraphStyle('S2', fontName=font_name, fontSize=8)), Paragraph("Satıcıya", ParagraphStyle('S3', fontName=font_name, fontSize=8))],
        [Paragraph("4- DİĞER ÖZEL ŞARTLAR", ParagraphStyle('S2', fontName=font_name, fontSize=8)), Paragraph("YOK", ParagraphStyle('S3', fontName=font_name, fontSize=8))],
        [Paragraph("5- UYULMASI GEREKEN STANDARTLAR", ParagraphStyle('S2', fontName=font_name, fontSize=8)), Paragraph("TSE", ParagraphStyle('S3', fontName=font_name, fontSize=8))],
        [Paragraph("6- TEKNİK ŞARTNAME", ParagraphStyle('S2', fontName=font_name, fontSize=8)), Paragraph("YOK", ParagraphStyle('S3', fontName=font_name, fontSize=8))],
        [Paragraph("7- DİĞER HUSUSLAR", ParagraphStyle('S2', fontName=font_name, fontSize=8)), Paragraph("YOK", ParagraphStyle('S3', fontName=font_name, fontSize=8))]
    ]
    sartlar_table = Table(sartlar_data, colWidths=[120*mm, 60*mm])
    sartlar_table.setStyle(TableStyle([
        ('ALIGN', (0,0), (-1,-1), 'LEFT'),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('INNERGRID', (0,0), (-1,-1), 0.25, colors.black),
        ('BOX', (0,0), (-1,-1), 0.25, colors.black),
        ('SPAN', (0,0), (1,0)),
        ('BOTTOMPADDING', (0,0), (-1,-1), 2),
        ('TOPPADDING', (0,0), (-1,-1), 2),
    ]))
    
    elements.append(sartlar_table)
    
    elements.append(Spacer(1, 4*mm))
    
    bottom_text = "Yukarıda belirtilen ve İdarenizce satın alınacak mal veya hizmetlerin cinsi, özellikleri, miktarı ve diğer şartları okudum, anladım. KDV hariç toplam .................................................. TL bedel ile belirtilen mal/hizmet satışını gerçekleştirmeyi kabul ve taahhüt ediyorum."
    elements.append(Paragraph(bottom_text, style_justify))
    
    elements.append(Spacer(1, 6*mm))
    
    sig_data = [
        ["", Paragraph(f"Teklif Eden<br/>..../..../........<br/>Adı, Soyadı--Ticaret Ünvanı--İmza--Mühür", style_center)]
    ]
    sig_tbl = Table(sig_data, colWidths=[100*mm, 80*mm])
    elements.append(sig_tbl)
    
    elements.append(Spacer(1, 10*mm))
    elements.append(Paragraph("NOT :-Bu Belge Piyasa Fiyat Araştırması Tutanağına Eklenecektir.", style_normal))
        
    doc.build(elements)
    return dosya_adi

def uret_ozel_fiyat_isteme_excel(veri, hedef_klasor):
    import datetime
    import os
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, Border, Side
    
    timestamp = datetime.datetime.now().strftime('%Y%m%d_%H%M%S')
    dosya_adi = os.path.join(hedef_klasor, f"Fiyat_Isteme_{timestamp}.xlsx")
    
    wb = Workbook()
    ws = wb.active
    ws.title = "Fiyat Isteme"
    
    normal_font = Font(name='Arial', size=9)
    bold_font = Font(name='Arial', size=9, bold=True)
    normal_font_10 = Font(name='Arial', size=10)
    bold_font_10 = Font(name='Arial', size=10, bold=True)
    
    # Yeni fontlar: Tablo icin 8
    normal_font_8 = Font(name='Arial', size=8)
    bold_font_8 = Font(name='Arial', size=8, bold=True)
    
    center_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
    left_align = Alignment(horizontal='left', vertical='center', wrap_text=True)
    right_align = Alignment(horizontal='right', vertical='center', wrap_text=True)
    thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
    
    def tr_capitalize(text):
        tr_lower = {'I': 'ı', 'İ': 'i'}
        tr_upper = {'ı': 'I', 'i': 'İ'}
        
        words = text.split()
        res = []
        for w in words:
            if not w: continue
            first = w[0]
            rest = w[1:]
            
            first = tr_upper.get(first, first.upper())
            rest_lower = ""
            for c in rest:
                rest_lower += tr_lower.get(c, c.lower())
                
            res.append(first + rest_lower)
        return " ".join(res)
    
    baslik = veri.get("resmi_baslik", "").replace("\\n", "\n")
    baslik_satirlari = [s.strip() for s in baslik.split("\n") if s.strip()]
    
    r = 1
    for i, satir in enumerate(baslik_satirlari):
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=7)
        if i == 2:
            satir = tr_capitalize(satir)
        c = ws.cell(row=r, column=1, value=satir)
        c.font = bold_font_10
        c.alignment = center_align
        r += 1
        
    r += 1
    tarih = format_date(veri.get("belge_tarihi", ""))
    sayi = veri.get("yazisma_kodu", "")
    
    ws.cell(row=r, column=1, value="Sayı").font = normal_font_10
    ws.cell(row=r, column=2, value=":").font = normal_font_10
    ws.cell(row=r, column=3, value=sayi).font = normal_font_10
    ws.cell(row=r, column=7, value=tarih).font = normal_font_10
    ws.cell(row=r, column=7).alignment = right_align
    r += 1
    ws.cell(row=r, column=1, value="Konu").font = normal_font_10
    ws.cell(row=r, column=2, value=":").font = normal_font_10
    ws.cell(row=r, column=3, value="Teklifiniz").font = normal_font_10
    r += 2
    
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=7)
    c = ws.cell(row=r, column=1, value="Sayın Yetkili")
    c.font = bold_font_10
    c.alignment = center_align
    r += 2
    
    metin = "Aşağıda cinsi, özellikleri ve miktarları yazılı mallar / hizmetler 4734 sayılı Kamu İhale Kanunu'nun 22/d Maddesi gereğince Doğrudan Temin Usulüyle satın alınacaktır. İlgilenmeniz halinde KDV hariç teklifinizin bildirilmesini rica ederim / ederiz."
    ws.merge_cells(start_row=r, start_column=1, end_row=r+2, end_column=7)
    c = ws.cell(row=r, column=1, value=metin)
    c.font = normal_font
    c.alignment = Alignment(horizontal='justify', vertical='top', wrap_text=True)
    r += 4
    
    komisyon_uyeleri = get_komisyon(veri)
    if len(komisyon_uyeleri) > 0:
        for idx, uye in enumerate(komisyon_uyeleri[:3]): 
            col_pos = 1 + idx*2
            c_isim = ws.cell(row=r, column=col_pos, value=uye)
            c_isim.font = normal_font
            c_isim.alignment = center_align
            
            c_unvan = ws.cell(row=r+1, column=col_pos, value="Öğretmen")
            c_unvan.font = normal_font
            c_unvan.alignment = center_align
    r += 3
    
    # Yeni Super Header
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
    c1 = ws.cell(row=r, column=1, value="Satın Alınacak Malın")
    c1.alignment = left_align
    c1.font = bold_font_8
    
    ws.merge_cells(start_row=r, start_column=6, end_row=r, end_column=7)
    c2 = ws.cell(row=r, column=6, value="Teklif Edilen KDV Hariç")
    c2.alignment = left_align
    c2.font = bold_font_8
    for i in range(1, 8): ws.cell(row=r, column=i).border = thin_border
    
    r += 1
    headers = ["S.No", "Cinsi", "Özellikleri", "Ölçüsü", "Miktarı", "Birim Fiyatı (TL)", "Toplam Fiyatı (TL)"]
    for i, h in enumerate(headers, 1):
        c = ws.cell(row=r, column=i, value=h)
        c.font = bold_font_8
        c.alignment = left_align
        c.border = thin_border
        
    kalemler = veri.get("kalemler", [])
    gecerli_kalemler = [k for k in kalemler if k.get("cins", "").strip()]
    if not gecerli_kalemler: gecerli_kalemler = [{}]
    
    r += 1
    for idx, k in enumerate(gecerli_kalemler):
        val = str(k.get("miktar", ""))
        try:
            if float(val) == int(float(val)): val = str(int(float(val)))
        except: pass
        
        row_data = [str(idx+1), k.get("cins", ""), k.get("ozellik", ""), k.get("birim", ""), val, "", ""]
        for i, v in enumerate(row_data, 1):
            c = ws.cell(row=r, column=i, value=v)
            c.font = normal_font_8
            c.alignment = center_align
            c.border = thin_border
        r += 1
        
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
    c = ws.cell(row=r, column=1, value="KDV Hariç Teklif Edilen Toplam Fiyat:")
    c.font = bold_font_8
    c.alignment = right_align
    for i in range(1, 8): ws.cell(row=r, column=i).border = thin_border
    
    # 0,00 kaldirildi
    ws.cell(row=r, column=7, value="").alignment = center_align
    
    r += 1
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=7)
    c = ws.cell(row=r, column=1, value="DİĞER ŞARTLAR")
    c.font = bold_font_8
    c.alignment = center_align
    for i in range(1, 8): ws.cell(row=r, column=i).border = thin_border
    
    sartlar = [
        ("1- TESLİM SÜRESİ", "1 gün"),
        ("2- TESLİM EDİLECEK PARTİ MİKTARI", "1"),
        ("3- NAKLİYE VE SİGORTANIN KİME AİT OLDUĞU", "Satıcıya"),
        ("4- DİĞER ÖZEL ŞARTLAR", "YOK"),
        ("5- UYULMASI GEREKEN STANDARTLAR", "TSE"),
        ("6- TEKNİK ŞARTNAME", "YOK"),
        ("7- DİĞER HUSUSLAR", "YOK")
    ]
    
    for s in sartlar:
        r += 1
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
        c1 = ws.cell(row=r, column=1, value=s[0])
        c1.font = normal_font_8
        c1.alignment = left_align
        
        ws.merge_cells(start_row=r, start_column=6, end_row=r, end_column=7)
        c2 = ws.cell(row=r, column=6, value=s[1])
        c2.font = normal_font_8
        c2.alignment = left_align
        
        for i in range(1, 8): ws.cell(row=r, column=i).border = thin_border
        
    r += 2
    ws.merge_cells(start_row=r, start_column=1, end_row=r+2, end_column=7)
    bt = "Yukarıda belirtilen ve İdarenizce satın alınacak mal veya hizmetlerin cinsi, özellikleri, miktarı ve diğer şartları okudum, anladım. KDV hariç toplam .................................................. TL bedel ile belirtilen mal/hizmet satışını gerçekleştirmeyi kabul ve taahhüt ediyorum."
    ws.cell(row=r, column=1, value=bt).alignment = Alignment(horizontal='justify', vertical='top', wrap_text=True)
    
    r += 4
    ws.cell(row=r, column=6, value="Teklif Eden").alignment = center_align
    ws.cell(row=r+1, column=6, value="..../..../........").alignment = center_align
    ws.cell(row=r+2, column=6, value="Adı, Soyadı--Ticaret Ünvanı--İmza--Mühür").alignment = center_align
    
    r += 4
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=7)
    ws.cell(row=r, column=1, value="NOT :-Bu Belge Piyasa Fiyat Araştırması Tutanağına Eklenecektir.").font = normal_font
    
    cols_width = [6, 18, 18, 12, 12, 15, 15]
    for i, w in enumerate(cols_width, 1):
        ws.column_dimensions[chr(64+i)].width = w
        
    wb.save(dosya_adi)
    return dosya_adi

def uret_muayene_kabul_pdf(veri, hedef_klasor):
    import datetime
    import os
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.pagesizes import A4, portrait
    from reportlab.lib import colors
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    
    timestamp = datetime.datetime.now().strftime('%Y%m%d_%H%M%S')
    dosya_adi = os.path.join(hedef_klasor, f"Muayene_Kabul_{timestamp}.pdf")
    
    try:
        pdfmetrics.registerFont(TTFont('Arial_TR', r'C:\Windows\Fonts\arial.ttf'))
        font_name = 'Arial_TR'
    except:
        font_name = 'Helvetica'
        
    doc = SimpleDocTemplate(
        dosya_adi,
        pagesize=portrait(A4),
        rightMargin=15*mm,
        leftMargin=15*mm,
        topMargin=10*mm,
        bottomMargin=15*mm
    )
    
    style_normal = ParagraphStyle('Normal_TR', fontName=font_name, fontSize=9, leading=12, alignment=0)
    style_center = ParagraphStyle('Center_TR', fontName=font_name, fontSize=9, alignment=1, leading=12)
    # Üst başlık ve Muayene Kabul başlığı için aynı font ve boyut (kalın olmayan)
    style_header = ParagraphStyle('Header_TR', fontName=font_name, fontSize=10, alignment=1, leading=14)
    
    elements = []
    
    # 1. Üstteki 3 satır bilgi girişinden
    baslik = veri.get('resmi_baslik', '')
    if baslik:
        for line in baslik.split('\n'):
            line = line.strip()
            if line:
                elements.append(Paragraph(line, style_header))
    else:
        elements.append(Paragraph("T.C.", style_header))
        elements.append(Paragraph("KAYMAKAMLIĞI", style_header))
        elements.append(Paragraph("Okul Müdürlüğü", style_header))
        
    elements.append(Spacer(1, 10*mm))
    
    # Tablo Verileri
    table_data = []
    
    # Satır 0: Belge Başlığı
    table_data.append([
        Paragraph("MUAYENE VE KABUL BELGESİ", ParagraphStyle('Title_TR', fontName=font_name, fontSize=10, alignment=1, spaceAfter=5, spaceBefore=5)), 
        "", "", ""
    ])
    
    # Satır 1: Başlıklar
    table_data.append([
        Paragraph(x, ParagraphStyle('tb', fontName=font_name, fontSize=9, alignment=1)) for x in ["Sıra\nNo", "Satın Alınacak Malın", "Özellikleri", "Miktarı"]
    ])
    
    # Kalemler
    kalemler = veri.get('kalemler', [])
    for idx, k in enumerate(kalemler):
        cins = k.get('cins', '')
        ozellik = k.get('aciklama', '') or k.get('ozellik', '') or ''
        miktar = str(k.get('miktar', ''))
        birim = k.get('birim', '')
        miktar_str = f"{miktar} {birim}" if birim else miktar
        table_data.append([
            str(idx+1), 
            Paragraph(cins, ParagraphStyle('tn', fontName=font_name, fontSize=9)), 
            Paragraph(ozellik, ParagraphStyle('tn', fontName=font_name, fontSize=9)), 
            Paragraph(miktar_str, ParagraphStyle('tn', fontName=font_name, fontSize=9, alignment=1))
        ])
        
    # Alt Metin
    tarih = format_date(veri.get('belge_tarihi', ''))
    text = f"&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;Yukarıda yazılı malların / malzemelerin / işlerin, yapılan muayene neticesinde dokümanlarda belirtilen niteliklere uygun olduğu tespit edilmiştir. {tarih}"
    table_data.append([
        Paragraph(text, ParagraphStyle('Normal_TR', fontName=font_name, fontSize=9, leading=12, alignment=0)),
        "", "", ""
    ])
    
    # 5. Komisyon Üyeleri ve Görevleri (Başlık ve İmzalar Tek Hücrede)
    komisyon = veri.get('komisyon', {})
    kom_isimler = [
        komisyon.get('ihale_kom_muayene_1', ''),
        komisyon.get('ihale_kom_muayene_2', ''),
        komisyon.get('ihale_kom_muayene_3', '')
    ]
    kom_titles = ['Başkan', 'Üye', 'Üye']
    
    sig_data = [[Paragraph(f"{kom_isimler[0]}<br/>{kom_titles[0]}", style_center), 
                 Paragraph(f"{kom_isimler[1]}<br/>{kom_titles[1]}", style_center), 
                 Paragraph(f"{kom_isimler[2]}<br/>{kom_titles[2]}", style_center)]]
    
    sig_table = Table(sig_data, colWidths=[60*mm, 60*mm, 60*mm])
    sig_table.setStyle(TableStyle([
        ('ALIGN', (0,0), (-1,-1), 'CENTER'),
        ('VALIGN', (0,0), (-1,-1), 'TOP'),
    ]))
    
    komisyon_cell = [
        Spacer(1, 10*mm),
        Paragraph("MUAYENE VE KABUL GÖREVLİLERİ", ParagraphStyle('kom', fontName=font_name, fontSize=10, alignment=1, spaceBefore=5, spaceAfter=15)),
        sig_table
    ]
    
    table_data.append([komisyon_cell, "", "", ""])
    
    text_row_idx = len(table_data) - 2
    kom_sig_idx = len(table_data) - 1
    
    row_heights = [15*mm, 10*mm] + [8*mm]*len(kalemler) + [None, 45*mm]
    
    t = Table(table_data, colWidths=[15*mm, 85*mm, 50*mm, 30*mm], rowHeights=row_heights)
    t.setStyle(TableStyle([
        ('GRID', (0,0), (-1,-1), 0.5, colors.black),
        
        # Span Title
        ('SPAN', (0,0), (3,0)),
        ('ALIGN', (0,0), (3,0), 'CENTER'),
        ('VALIGN', (0,0), (3,0), 'MIDDLE'),
        
        # Span Text row
        ('SPAN', (0, text_row_idx), (3, text_row_idx)),
        ('VALIGN', (0, text_row_idx), (3, text_row_idx), 'TOP'),
        ('TOPPADDING', (0, text_row_idx), (3, text_row_idx), 8),
        ('BOTTOMPADDING', (0, text_row_idx), (3, text_row_idx), 8),
        ('LEFTPADDING', (0, text_row_idx), (3, text_row_idx), 10),
        ('RIGHTPADDING', (0, text_row_idx), (3, text_row_idx), 10),
        
        # Span Kom_sig (Başlık + İmzalar)
        ('SPAN', (0, kom_sig_idx), (3, kom_sig_idx)),
        ('VALIGN', (0, kom_sig_idx), (3, kom_sig_idx), 'TOP'),
        ('BOTTOMPADDING', (0, kom_sig_idx), (3, kom_sig_idx), 10),
        
        # Header alignment
        ('ALIGN', (0,1), (-1,1), 'CENTER'),
        ('VALIGN', (0,1), (-1,1), 'MIDDLE'),
        
        # Items alignment
        ('ALIGN', (0,2), (0, text_row_idx-1), 'CENTER'), # Sıra No
        ('ALIGN', (1,2), (2, text_row_idx-1), 'LEFT'),   # Mal ve Özellik
        ('ALIGN', (3,2), (3, text_row_idx-1), 'CENTER'), # Miktar
        ('VALIGN', (0,2), (-1, text_row_idx-1), 'MIDDLE'),
        
        ('FONTNAME', (0,0), (-1,-1), font_name),
        ('FONTSIZE', (0,0), (-1,-1), 9),
        ('WORDWRAP', (0,0), (-1,-1), True),
    ]))
    
    elements.append(t)
    
    doc.build(elements)
    return dosya_adi
def uret_muayene_kabul_excel(veri, hedef_klasor):
    import datetime
    import os
    import openpyxl
    from openpyxl.styles import Font, Alignment, Border, Side
    
    timestamp = datetime.datetime.now().strftime('%Y%m%d_%H%M%S')
    dosya_adi = os.path.join(hedef_klasor, f"Muayene_Kabul_{timestamp}.xlsx")
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Muayene Raporu"
    
    font_bold = Font(name='Times New Roman', bold=True, size=11)
    font_normal = Font(name='Times New Roman', size=10)
    font_title = Font(name='Times New Roman', bold=True, size=12)
    
    align_center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    align_left = Alignment(horizontal='left', vertical='center', wrap_text=True)
    align_justify = Alignment(horizontal='justify', vertical='top', wrap_text=True)
    
    thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
    
    ws.merge_cells('A1:E1')
    ws['A1'] = "M U A Y E N E   R A P O R U"
    ws['A1'].font = font_title
    ws['A1'].alignment = align_center
    
    headers = ["Sıra No", "Malzemenin Adı", "Miktarı", "Kabul Olunan Miktar", "Reddedilen Miktar"]
    ws.append([]) # row 2 empty
    ws.append(headers)
    for col in range(1, 6):
        cell = ws.cell(row=3, column=col)
        cell.font = font_bold
        cell.alignment = align_center
        cell.border = thin_border
        
    ws.column_dimensions['A'].width = 8
    ws.column_dimensions['B'].width = 40
    ws.column_dimensions['C'].width = 15
    ws.column_dimensions['D'].width = 20
    ws.column_dimensions['E'].width = 15
    
    kalemler = veri.get('kalemler', [])
    row_idx = 4
    for idx, k in enumerate(kalemler):
        cins = k.get('cins', '')
        miktar = str(k.get('miktar', ''))
        birim = k.get('birim', '')
        miktar_str = f"{miktar} {birim}" if birim else miktar
        
        ws.cell(row=row_idx, column=1, value=idx+1).border = thin_border
        ws.cell(row=row_idx, column=1).alignment = align_center
        ws.cell(row=row_idx, column=1).font = font_normal
        
        ws.cell(row=row_idx, column=2, value=cins).border = thin_border
        ws.cell(row=row_idx, column=2).alignment = align_left
        ws.cell(row=row_idx, column=2).font = font_normal
        
        ws.cell(row=row_idx, column=3, value=miktar_str).border = thin_border
        ws.cell(row=row_idx, column=3).alignment = align_center
        ws.cell(row=row_idx, column=3).font = font_normal
        
        ws.cell(row=row_idx, column=4, value=miktar_str).border = thin_border
        ws.cell(row=row_idx, column=4).alignment = align_center
        ws.cell(row=row_idx, column=4).font = font_normal
        
        ws.cell(row=row_idx, column=5, value="").border = thin_border
        ws.cell(row=row_idx, column=5).alignment = align_center
        ws.cell(row=row_idx, column=5).font = font_normal
        
        row_idx += 1
        
    # Boş satırlar kaldırıldı
        
    row_idx += 1
    tarih = format_date(veri.get('tarih', ''))
    kalem_sayisi = len(kalemler)
    text = f"Yukarıda yazılı {kalem_sayisi} kalem malzeme, Mal Alımlarında Uygulanacak Ara Denetim ile Mal Alımları Muayene ve Kabul Yönetmeliğine göre İhale Dokümanında yazılı şartlar dahilinde muayene edilerek neticesi yazılmıştır.   {tarih}"
    
    ws.merge_cells(f'A{row_idx}:E{row_idx+2}')
    cell = ws.cell(row=row_idx, column=1, value=text)
    cell.font = font_normal
    cell.alignment = align_justify
    
    row_idx += 4
    ws.merge_cells(f'A{row_idx}:E{row_idx}')
    cell = ws.cell(row=row_idx, column=1, value="M U A Y E N E   V E   K A B U L   K O M İ S Y O N U")
    cell.font = font_bold
    cell.alignment = align_center
    
    row_idx += 2
    komisyon = veri.get('komisyon', {})
    kom_isimler = [
        komisyon.get('ihale_kom_muayene_1', ''),
        komisyon.get('ihale_kom_muayene_2', ''),
        komisyon.get('ihale_kom_muayene_3', '')
    ]
    kom_titles = ["Başkan", "Üye", "Üye"]
    
    ws.cell(row=row_idx, column=1, value=kom_isimler[0]).font = font_normal
    ws.cell(row=row_idx, column=3, value=kom_isimler[1]).font = font_normal
    ws.cell(row=row_idx, column=5, value=kom_isimler[2]).font = font_normal
    ws.cell(row=row_idx, column=1).alignment = align_center
    ws.cell(row=row_idx, column=3).alignment = align_center
    ws.cell(row=row_idx, column=5).alignment = align_center
    
    ws.cell(row=row_idx+1, column=1, value=kom_titles[0]).font = font_normal
    ws.cell(row=row_idx+1, column=3, value=kom_titles[1]).font = font_normal
    ws.cell(row=row_idx+1, column=5, value=kom_titles[2]).font = font_normal
    ws.cell(row=row_idx+1, column=1).alignment = align_center
    ws.cell(row=row_idx+1, column=3).alignment = align_center
    ws.cell(row=row_idx+1, column=5).alignment = align_center
    
    wb.save(dosya_adi)
    return dosya_adi

def belge_uret(veri, pdf_yol):

    belge_tipi = veri.get("belge_tipi", "")
    fmt = veri.get("format", "pdf")
    hedef_klasor = get_ihale_klasoru(veri, pdf_yol)

    if not os.path.exists(hedef_klasor):
        os.makedirs(hedef_klasor, exist_ok=True)

    if belge_tipi == "fiyat_isteme":
        if fmt == "pdf":
            return uret_fiyat_isteme_pdf(veri, hedef_klasor)
        elif fmt == "excel":
            return uret_fiyat_isteme_excel(veri, hedef_klasor)
    elif belge_tipi == "yaklasik_maliyet":
        if fmt == "pdf":
            return uret_yaklasik_maliyet_pdf(veri, hedef_klasor)
        else:
            return uret_yaklasik_maliyet_excel(veri, hedef_klasor)

    elif belge_tipi == "ozel_fiyat_isteme":
        if fmt == "pdf":
            return uret_ozel_fiyat_isteme_pdf(veri, hedef_klasor)
        else:
            return uret_ozel_fiyat_isteme_excel(veri, hedef_klasor)

    elif belge_tipi == "piyasa_arastirmasi":
        if fmt == "pdf":
            return uret_piyasa_arastirmasi_pdf(veri, hedef_klasor)
        else:
            return uret_piyasa_arastirmasi_excel(veri, hedef_klasor)

    elif belge_tipi == "muayene_kabul":
        if fmt == "pdf":
            return uret_muayene_kabul_pdf(veri, hedef_klasor)
        else:
            return uret_muayene_kabul_excel(veri, hedef_klasor)

    return None


def uret_piyasa_arastirmasi_pdf(veri, hedef_klasor):
    import datetime
    import os
    from reportlab.pdfgen import canvas
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib import colors
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    
    timestamp = datetime.datetime.now().strftime('%Y%m%d_%H%M%S')
    dosya_adi = os.path.join(hedef_klasor, f"Piyasa_Fiyat_Arastirmasi_{timestamp}.pdf")
    
    try:
        pdfmetrics.registerFont(TTFont('Arial_TR', r'C:\Windows\Fonts\arial.ttf'))
        font_name = 'Arial_TR'
    except:
        font_name = 'Helvetica'
        
    def add_page_number(canvas, doc):
        canvas.saveState()
        canvas.setFont(font_name, 6)
        canvas.translate(8*mm, 20*mm)
        canvas.rotate(90)
        canvas.drawString(0, 0, "M.H.B.Y. Örnek No: 2")
        canvas.restoreState()
        
        canvas.saveState()
        canvas.setFont(font_name, 6)
        canvas.drawString(10*mm, 8*mm, "*Piyasa Fiyat Araştırması yapılacak kişi / firma, yer sayısına ve Piyasa Fiyat Araştırması için görevlendirilecek personelin sayısına ihale yetkilisi karar verebilecektir.")
        canvas.restoreState()
        
    doc = SimpleDocTemplate(
        dosya_adi,
        pagesize=landscape(A4),
        rightMargin=10*mm,
        leftMargin=15*mm,
        topMargin=10*mm,
        bottomMargin=15*mm
    )
    
    style_normal = ParagraphStyle('Normal_TR', fontName=font_name, fontSize=8, leading=10)
    style_center = ParagraphStyle('Center_TR', fontName=font_name, fontSize=8, alignment=1, leading=10)
    style_right = ParagraphStyle('Right_TR', fontName=font_name, fontSize=8, alignment=2, leading=10)
    style_title = ParagraphStyle('Title_TR', fontName=font_name, fontSize=11, alignment=1, leading=14, spaceAfter=8)
    
    elements = []
    
    tarih = format_date(veri.get("belge_tarihi", ""))
    konu = veri.get("ihale_konusu", "")
    komisyon_onaylari = veri.get("komisyon_onaylari", {}) or {}

    idare_adi = _idare_adi_olustur(veri.get("resmi_baslik", ""))

    onay_metni = ""
    for kom_id in ["ihale_kom_piyasa"]:
        onay = komisyon_onaylari.get(kom_id, {})
        if onay:
            kurum = (onay.get("kurum") or "").strip()
            sayi = (onay.get("sayi") or "").strip()
            tarih_onay = (onay.get("tarih") or "").strip()
            if kurum or sayi or tarih_onay:
                kurum_metni = _idare_adi_olustur(kurum)
                
                kurum_temiz = kurum_metni
                if kurum_temiz.lower().endswith("müdürlüğü"):
                    kurum_temiz = kurum_temiz[:-9].strip()

                if kurum and sayi and tarih_onay:
                    onay_metni = f"{kurum_temiz} Müdürlüğünün {sayi} sayılı ve {tarih_onay} tarihli onayı"
                elif kurum and sayi:
                    onay_metni = f"{kurum_temiz} Müdürlüğünün {sayi} sayılı onayı"
                elif kurum and tarih_onay:
                    onay_metni = f"{kurum_temiz} Müdürlüğünün {tarih_onay} tarihli onayı"
                elif sayi and tarih_onay:
                    onay_metni = f"Müdürlüğünün {sayi} sayılı ve {tarih_onay} tarihli onayı"
                break

    if not onay_metni:
        onay_metni = "İdare onayı"

    elements.append(Paragraph("PİYASA FİYAT ARAŞTIRMASI TUTANAĞI", style_title))
    
    firmalar = veri.get("firmalar", ["", "", "", ""])
    firma_vergiler = veri.get("firma_vergiler", ["", "", "", ""])
    firma_adresleri = veri.get("firma_adresleri", ["", "", "", ""])
    while len(firmalar) < 4: firmalar.append("")
    while len(firma_vergiler) < 4: firma_vergiler.append("")
    while len(firma_adresleri) < 4: firma_adresleri.append("")
    
    table_data = []
    
    row0 = ["İdarenin Adı", "", "", "", "", idare_adi, "", "", "", "", "", "", ""]
    table_data.append([Paragraph(x, style_normal) if isinstance(x, str) and x else x for x in row0])
    
    row1 = ["Yapılan İş / Mal / Hizmetin Adı, Niteliği", "", "", "", "", konu, "", "", "", "", "", "", ""]
    table_data.append([Paragraph(x, style_normal) if isinstance(x, str) and x else x for x in row1])
    
    row2 = ["Yetkilendirilen Görevlilere ilişkin Onayın Tarih ve Nosu", "", "", "", "", onay_metni, "", "", "", "", "", "", ""]
    table_data.append([Paragraph(x.replace("\n", "<br/>"), style_normal) if isinstance(x, str) and x else x for x in row2])
    
    # Headers (4 rows: 3, 4, 5, 6)
    row3 = ["Sıra<br/>No", "Satın Alınacak Malın", "", "", "", "Kişiler / Firmalar ve Fiyat Teklifleri ( KDV Hariç )", "", "", "", "", "", "", ""]
    table_data.append([Paragraph(x, style_center) if isinstance(x, str) and x else x for x in row3])
    
    row4 = ["", "", "", "", "", "1", "", "2", "", "3", "", "4", ""]
    table_data.append([Paragraph(x, style_center) if isinstance(x, str) and x else x for x in row4])
    
    f_headers = []
    for i in range(4):
        f = firmalar[i].strip()
        v = firma_vergiler[i].strip()
        f_text = f"{f}" if f else ""
        if v:
            if len(v) == 10: f_text += f"<br/>(Vergi No: {v})"
            elif len(v) == 11: f_text += f"<br/>(T.C. No: {v})"
            else: f_text += f"<br/>({v})"
        f_headers.append(Paragraph(f_text, style_center))
        
    row5 = ["", "Cinsi", "Özelliği", "Miktarı", "Ölçüsü", f_headers[0], "", f_headers[1], "", f_headers[2], "", f_headers[3], ""]
    table_data.append([Paragraph(x, style_center) if isinstance(x, str) and x else x for x in row5])
    
    row6 = ["", "", "", "", ""]
    for i in range(4):
        row6.extend(["Birim<br/>Fiyat", "Toplam<br/>Fiyat"])
    table_data.append([Paragraph(x, style_center) if isinstance(x, str) and x else x for x in row6])
    
    kalemler = veri.get("kalemler", [])
    gecerli_kalemler = [k for k in kalemler if k.get("cins", "").strip()]
    if not gecerli_kalemler: gecerli_kalemler = [{}]
        
    col_sums = [0.0, 0.0, 0.0, 0.0]
    
    for idx, k in enumerate(gecerli_kalemler):
        val = str(k.get("miktar", ""))
        miktar_float = 0.0
        try:
            miktar_float = float(val)
            if miktar_float == int(miktar_float): val = str(int(miktar_float))
        except: pass
        
        row = [
            Paragraph(str(idx+1), style_center),
            Paragraph(k.get("cins", ""), style_normal),
            Paragraph(k.get("ozellik", ""), style_normal),
            Paragraph(val, style_center),
            Paragraph(k.get("birim", ""), style_center)
        ]
        
        fiyatlar = k.get("fiyatlar", [])
        
        for fidx in range(4):
            bf = ""
            tf = ""
            if fidx < len(fiyatlar) and fiyatlar[fidx] not in ["", None]:
                try:
                    fiyat_val = float(fiyatlar[fidx])
                    if fiyat_val > 0:
                        bf = f"{fiyat_val:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
                        tf_float = fiyat_val * miktar_float
                        tf = f"{tf_float:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
                        col_sums[fidx] += tf_float
                except:
                    pass
            row.extend([Paragraph(bf, style_right), Paragraph(tf, style_right)])
            
        table_data.append(row)
        
    t_row = [
        Paragraph("Kişilerin / Firmaların Teklifleri Toplamı ( KDV HARİÇ )", style_normal), "", "", "", "", 
        Paragraph("TEKLİF<br/>TOPLAMI", style_center), Paragraph(f"{col_sums[0]:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".") if col_sums[0]>0 else "", style_center), 
        Paragraph("TEKLİF<br/>TOPLAMI", style_center), Paragraph(f"{col_sums[1]:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".") if col_sums[1]>0 else "", style_center), 
        Paragraph("TEKLİF<br/>TOPLAMI", style_center), Paragraph(f"{col_sums[2]:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".") if col_sums[2]>0 else "", style_center), 
        Paragraph("TEKLİF<br/>TOPLAMI", style_center), Paragraph(f"{col_sums[3]:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".") if col_sums[3]>0 else "", style_center)
    ]
    table_data.append(t_row)
    
    # Bottom block (3 rows)
    min_idx = -1
    min_val = float('inf')
    for i, v in enumerate(col_sums):
        if v > 0 and v < min_val:
            min_val = v
            min_idx = i
            
    uygun_firma = firmalar[min_idx] if min_idx != -1 else ""
    uygun_adres = firma_adresleri[min_idx] if min_idx != -1 else ""
    uygun_tutar = f"{min_val:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".") if min_idx != -1 else ""
    
    b_row1 = ["", "", "", "", "", Paragraph("Teklifi Uygun Görülen Kişi / Firma", style_center), "", "", "", "", "", "", ""]
    table_data.append(b_row1)
    
    b_row2 = [Paragraph("Satın Alınacak Malın", style_center), "", "", "", "", Paragraph("Adı", style_center), "", Paragraph("Adresi", style_center), "", "", Paragraph("Teklif Ettiği Toplam Fiyat (KDV Hariç)", style_center), "", ""]
    table_data.append(b_row2)
    
    b_row3 = [Paragraph("Tümünün Bu Kişi / Firmadan Alımı Uygun Görülmüştür.", style_center), "", "", "", "", Paragraph(uygun_firma, style_center), "", Paragraph(uygun_adres, style_center), "", "", Paragraph(uygun_tutar, style_center), "", ""]
    table_data.append(b_row3)
    
    col_widths = [10*mm, 35*mm, 35*mm, 15*mm, 15*mm, 20*mm, 21*mm, 20*mm, 21*mm, 20*mm, 21*mm, 20*mm, 21*mm]
    
    t = Table(table_data, colWidths=col_widths, repeatRows=7)
    
    ts = [
        ('FONTNAME', (0,0), (-1,-1), font_name),
        ('FONTSIZE', (0,0), (-1,-1), 8),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('INNERGRID', (0,0), (-1,-1), 0.25, colors.black),
        ('BOX', (0,0), (-1,-1), 0.25, colors.black),
        
        ('SPAN', (0,0), (4,0)),
        ('SPAN', (5,0), (12,0)),
        
        ('SPAN', (0,1), (4,1)),
        ('SPAN', (5,1), (12,1)),
        
        ('SPAN', (0,2), (4,2)),
        ('SPAN', (5,2), (12,2)),
        
        # Row 3 (Headers)
        ('SPAN', (0,3), (0,6)), # S.No (spans row 3,4,5,6)
        ('SPAN', (1,3), (4,4)), # Satin Alinacak Malin (spans col 1-4, row 3-4)
        ('SPAN', (5,3), (12,3)), # Kisiler/Firmalar (spans col 5-12, row 3)
        
        # Row 4 (1,2,3,4)
        ('SPAN', (5,4), (6,4)),
        ('SPAN', (7,4), (8,4)),
        ('SPAN', (9,4), (10,4)),
        ('SPAN', (11,4), (12,4)),
        
        # Row 5 (Cinsi, Firm Names)
        ('SPAN', (1,5), (1,6)), # Cinsi
        ('SPAN', (2,5), (2,6)), # Ozelligi
        ('SPAN', (3,5), (3,6)), # Miktari
        ('SPAN', (4,5), (4,6)), # Olcusu
        
        ('SPAN', (5,5), (6,5)), # Firm 1
        ('SPAN', (7,5), (8,5)), # Firm 2
        ('SPAN', (9,5), (10,5)), # Firm 3
        ('SPAN', (11,5), (12,5)), # Firm 4
        
        # Totals row
        ('SPAN', (0, -4), (4, -4)),
        
        # Bottom Block
        ('SPAN', (0, -3), (4, -3)), # blank lead row for the moved block
        ('SPAN', (5, -3), (12, -3)), # Teklifi uygun gorulen...

        ('SPAN', (0, -2), (4, -2)), # Satın Alınacak Malın
        ('SPAN', (5, -2), (6, -2)), # Adi
        ('SPAN', (7, -2), (9, -2)), # Adresi
        ('SPAN', (10, -2), (12, -2)), # Teklif Ettigi...

        ('SPAN', (0, -1), (4, -1)), # Tümünün Bu Kişi / Firmadan Alımı Uygun Görülmüştür.
        ('SPAN', (5, -1), (6, -1)), # Adi value
        ('SPAN', (7, -1), (9, -1)), # Adresi value
        ('SPAN', (10, -1), (12, -1)), # Teklif Ettigi value
        
        ('LEFTPADDING', (0,0), (-1,-1), 2),
        ('RIGHTPADDING', (0,0), (-1,-1), 2),
    ]
    t.setStyle(TableStyle(ts))
    elements.append(t)
    elements.append(Spacer(1, 4*mm))
    
    p1 = f"4734 Sayılı Kamu İhale Kanununun 22'nci Maddesi uyarınca Doğrudan Temin Usulüyle yapılacak alımlara ilişkin yapılan piyasa araştırmasında firmalarca / kişilerce teklif edilen fiyatlar tarafımca / tarafımızca değerlendirilerek yukarıda adı ve adresi belirtilen kişi / firmadan alım yapılması uygun görülmüştür. {tarih}"
    elements.append(Paragraph(p1, ParagraphStyle('P1', fontName=font_name, fontSize=9, alignment=4, leading=11)))
    elements.append(Spacer(1, 8*mm))
    
    elements.append(Paragraph("<b>PİYASA FİYAT ARAŞTIRMASI GÖREVLİSİ / GÖREVLİLERİ</b>", style_center))
    elements.append(Spacer(1, 4*mm))
    
    komisyon_uyeleri = get_komisyon(veri)
    if not komisyon_uyeleri: komisyon_uyeleri = ["", "", ""]
    
    imza_data = [[]]
    for uye in komisyon_uyeleri:
        imza_data[0].append(Paragraph(f"{uye}<br/>Öğretmen", style_center))
        
    col_w = 277*mm / max(len(komisyon_uyeleri), 1)
    sig_table = Table(imza_data, colWidths=[col_w] * len(komisyon_uyeleri))
    sig_table.setStyle(TableStyle([('ALIGN', (0,0), (-1,-1), 'CENTER'), ('VALIGN', (0,0), (-1,-1), 'TOP')]))
    elements.append(sig_table)
        
    doc.build(elements, onFirstPage=add_page_number, onLaterPages=add_page_number)
    return dosya_adi

def uret_piyasa_arastirmasi_excel(veri, hedef_klasor):
    import datetime
    import os
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, Border, Side
    
    timestamp = datetime.datetime.now().strftime('%Y%m%d_%H%M%S')
    dosya_adi = os.path.join(hedef_klasor, f"Piyasa_Fiyat_Arastirmasi_{timestamp}.xlsx")
    
    wb = Workbook()
    ws = wb.active
    ws.title = "Piyasa Fiyat Arastirmasi"
    
    normal_font = Font(name='Arial', size=9)
    bold_font = Font(name='Arial', size=9, bold=True)
    title_font = Font(name='Arial', size=11, bold=True)
    
    center_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
    left_align = Alignment(horizontal='left', vertical='center', wrap_text=True)
    right_align = Alignment(horizontal='right', vertical='center', wrap_text=True)
    thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
    
    ws.merge_cells('A1:M1')
    ws['A1'] = "PİYASA FİYAT ARAŞTIRMASI TUTANAĞI"
    ws['A1'].font = title_font
    ws['A1'].alignment = center_align
    
    tarih = format_date(veri.get("belge_tarihi", ""))
    konu = veri.get("ihale_konusu", "")
    komisyon_onaylari = veri.get("komisyon_onaylari", {}) or {}
    idare_adi = _idare_adi_olustur(veri.get("resmi_baslik", ""))
    onay_metni = "İdare onayı"
    for kom_id in ["ihale_kom_piyasa"]:
        onay = komisyon_onaylari.get(kom_id, {})
        if onay:
            kurum = (onay.get("kurum") or "").strip()
            sayi = (onay.get("sayi") or "").strip()
            tarih_onay = (onay.get("tarih") or "").strip()
            if kurum or sayi or tarih_onay:
                kurum_metni = _idare_adi_olustur(kurum)
                
                kurum_temiz = kurum_metni
                if kurum_temiz.lower().endswith("müdürlüğü"):
                    kurum_temiz = kurum_temiz[:-9].strip()

                if kurum and sayi and tarih_onay:
                    onay_metni = f"{kurum_temiz} Müdürlüğünün {sayi} sayılı ve {tarih_onay} tarihli onayı"
                elif kurum and sayi:
                    onay_metni = f"{kurum_temiz} Müdürlüğünün {sayi} sayılı onayı"
                elif kurum and tarih_onay:
                    onay_metni = f"{kurum_temiz} Müdürlüğünün {tarih_onay} tarihli onayı"
                elif sayi and tarih_onay:
                    onay_metni = f"Müdürlüğünün {sayi} sayılı ve {tarih_onay} tarihli onayı"
                break
    
    r = 3
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
    ws.cell(row=r, column=1, value="İdarenin Adı").font = normal_font
    ws.merge_cells(start_row=r, start_column=6, end_row=r, end_column=13)
    ws.cell(row=r, column=6, value=idare_adi).font = normal_font
    for i in range(1, 14): ws.cell(row=r, column=i).border = thin_border
    r += 1
    
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
    ws.cell(row=r, column=1, value="Yapılan İş / Mal / Hizmetin Adı, Niteliği").font = normal_font
    ws.merge_cells(start_row=r, start_column=6, end_row=r, end_column=13)
    ws.cell(row=r, column=6, value=konu).font = normal_font
    for i in range(1, 14): ws.cell(row=r, column=i).border = thin_border
    r += 1
    
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
    ws.cell(row=r, column=1, value="Yetkilendirilen Görevlilere ilişkin Onayın Tarih ve Nosu").font = normal_font
    ws.cell(row=r, column=1).alignment = left_align
    ws.merge_cells(start_row=r, start_column=6, end_row=r, end_column=13)
    ws.cell(row=r, column=6, value=onay_metni).font = normal_font
    for i in range(1, 14): ws.cell(row=r, column=i).border = thin_border
    r += 1
    
    firmalar = veri.get("firmalar", ["", "", "", ""])
    firma_vergiler = veri.get("firma_vergiler", ["", "", "", ""])
    firma_adresleri = veri.get("firma_adresleri", ["", "", "", ""])
    while len(firmalar) < 4: firmalar.append("")
    while len(firma_vergiler) < 4: firma_vergiler.append("")
    while len(firma_adresleri) < 4: firma_adresleri.append("")
    
    # Headers Row 1
    ws.merge_cells(start_row=r, start_column=1, end_row=r+3, end_column=1)
    ws.cell(row=r, column=1, value="Sıra\nNo").alignment = center_align
    ws.merge_cells(start_row=r, start_column=2, end_row=r+1, end_column=5)
    ws.cell(row=r, column=2, value="Satın Alınacak Malın").alignment = center_align
    ws.merge_cells(start_row=r, start_column=6, end_row=r, end_column=13)
    ws.cell(row=r, column=6, value="Kişiler / Firmalar ve Fiyat Teklifleri ( KDV Hariç )").alignment = center_align
    for i in range(1, 14): ws.cell(row=r, column=i).border = thin_border
    r += 1
    
    # Headers Row 2 (1,2,3,4)
    for i in range(4):
        col_start = 6 + i*2
        ws.merge_cells(start_row=r, start_column=col_start, end_row=r, end_column=col_start+1)
        ws.cell(row=r, column=col_start, value=str(i+1)).alignment = center_align
    for i in range(1, 14): ws.cell(row=r, column=i).border = thin_border
    r += 1
    
    # Headers Row 3 (Cinsi, Ozelligi..., Firm Names)
    ws.merge_cells(start_row=r, start_column=2, end_row=r+1, end_column=2)
    ws.cell(row=r, column=2, value="Cinsi").alignment = center_align
    ws.merge_cells(start_row=r, start_column=3, end_row=r+1, end_column=3)
    ws.cell(row=r, column=3, value="Özelliği").alignment = center_align
    ws.merge_cells(start_row=r, start_column=4, end_row=r+1, end_column=4)
    ws.cell(row=r, column=4, value="Miktarı").alignment = center_align
    ws.merge_cells(start_row=r, start_column=5, end_row=r+1, end_column=5)
    ws.cell(row=r, column=5, value="Ölçüsü").alignment = center_align
    
    for i in range(4):
        col_start = 6 + i*2
        ws.merge_cells(start_row=r, start_column=col_start, end_row=r, end_column=col_start+1)
        f = firmalar[i].strip()
        v = firma_vergiler[i].strip()
        f_text = f"{f}" if f else ""
        if v:
            if len(v) == 10: f_text += f"\n(Vergi No: {v})"
            elif len(v) == 11: f_text += f"\n(T.C. No: {v})"
            else: f_text += f"\n({v})"
        ws.cell(row=r, column=col_start, value=f_text).alignment = center_align
        ws.cell(row=r+1, column=col_start, value="Birim Fiyat").alignment = center_align
        ws.cell(row=r+1, column=col_start+1, value="Toplam Fiyat").alignment = center_align
        
    for i in range(1, 14):
        ws.cell(row=r, column=i).border = thin_border
        ws.cell(row=r, column=i).font = bold_font
        ws.cell(row=r+1, column=i).border = thin_border
        ws.cell(row=r+1, column=i).font = bold_font
        
    r += 2
    
    kalemler = veri.get("kalemler", [])
    gecerli_kalemler = [k for k in kalemler if k.get("cins", "").strip()]
    if not gecerli_kalemler: gecerli_kalemler = [{}]
        
    col_sums = [0.0, 0.0, 0.0, 0.0]
    
    for idx, k in enumerate(gecerli_kalemler):
        val = str(k.get("miktar", ""))
        miktar_float = 0.0
        try:
            miktar_float = float(val)
            if miktar_float == int(miktar_float): val = str(int(miktar_float))
        except: pass
        
        ws.cell(row=r, column=1, value=str(idx+1)).alignment = center_align
        ws.cell(row=r, column=2, value=k.get("cins", "")).alignment = left_align
        ws.cell(row=r, column=3, value=k.get("ozellik", "")).alignment = left_align
        ws.cell(row=r, column=4, value=val).alignment = center_align
        ws.cell(row=r, column=5, value=k.get("birim", "")).alignment = center_align
        
        fiyatlar = k.get("fiyatlar", [])
        
        for fidx in range(4):
            col_start = 6 + fidx*2
            bf = ""
            tf = ""
            if fidx < len(fiyatlar) and fiyatlar[fidx] not in ["", None]:
                try:
                    fiyat_val = float(fiyatlar[fidx])
                    if fiyat_val > 0:
                        bf = fiyat_val
                        tf_float = fiyat_val * miktar_float
                        tf = tf_float
                        col_sums[fidx] += tf_float
                except:
                    pass
            ws.cell(row=r, column=col_start, value=bf).alignment = right_align
            if isinstance(bf, float): ws.cell(row=r, column=col_start).number_format = '#,##0.00'
            ws.cell(row=r, column=col_start+1, value=tf).alignment = right_align
            if isinstance(tf, float): ws.cell(row=r, column=col_start+1).number_format = '#,##0.00'
            
        for i in range(1, 14):
            ws.cell(row=r, column=i).border = thin_border
            ws.cell(row=r, column=i).font = normal_font
        r += 1
        
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
    ws.cell(row=r, column=1, value="Kişilerin / Firmaların Teklifleri Toplamı ( KDV HARİÇ )").font = bold_font
    
    for fidx in range(4):
        col_start = 6 + fidx*2
        ws.cell(row=r, column=col_start, value="TEKLİF TOPLAMI").alignment = center_align
        val = col_sums[fidx] if col_sums[fidx] > 0 else ""
        ws.cell(row=r, column=col_start+1, value=val).alignment = center_align
        if isinstance(val, float): ws.cell(row=r, column=col_start+1).number_format = '#,##0.00'
        
    for i in range(1, 14):
        ws.cell(row=r, column=i).border = thin_border
        ws.cell(row=r, column=i).font = bold_font
    r += 1
    
    # Bottom Block
    min_idx = -1
    min_val = float('inf')
    for i, v in enumerate(col_sums):
        if v > 0 and v < min_val:
            min_val = v
            min_idx = i
            
    uygun_firma = firmalar[min_idx] if min_idx != -1 else ""
    uygun_adres = firma_adresleri[min_idx] if min_idx != -1 else ""
    uygun_tutar = min_val if min_idx != -1 else ""
    
    # Row 1
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
    ws.cell(row=r, column=1, value="Satın Alınacak Malın").alignment = center_align
    ws.merge_cells(start_row=r, start_column=6, end_row=r, end_column=13)
    ws.cell(row=r, column=6, value="Teklifi Uygun Görülen Kişi / Firma").alignment = center_align
    for i in range(1, 14): ws.cell(row=r, column=i).border = thin_border
    r += 1
    
    # Row 2
    ws.merge_cells(start_row=r, start_column=1, end_row=r+1, end_column=1)
    ws.cell(row=r, column=1, value="Tümünün").alignment = center_align
    ws.merge_cells(start_row=r, start_column=2, end_row=r+1, end_column=5)
    ws.cell(row=r, column=2, value="Bu Kişi / Firmadan Alımı Uygun Görülmüştür.").alignment = center_align
    
    ws.merge_cells(start_row=r, start_column=6, end_row=r, end_column=7)
    ws.cell(row=r, column=6, value="Adı").alignment = center_align
    ws.merge_cells(start_row=r, start_column=8, end_row=r, end_column=10)
    ws.cell(row=r, column=8, value="Adresi").alignment = center_align
    ws.merge_cells(start_row=r, start_column=11, end_row=r, end_column=13)
    ws.cell(row=r, column=11, value="Teklif Ettiği Toplam Fiyat (KDV Hariç)").alignment = center_align
    for i in range(1, 14): ws.cell(row=r, column=i).border = thin_border
    r += 1
    
    # Row 3
    ws.merge_cells(start_row=r, start_column=6, end_row=r, end_column=7)
    ws.cell(row=r, column=6, value=uygun_firma).alignment = center_align
    ws.merge_cells(start_row=r, start_column=8, end_row=r, end_column=10)
    ws.cell(row=r, column=8, value=uygun_adres).alignment = center_align
    ws.merge_cells(start_row=r, start_column=11, end_row=r, end_column=13)
    ws.cell(row=r, column=11, value=uygun_tutar).alignment = center_align
    if isinstance(uygun_tutar, float): ws.cell(row=r, column=11).number_format = '#,##0.00'
    for i in range(1, 14): ws.cell(row=r, column=i).border = thin_border
    
    # Fonts for bottom block
    for i in range(r-2, r+1):
        for j in range(1, 14):
            ws.cell(row=i, column=j).font = normal_font
            
    r += 2
    ws.merge_cells(start_row=r, start_column=1, end_row=r+1, end_column=13)
    p1 = f"4734 Sayılı Kamu İhale Kanununun 22'nci Maddesi uyarınca Doğrudan Temin Usulüyle yapılacak alımlara ilişkin yapılan piyasa araştırmasında firmalarca / kişilerce teklif edilen fiyatlar tarafımca / tarafımızca değerlendirilerek yukarıda adı ve adresi belirtilen kişi / firmadan alım yapılması uygun görülmüştür. {tarih}"
    ws.cell(row=r, column=1, value=p1).alignment = left_align
    r += 3
    
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=13)
    ws.cell(row=r, column=1, value="PİYASA FİYAT ARAŞTIRMASI GÖREVLİSİ / GÖREVLİLERİ").alignment = center_align
    ws.cell(row=r, column=1).font = bold_font
    
    r += 2
    komisyon_uyeleri = get_komisyon(veri)
    if not komisyon_uyeleri: komisyon_uyeleri = ["", "", ""]
    
    step = 13 // max(len(komisyon_uyeleri), 1)
    for idx, uye in enumerate(komisyon_uyeleri):
        col = 1 + idx*step
        ws.cell(row=r, column=col, value=uye).alignment = center_align
        ws.cell(row=r+1, column=col, value="Öğretmen").alignment = center_align
        
    ws.cell(row=r+4, column=1, value="*Piyasa Fiyat Araştırması yapılacak kişi / firma, yer sayısına ve Piyasa Fiyat Araştırması için görevlendirilecek personelin sayısına ihale yetkilisi karar verebilecektir.").font = Font(name='Arial', size=7)
        
    cols_width = [5, 20, 20, 10, 10, 12, 12, 12, 12, 12, 12, 12, 12]
    for i, w in enumerate(cols_width, 1):
        ws.column_dimensions[chr(64+i)].width = w
        
    wb.save(dosya_adi)
    return dosya_adi