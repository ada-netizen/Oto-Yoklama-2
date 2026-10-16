import os, logging, tempfile, re, shutil, uuid, json
import platform, subprocess
from datetime import datetime, timedelta
import pandas as pd
from dependencies import yollar, db, islem_loglari, islem_logla
from sistem_motoru import SistemMotoru
from models import *
from araclar import VeriAraclari

MAX_IMPORT_SIZE = 25 * 1024 * 1024
ALLOWED_IMPORT_EXTENSIONS = {".xlsx", ".xls", ".csv"}

def dosyayi_otomatik_ac(dosya_yolu):
    """Oluşturulan PDF veya Excel dosyasını bilgisayarın varsayılan programıyla anında açar"""
    try:
        if platform.system() == 'Windows':
            norm_yol = os.path.normpath(dosya_yolu)
            os.startfile(norm_yol)
        elif platform.system() == 'Darwin':
            subprocess.call(('open', dosya_yolu))
        else:
            subprocess.call(('xdg-open', dosya_yolu))
    except Exception as e:
        logging.error(f"dosyayi_otomatik_ac hatasi ({dosya_yolu}): {e}")

def _sablon_excel_yolu_olustur():
    hedef_yol = os.path.join(os.path.dirname(__file__), "sablonlar", "sablon.xlsx")

    if not os.path.exists(hedef_yol):
        try:
            import openpyxl
            from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

            os.makedirs(os.path.dirname(hedef_yol), exist_ok=True)
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Sablon"

            # MYS uyumlu başlık satırı
            basliklar = [
                "Ürün Adı", "Miktar", "Ölçü Birimi",
                "Ürün No", "Model", "Marka",
                "Birim Tutar", "Firma VKN/TCKN"
            ]
            thin = Side(style="thin")
            border = Border(left=thin, right=thin, top=thin, bottom=thin)
            fill  = PatternFill("solid", fgColor="DDEEFF")

            for col_idx, baslik in enumerate(basliklar, start=1):
                cell = ws.cell(row=1, column=col_idx, value=baslik)
                cell.font      = Font(name="Calibri", bold=True, size=10)
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                cell.border    = border
                cell.fill      = fill

            # Sütun genişlikleri
            genislikler = [40, 10, 15, 15, 15, 15, 18, 20]
            for i, g in enumerate(genislikler, start=1):
                ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = g

            ws.row_dimensions[1].height = 25
            ws.freeze_panes = "A2"
            wb.save(hedef_yol)
            logging.info(f"MYS uyumlu şablon oluşturuldu: {hedef_yol}")
        except Exception as e:
            logging.error(f"Şablon oluşturulamadı: {e}")

    return hedef_yol if os.path.exists(hedef_yol) else None



def excel_yukleme_dogrula(dosya):
    """Excel/CSV yüklemelerini boyut ve uzantı açısından kontrol eder."""
    dosya_adi = dosya.filename or ""
    uzanti = os.path.splitext(dosya_adi)[1].lower()
    if uzanti not in ALLOWED_IMPORT_EXTENSIONS:
        return None, "Yalnızca XLSX, XLS veya CSV dosyaları yüklenebilir."

    try:
        dosya.file.seek(0, os.SEEK_END)
        boyut = dosya.file.tell()
        dosya.file.seek(0)
    except (AttributeError, OSError) as e:
        return None, f"Dosya okunamadı: {e}"

    if boyut == 0:
        return None, "Yüklenen dosya boş."
    if boyut > MAX_IMPORT_SIZE:
        return None, "Dosya boyutu 25 MB sınırını aşamaz."
    return uzanti, None


from contextlib import contextmanager
from fastapi import UploadFile

@contextmanager
def gecici_dosya_olustur(dosya: UploadFile, prefix="temp_"):
    # Güvenli geçici dosya adı: Türkçe karakter veya özel karakter içeren
    # orijinal dosya adı yerine uuid tabanlı güvenli bir ad kullanıyoruz.
    uzanti = os.path.splitext(dosya.filename or ".pdf")[1] or ".pdf"
    temp_yol = f"{prefix}{uuid.uuid4().hex[:8]}{uzanti}"
    with open(temp_yol, "wb") as buffer:
        shutil.copyfileobj(dosya.file, buffer)
    try:
        yield temp_yol
    finally:
        if os.path.exists(temp_yol):
            try:
                os.remove(temp_yol)
            except Exception:
                pass


_son_ayarlar_mtime = 0
ayarlar_cache = None

def ayarlari_al():
    """Ayarları sadece dosya degistiginde okur (Cache)."""
    global ayarlar_cache, _son_ayarlar_mtime
    try:
        mtime = os.path.getmtime(yollar["AYARLAR"])
    except OSError:
        mtime = 0
        
    if mtime != _son_ayarlar_mtime or _son_ayarlar_mtime == 0 or ayarlar_cache is None:
        ayarlar_cache = SistemMotoru.ayarlari_yukle(yollar["AYARLAR"])
        _son_ayarlar_mtime = mtime
    return ayarlar_cache


def pdf_klasoru_hazirla():
    ayar = ayarlari_al()
    ana_klasor = ayar.get("pdf_kayit_klasoru", yollar["PDF"])
    if not os.path.exists(ana_klasor):
        os.makedirs(ana_klasor)
    return ana_klasor, ayar


def _personel_grup_tahmin_et(gorev):
    g = str(gorev).upper()
    if any(k in g for k in ["SÖZLEŞMELİ", "SOZLESMELI"]):
        return "Sözleşmeli"
    if any(k in g for k in ["ÜCRETLİ", "UCRETLI"]):
        return "Ücretli"
    if any(k in g for k in ["GÖREVLENDİRME", "GOREVLENDIRME"]):
        return "Görevlendirme"
    if any(k in g for k in ["SÜREKLİ İŞÇİ", "SUREKLI ISCI", "İŞÇİ", "ISCI"]):
        return "Sürekli İşçi"
    if "TYP" in g:
        return "TYP"
    if "İUP" in g or "IUP" in g:
        return "İUP"
    if any(k in g for k in ["HİZMET ALIMI", "HIZMET ALIMI"]):
        return "Hizmet Alımı"
    return "Kadrolu"


def _excel_sutunu_bul(sutunlar, adaylar):
    """Excel başlıklarındaki boşluk ve Türkçe büyük/küçük harf farklarını tolere eder."""
    normalize = lambda deger: re.sub(r"[^A-ZÇĞIÖŞÜ0-9]", "", str(deger).strip().upper().replace('İ', 'I'))
    aday_seti = {normalize(aday) for aday in adaylar}
    for sutun in sutunlar:
        if normalize(sutun) in aday_seti:
            return sutun
    return None


def _rapor_verisi_hazirla(tur, ozel_deger):
    ogrenci_listesi, devamsizlik_listesi = db.yukle()
    gruplu = {}
    for dev in devamsizlik_listesi:
        no = str(dev['no']).strip()
        gruplu.setdefault(no, []).append(dev)

    hedef_tarih_obj = None
    secilen_dev_turu = "TÜMÜ"
    if tur in ["tarih_bazli", "tarih_tur_bazli"]:
        tarih_str = ""
        if tur == "tarih_bazli":
            tarih_str = str(ozel_deger)
        elif tur == "tarih_tur_bazli":
            if isinstance(ozel_deger, dict):
                tarih_str = str(ozel_deger.get("tarih", ""))
                secilen_dev_turu = str(ozel_deger.get("tur", "TÜMÜ")).strip().upper()
            elif isinstance(ozel_deger, str) and "|" in ozel_deger:
                parcalar = ozel_deger.split("|")
                tarih_str = parcalar[0].strip()
                secilen_dev_turu = parcalar[1].strip().upper()
            else:
                tarih_str = str(ozel_deger).strip()

        try:
            t_temiz = VeriAraclari.tarih_formatla(tarih_str)
            gg, aa, yy = map(int, t_temiz.split('/'))
            hedef_tarih_obj = datetime(yy, aa, gg)
        except Exception as e:
            logging.error(f"_rapor_verisi_hazirla tarih formati hatasi: {e}")
            return None, "Tarih formatı hatalı! (GG/AA/YYYY olmalı)"

    veri = []
    for ogr in ogrenci_listesi:
        ogr_no = str(ogr['no']).strip()
        ogr_devleri = gruplu.get(ogr_no, [])
        ozsz, ozrl = VeriAraclari.hesapla_devamsizlik(ogr_no, ogr_devleri, [])
        gosterilen_sube = VeriAraclari.kisa_sube_adi(ogr['sube'])

        if tur == "ozursuz" and ozsz < 10: continue
        if tur == "ozurlu" and ozrl < 20: continue
        if tur == "gun_siniri" and ozsz < float(ozel_deger): continue
        if tur == "sube_bazli" and gosterilen_sube.replace(" ", "") != str(ozel_deger).replace(" ", ""): continue

        if tur == "tarih_bazli":
            bulundu = False
            dev_turu = ""
            for dev in ogr_devleri:
                if dev['tur'].upper() not in ["D", "ÖY", "SY"]:
                    continue
                try:
                    gun_mik = float(VeriAraclari.temiz_sure(dev['gun']))
                    tam_g = int(gun_mik) if gun_mik >= 1 else 1
                    bd, bm, by = map(int, VeriAraclari.tarih_formatla(dev['tarih']).split('/'))
                    bas_tarih = datetime(by, bm, bd)
                    for i in range(tam_g):
                        g_t = bas_tarih + timedelta(days=i)
                        if g_t.weekday() < 5 and g_t.date() == hedef_tarih_obj.date():
                            bulundu = True
                            dev_turu = dev['tur'].upper()
                            break
                    if bulundu:
                        break
                except Exception as e:
                    logging.error(f"_rapor_verisi_hazirla dongu hatasi: {e}")
                    pass
            if not bulundu:
                continue
            veri.append([gosterilen_sube, ogr_no, ogr['ad_soyad'], dev_turu])
        elif tur == "tarih_tur_bazli":
            bulundu = False
            bulunan_bilgi = ""
            for dev in ogr_devleri:
                dev_tur = str(dev.get('tur', '')).strip().upper()
                if secilen_dev_turu != "TÜMÜ":
                    if secilen_dev_turu in ["İ", "I"] and dev_tur in ["İ", "I"]:
                        eslesti = True
                    elif secilen_dev_turu == "S" and dev_tur in ["S", "SV"]:
                        eslesti = True
                    elif dev_tur == secilen_dev_turu:
                        eslesti = True
                    else:
                        eslesti = False
                    if not eslesti:
                        continue

                try:
                    gun_mik = float(VeriAraclari.temiz_sure(dev.get('gun', 1)))
                    tam_g = int(gun_mik) if gun_mik >= 1 else 1
                    t_str = VeriAraclari.tarih_formatla(dev['tarih'])
                    bd, bm, by = map(int, t_str.split('/'))
                    bas_tarih = datetime(by, bm, bd)
                    for i in range(tam_g):
                        g_t = bas_tarih + timedelta(days=i)
                        if g_t.weekday() < 5 and g_t.date() == hedef_tarih_obj.date():
                            bulundu = True
                            gun_gosterim = f"{gun_mik:g} Gün" if gun_mik != 1 else "1 Gün"
                            bulunan_bilgi = f"{dev_tur} ({gun_gosterim})"
                            break
                    if bulundu:
                        break
                except Exception as e:
                    logging.error(f"_rapor_verisi_hazirla tarih_tur dongu hatasi: {e}")
            if not bulundu:
                continue
            veri.append([gosterilen_sube, ogr_no, ogr['ad_soyad'], bulunan_bilgi])
        else:
            veri.append([gosterilen_sube, ogr_no, ogr['ad_soyad'], ozsz])

    if tur in ["tarih_bazli", "tarih_tur_bazli"]:
        def anahtar(x):
            r = re.findall(r'\d+', str(x[0]))
            return (int(r[0]) if r else 99, x[0], x[1])
        veri.sort(key=anahtar)
    else:
        veri.sort(key=lambda x: x[3], reverse=True)
    return veri, None



def evrak_kaydet(tur, personel_ad, brans):
    try:
        from dependencies import db
        from datetime import datetime
        db.cursor.execute("INSERT INTO uretilen_evraklar (tur, personel_ad, brans, tarih) VALUES (?, ?, ?, ?)", (tur, personel_ad, brans, datetime.now().isoformat()))
        db.conn.commit()
    except Exception as e:
        import logging
        logging.error(f"Evrak loglanamadi: {e}")
